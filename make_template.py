"""Generate Freight Loader Excel import template."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Packing List"

# Column layout matches freight_loader.py _load() parser:
#   A=row[0] Crate #, B=row[1] Part Number, C=row[2] (unused/description),
#   D=row[3] Length", E=row[4] Width", F=row[5] Height",
#   G=row[6] Weight lbs, H=row[7] PCS, I=row[8] (unused), J=row[9] Notes
headers = [
    "Crate #",       # A  row[0]
    "Part Number",   # B  row[1]
    "Description",   # C  row[2] — not imported, for reference only
    'Length (in)',   # D  row[3]
    'Width (in)',    # E  row[4]
    'Height (in)',   # F  row[5]
    "Weight (lbs)",  # G  row[6]
    "PCS",           # H  row[7]
    "Ref / PO #",    # I  row[8] — not imported, for reference only
    "Notes",         # J  row[9]
    "Stackable?",    # K  row[10] — Y/N or Yes/No; blank = Yes
]

col_widths = [10, 18, 28, 13, 13, 13, 14, 7, 16, 30, 12]

header_fill   = PatternFill("solid", fgColor="2A7D6E")
unused_fill   = PatternFill("solid", fgColor="C0A070")
header_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
unused_font   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
data_font     = Font(name="Calibri", size=11)
center        = Alignment(horizontal="center", vertical="center")
left          = Alignment(horizontal="left",   vertical="center")

thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Unused columns (C=3, I=9) get a different color to signal "not imported"
unused_cols = {3, 9}

for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font        = unused_font if col_idx in unused_cols else header_font
    cell.fill        = unused_fill if col_idx in unused_cols else header_fill
    cell.alignment   = center
    cell.border      = border
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 22

# Add a note row in row 2 (italic, gray) to explain unused columns
note_fill = PatternFill("solid", fgColor="FFF3CD")
note_font = Font(name="Calibri", italic=True, color="555555", size=10)
for col_idx in range(1, 12):
    cell = ws.cell(row=2, column=col_idx)
    cell.fill      = note_fill
    cell.font      = note_font
    cell.alignment = center
    cell.border    = border

ws.cell(row=2, column=3).value  = "← not imported (your ref)"
ws.cell(row=2, column=9).value  = "← not imported (your ref)"
ws.cell(row=2, column=1).value  = "← sample row below ↓"

# Sample data rows
sample_rows = [
    [1, "1234-FRAME-A",  "Main Frame Assembly",    96,  48,  60,  2400, 1, "PO-10001", "Fork pockets on long side", "Y"],
    [2, "5678-PANEL-B",  "Side Panel Set",         84,  48,  72,  1850, 2, "PO-10001", "Fragile — glass panels",    "N"],
    [3, "9012-BASE-C",   "Equipment Base Skid",    72,  48,  36,  3100, 1, "PO-10002", "",                          "Y"],
    [4, "3456-CRATE-D",  "Electrical Components",  48,  48,  48,   750, 4, "PO-10002", "This end up",               "Y"],
]

data_fill  = PatternFill("solid", fgColor="FFFFFF")
alt_fill   = PatternFill("solid", fgColor="F2F9F8")

for row_offset, row_data in enumerate(sample_rows):
    row_num = row_offset + 3
    fill = alt_fill if row_offset % 2 else data_fill
    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=row_num, column=col_idx, value=value)
        cell.font      = data_font
        cell.fill      = fill
        cell.alignment = center if col_idx != 3 else left
        cell.border    = border

ws.row_dimensions[2].height = 16
for r in range(3, 3 + len(sample_rows)):
    ws.row_dimensions[r].height = 18

ws.freeze_panes = "A3"

out = "Freight_Loader_Import_Template.xlsx"
wb.save(out)
print(f"Saved: {out}")
