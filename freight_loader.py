"""
Freight Loader — 53' Dry Van Load Planning Tool
Imports Excel packing lists, calculates balanced load plans, exports floor plan PDF and crate labels.
"""

APP_VERSION = "2.1.0"

import sys
import math
import datetime
import json
import os
import re
import shutil
import subprocess
import itertools
import tempfile
import time
from dataclasses import dataclass, field
from typing import Optional

SAVE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SAVE_DIR, "Freight_Loader_Import_Template.xlsx")
PACKING_LIST_DIR = os.path.join(SAVE_DIR, "Saved Load Templates")   # filled-in Excel packing lists
LOAD_PLANS_DIR = os.path.join(SAVE_DIR, "Load Plans")                # saved plans (.json) and PDFs
OLD_PACKING_LIST_DIR = os.path.join(SAVE_DIR, "Packing Lists")      # name used before v2.0


def ensure_folders():
    """
    Create the Load Plans folder and move the old "Packing Lists" folder to its
    new name. If Windows has it locked (e.g. a list is open in Excel) the old
    folder keeps being used and the rename is tried again next start.
    """
    global PACKING_LIST_DIR
    os.makedirs(LOAD_PLANS_DIR, exist_ok=True)
    if os.path.isdir(OLD_PACKING_LIST_DIR) and not os.path.exists(PACKING_LIST_DIR):
        try:
            os.rename(OLD_PACKING_LIST_DIR, PACKING_LIST_DIR)
        except OSError:
            PACKING_LIST_DIR = OLD_PACKING_LIST_DIR

try:
    from PyQt6.QtPdf import QPdfDocument
    from PyQt6.QtPdfWidgets import QPdfView
    HAS_QTPDF = True
except ImportError:          # older PyQt6 builds without the PDF module
    HAS_QTPDF = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QTableWidget, QTableWidgetItem, QPushButton, QLabel, QSpinBox,
    QLineEdit, QGroupBox, QSplitter, QScrollArea, QMessageBox,
    QHeaderView, QCheckBox, QSizePolicy, QDialog, QDialogButtonBox,
    QFileDialog, QFormLayout, QAbstractItemView, QFrame, QComboBox, QTabWidget,
    QButtonGroup, QToolTip
)
from PyQt6.QtCore import (Qt, QRectF, QPointF, QSize, QEvent, QSettings, QBuffer, QByteArray,
                          QIODevice, pyqtSignal)
from PyQt6.QtGui import (
    QPainter, QColor, QPen, QBrush, QFont, QPageLayout, QPageSize, QPdfWriter, QPolygonF,
    QFontMetrics, QCursor
)

# ── Colors ───────────────────────────────────────────────────────────────────
TEAL_DARK  = QColor("#2a7d6e")
TEAL_MED   = QColor("#3dab96")
TEAL_LIGHT = QColor("#cdeae5")
TAN        = QColor("#c0a070")
GRAY_LINE  = QColor("#b0b0b0")
VOID_BG    = QColor("#f8f8f8")

MAX_STOPS = 10
# Delivery-stop badge colors, indexed by stop - 1
STOP_COLORS = ["#1f4e79", "#7b2d8e", "#1e7b34", "#b03a2e", "#7a5c12",
               "#0e6e6e", "#a0522d", "#4b0082", "#556b2f", "#8b1c62"]

# Per-crate colors, indexed by piece_id (shared by top view and 3D view)
CRATE_COLORS = [
    "#4a90d9","#e67e22","#27ae60","#8e44ad","#c0392b","#16a085",
    "#d35400","#2980b9","#7f8c8d","#f39c12","#1abc9c","#e74c3c",
]

# ── 53' Dry Van specs ────────────────────────────────────────────────────────
TRAILER = {
    "name": "53' Dry Van",
    "interior_length_in": 636,
    "interior_width_in":  98,
    "interior_height_in": 110,
    "steer_limit": 12000,
    "drive_limit": 34000,
    "trailer_limit": 34000,
    "gross_limit":  80000,
}

# ── Tractor and trailer running gear (Truck & Trailer… dialog) ───────────────
# Empty weights are typical for a sleeper tractor and a 53' dry van; the
# trailer body's weight sits at its centre of gravity, the slider (axles,
# wheels, brakes) moves with the tandem. Positions are in inches.
RIG_DEFAULTS = {
    "steer_tare":        10800,   # tractor alone (bobtail), steer axle
    "drive_tare":         8700,   # tractor alone, drive tandem
    "wheelbase_in":        230,   # steer axle to drive tandem centre
    "fw_min_in":             0,   # fifth wheel ahead of drive tandem centre: slide range
    "fw_max_in":            24,
    "fw_step_in":            2,
    "fw_home_in":           12,   # usual position; auto stays here unless moving helps
    "kingpin_in":           36,   # kingpin behind the trailer nose
    "trailer_tare":      14500,   # empty trailer
    "slider_tare":        3000,   #   of which the sliding tandem assembly
    "trailer_cg_in":       270,   # empty body's centre of gravity behind the nose
    "tandem_min_in":       420,   # trailer tandem centre behind kingpin: slide range
    "tandem_max_in":       540,
    "tandem_step_in":        6,   # slider hole spacing
    "tandem_spread_in":     49,   # front to rear trailer axle
    "kpra_max_in":         480,   # kingpin to rear axle legal max (40' — CA; many states 41'); 0 = none
}
RIG = dict(RIG_DEFAULTS)


# ── Data model ────────────────────────────────────────────────────────────────
@dataclass
class FreightPiece:
    piece_id: int
    crate_label: str     # "Crate 9"
    part_number: str
    length_in: float
    width_in: float
    height_in: float
    weight_lbs: float
    pcs: int = 0
    notes: str = ""
    placement: str = "auto"
    stackable: bool = True
    stop: int = 1          # delivery stop ID (1..MAX_STOPS); see the shipment's stop list
    seq: int = 0           # delivery order on its truck, 1 = first (set by plan_load)

@dataclass
class ShipmentInfo:
    customer: str = ""
    shipment_name: str = ""
    trailer_num: str = ""
    bol_num: str = ""
    date: str = ""
    notes: str = ""
    stop_names: list = field(default_factory=lambda: [""] * MAX_STOPS)
    truck_label: str = ""      # e.g. "Truck 2 of 3" on PDFs (blank for a single truck)

@dataclass
class PalletSlot:
    row: int
    col: int                     # 0 = Position A / Left, 1 = Position B / Right
    piece: Optional[FreightPiece]
    row_start_in: float          # inches from nose to front of this row
    is_center: bool = False
    tier: int = 0                # 0 = floor level, 1 = stacked on top
    side: str = ""               # alt-center rows: "left"/"right" wall or "center"; "" = legacy
    x_in: Optional[float] = None # manual position of a floor crate (front edge from nose, in);
    y_in: Optional[float] = None #   left edge from the left wall — stacked crates follow it

@dataclass
class LoadPlan:
    slots: list
    axle_weights: dict
    violations: list
    total_length_used_in: float
    override: bool = False
    excluded_ids: list = None
    unplaced: list = None        # [(FreightPiece, reason)] that could not be loaded
    stop_axles: list = None      # [(stop just delivered, steer, drive, trailer)] for remaining load
    stop_names: list = None      # display names for stops 1..MAX_STOPS (set by the UI)
    stop_order: list = None      # stop IDs in delivery order (first delivered = at the doors)
    rig: dict = None             # {"fw_in", "tandem_in"}: fifth wheel and trailer tandem positions
    rig_auto: bool = True        # positions chosen by the planner (False = set by the user)


# ── Algorithm ─────────────────────────────────────────────────────────────────
# 1. Pieces that can't physically go in the trailer are set aside with a reason.
# 2. Each floor position is a "column": a bottom piece plus at most one piece
#    stacked on it. Wide (or forced-center) columns take a full row; narrow ones
#    pair two across, longest with next-longest, so short pieces don't sit
#    beside long ones wasting floor.
# 3. While the rows are longer than the trailer, stack the piece whose removal
#    from the floor saves the most length (stacking is a last resort).
# 4. If even maximum stacking won't fit, load the longest prefix of the freight
#    list that fits; the rest is reported as not loaded.
# 5. Order rows nose→tail to keep every axle legal and drive/trailer axles
#    evenly loaded, then pick each piece's wall to balance left/right weight.
# Multi-stop: rows are zoned by delivery stop — the last stop at the nose,
# stop 1 at the doors. Pieces only pair or stack with pieces for the same
# stop, rows only move within their zone, and axles are also checked for the
# load that remains after each delivery.
FIT_TOLERANCE_IN = 0.01


@dataclass
class _Column:
    """One floor position: a bottom piece and optionally one piece stacked on it."""
    bottom: FreightPiece
    top: Optional[FreightPiece] = None

    @property
    def length(self) -> float:
        return self.bottom.length_in

    @property
    def width(self) -> float:
        return min(self.bottom.width_in, TRAILER["interior_width_in"])

    @property
    def weight(self) -> float:
        return self.bottom.weight_lbs + (self.top.weight_lbs if self.top else 0.0)

    @property
    def seq(self) -> int:
        return self.bottom.seq

    @property
    def full_row(self) -> bool:
        return self.bottom.placement in ("center", "alt_center")

    @property
    def moment_in_row(self) -> float:
        """Σ weight × centre distance from the front of its row."""
        m = self.bottom.weight_lbs * self.bottom.length_in / 2
        if self.top:
            m += self.top.weight_lbs * self.top.length_in / 2
        return m


def _can_stack(bottom: FreightPiece, top: FreightPiece) -> bool:
    return (bottom.stackable
            and top.stop == bottom.stop
            and top.length_in <= bottom.length_in
            and top.width_in <= bottom.width_in
            and bottom.height_in + top.height_in <= TRAILER["interior_height_in"]
            and top.weight_lbs <= bottom.weight_lbs)


def _floor_rows(columns: list) -> list:
    """
    Group columns into rows (lists of 1–2 columns from the same stop). A stop's
    leftover unpaired column may share a row with the next stop's leftover —
    that row sits on the zone boundary, so either piece unloads without moving
    the other.
    """
    W = TRAILER["interior_width_in"]
    rows = [[c] for c in columns if c.full_row]
    stops = sorted({c.seq for c in columns})
    leftover = {}
    for stop in stops:
        narrow = sorted((c for c in columns if not c.full_row and c.seq == stop),
                        key=lambda c: c.length, reverse=True)
        i = 0
        while i < len(narrow):
            a = narrow[i]
            b = narrow[i + 1] if i + 1 < len(narrow) else None
            if b is not None and a.width + b.width <= W:
                rows.append([a, b])
                i += 2
            else:
                leftover.setdefault(stop, []).append(a)
                i += 1
    k = 0
    while k < len(stops):
        st = stops[k]
        nxt = stops[k + 1] if k + 1 < len(stops) else None
        if (st in leftover and nxt in leftover and len(leftover[st]) == 1
                and len(leftover[nxt]) == 1
                and leftover[st][0].width + leftover[nxt][0].width <= W):
            rows.append([leftover.pop(st)[0], leftover.pop(nxt)[0]])
            k += 2
        else:
            k += 1
    for cols in leftover.values():
        rows.extend([c] for c in cols)
    return rows


def _row_stop(row: list) -> float:
    """Zone key for a row; a boundary row shared by two stops sits between them."""
    return sum(c.seq for c in row) / len(row)


def _rows_length(rows: list) -> float:
    return sum(max(c.length for c in r) for r in rows)


def _pack_columns(pieces: list) -> tuple:
    """Columns for `pieces` with the least stacking needed to fit; (columns, fits)."""
    L = TRAILER["interior_length_in"]
    columns = [_Column(p) for p in pieces]
    current = _rows_length(_floor_rows(columns))
    while current > L + FIT_TOLERANCE_IN:
        best = None
        for t in columns:
            if t.top is not None:
                continue   # a column already carrying a piece can't go on top (no 3-high)
            bottoms = [b for b in columns
                       if b is not t and b.top is None and _can_stack(b.bottom, t.bottom)]
            if not bottoms:
                continue
            rest_len = _rows_length(_floor_rows([c for c in columns if c is not t]))
            gain = current - rest_len
            # Most floor saved first, then lighter tops (lower centre of gravity)
            key = (round(gain, 3), -t.bottom.weight_lbs)
            if best is None or key > best[0]:
                # Smallest suitable base keeps big bases free for big tops
                base = min(bottoms, key=lambda b: (b.length * b.width, -b.bottom.weight_lbs))
                best = (key, t, base, rest_len)
        if best is None:
            return columns, False
        _key, t, base, rest_len = best
        base.top = t.bottom
        columns.remove(t)
        current = rest_len
    return columns, True


# ── Axle weights ──────────────────────────────────────────────────────────────
# Two beams. The trailer rests on the kingpin and its tandem centre and carries
# the freight plus its own body; the slider's weight sits straight on the
# tandem. The kingpin load rides the fifth wheel, which sits fw_in ahead of
# the drive tandem centre, so it splits between drive and steer by lever arm:
# sliding the fifth wheel forward moves weight onto the steer axle, sliding
# the trailer tandem back moves weight from the trailer axles onto the drives.
def _rear_axle_ok(tandem_in: float) -> bool:
    """Kingpin-to-rear-axle distance within the legal limit (if one is set)."""
    return (not RIG["kpra_max_in"]
            or tandem_in + RIG["tandem_spread_in"] / 2 <= RIG["kpra_max_in"] + FIT_TOLERANCE_IN)


def rig_positions(legal_only: bool = True) -> tuple:
    """(fifth wheel notches, tandem holes) in inches, from RIG."""
    fw = [RIG["fw_min_in"] + k * RIG["fw_step_in"]
          for k in range(int((RIG["fw_max_in"] - RIG["fw_min_in"]) / RIG["fw_step_in"]) + 1)]
    holes = [RIG["tandem_min_in"] + k * RIG["tandem_step_in"]
             for k in range(int((RIG["tandem_max_in"] - RIG["tandem_min_in"]) / RIG["tandem_step_in"]) + 1)]
    if not legal_only:
        return fw, holes
    return fw, [t for t in holes if _rear_axle_ok(t)] or holes[:1]


def default_rig() -> dict:
    """Usual fifth wheel spot, tandem as far back as the kingpin-to-rear-axle limit allows."""
    fw, holes = rig_positions()
    return {"fw_in": min(fw, key=lambda f: abs(f - RIG["fw_home_in"])), "tandem_in": holes[-1]}


def rig_tare() -> float:
    """Empty tractor + trailer weight."""
    return RIG["steer_tare"] + RIG["drive_tare"] + RIG["trailer_tare"]


def max_payload() -> float:
    """Freight weight that brings the rig to the gross limit."""
    return TRAILER["gross_limit"] - rig_tare()


def axle_marks(plan: Optional["LoadPlan"]) -> list:
    """[(inches from nose, label, colour, kind)] for the kingpin, drive and trailer tandem."""
    rig = plan.rig if plan and plan.rig else default_rig()
    kp = RIG["kingpin_in"]
    return [(kp, "Kingpin", "#666666", "kingpin"),
            (kp + rig["fw_in"], "Drive axles", "#e67e00", "drive"),
            (kp + rig["tandem_in"], "Trailer axles", "#c0392b", "tandem")]


def rig_text(rig: dict) -> str:
    """'5th wheel 12" ahead of drives · Tandem 450" behind kingpin (rear axle 39'7")'."""
    rear = rig["tandem_in"] + RIG["tandem_spread_in"] / 2
    return (f"5th wheel {rig['fw_in']:.0f}\" ahead of drives  ·  Tandem {rig['tandem_in']:.0f}\" "
            f"behind kingpin (rear axle {int(rear // 12)}'{rear % 12:.0f}\")")


def _axles_for(total: float, moment: float, rig: dict) -> tuple:
    """(steer, drive, trailer) for freight `total` lb with `moment` (lb·in about the nose)."""
    kp, span = RIG["kingpin_in"], rig["tandem_in"]
    body = RIG["trailer_tare"] - RIG["slider_tare"]
    about_kp = moment - total * kp + body * (RIG["trailer_cg_in"] - kp)
    on_tandem = about_kp / span
    pin = total + body - on_tandem
    wb = RIG["wheelbase_in"]
    steer = RIG["steer_tare"] + pin * rig["fw_in"] / wb
    drive = RIG["drive_tare"] + pin * (wb - rig["fw_in"]) / wb
    return steer, drive, on_tandem + RIG["slider_tare"]


def _overload(steer: float, drive: float, trailer: float) -> float:
    return (max(0, steer - TRAILER["steer_limit"]) + max(0, drive - TRAILER["drive_limit"])
            + max(0, trailer - TRAILER["trailer_limit"]))


AUTO_RIG_COMFORT = 0.90   # axles under 90% of their limit don't need the rig moved


def auto_rig(loads: list) -> dict:
    """
    Fifth wheel and tandem positions for loads [(total, moment)] — the full
    load first, then what remains after each delivery. Keeps every axle legal
    if possible; when drive or trailer tandem is above AUTO_RIG_COMFORT of its
    limit, slides to bring the heavier one down; otherwise leaves the fifth
    wheel at its usual spot and the tandem as far back as allowed.
    """
    fw_opts, holes = rig_positions()
    best = None
    for t in holes:
        for f in fw_opts:
            rig = {"fw_in": f, "tandem_in": t}
            over = 0.0
            for total, moment in loads:
                over += _overload(*_axles_for(total, moment, rig))
            _s, drive, trailer = _axles_for(*loads[0], rig)
            peak = max(drive / TRAILER["drive_limit"], trailer / TRAILER["trailer_limit"])
            key = (round(over), round(max(0.0, peak - AUTO_RIG_COMFORT) * 200),
                   abs(f - RIG["fw_home_in"]), -t)
            if best is None or key < best[0]:
                best = (key, rig)
    return best[1]


def _order_loads(rows: list, order: list) -> list:
    """
    [(total, moment about the nose)] for rows laid nose→tail in `order`: the
    full load, then what remains after each delivery.
    rows[i] = (depth, weight, inner_moment, stop, stacked).
    """
    start = 0.0
    by_stop = {}                      # stop -> [weight, moment]
    for i in order:
        depth, weight, inner, stop, _stacked = rows[i]
        acc = by_stop.setdefault(stop, [0.0, 0.0])
        acc[0] += weight
        acc[1] += weight * start + inner
        start += depth
    out = []
    for delivered in [None] + sorted(by_stop)[:-1]:
        rem = [a for st, a in by_stop.items() if delivered is None or st > delivered]
        out.append((sum(a[0] for a in rem), sum(a[1] for a in rem)))
    return out


def _order_score(rows: list, order: list, rig: dict) -> float:
    """
    Lower is better: axle overload dominates (full load, and what remains after
    each delivery stop), then drive/trailer imbalance of the full load.
    """
    loads = _order_loads(rows, order)
    steer, drive, trailer = _axles_for(*loads[0], rig)
    over = sum(_overload(*_axles_for(t, m, rig)) for t, m in loads)
    balance = abs(drive / TRAILER["drive_limit"] - trailer / TRAILER["trailer_limit"])
    return over * 10 + balance * 10_000


def _order_overload(rows: list, order: list, rig: dict, auto: bool) -> float:
    """Pounds over any axle limit for this order, with the rig slid (auto) or fixed."""
    loads = _order_loads(rows, order)
    if auto:
        rig = auto_rig(loads)
    return sum(_overload(*_axles_for(t, m, rig)) for t, m in loads)


def _zone_ok(rows: list, order: list) -> bool:
    """Stops never increase toward the doors (last stop at the nose)."""
    return all(rows[order[k]][3] >= rows[order[k + 1]][3] for k in range(len(order) - 1))


def _stack_block_zone(rows: list, zone: list, front: int) -> list:
    """
    One stop zone nose→doors: `front` single rows, the block of stacked rows,
    then the remaining singles. Heaviest rows sit next to / in the middle of
    the block.
    """
    singles = sorted((i for i in zone if not rows[i][4]), key=lambda i: rows[i][1], reverse=True)
    doubles = sorted((i for i in zone if rows[i][4]), key=lambda i: rows[i][1], reverse=True)
    back = len(singles) - front
    # Single positions nearest the block first (ties: nose side)
    spots = sorted([("f", k) for k in range(front)] + [("b", k) for k in range(back)],
                   key=lambda sp: (front - sp[1] if sp[0] == "f" else sp[1] + 1, sp[0] == "b"))
    front_rows, back_rows = [None] * front, [None] * back
    for (side, k), idx in zip(spots, singles):
        (front_rows if side == "f" else back_rows)[k] = idx
    # Doubles: heaviest in the middle of the block
    m = len(doubles)
    mid_out = sorted(range(m), key=lambda k: (abs(k - (m - 1) / 2), k))
    block = [None] * m
    for k, idx in zip(mid_out, doubles):
        block[k] = idx
    return front_rows + block + back_rows


def _improve_order(rows: list, order: list, rig: dict, keep_pattern: bool) -> list:
    """
    Swaps and moves within each stop zone while they improve _order_score.
    keep_pattern: only changes that leave every stacked row where it is in the
    single / double pattern (so the stack block stays together and in place).
    """
    n = len(order)
    pattern = [rows[i][4] for i in order]

    def ok(cand):
        return _zone_ok(rows, cand) and (not keep_pattern or
                                         all(rows[i][4] == pt for i, pt in zip(cand, pattern)))
    best = _order_score(rows, order, rig)
    for _ in range(50):
        improved = False
        for i in range(n):
            for j in range(i + 1, n):
                if rows[order[i]][3] != rows[order[j]][3]:
                    continue   # rows only trade places within their stop's zone
                cand = order[:]
                cand[i], cand[j] = cand[j], cand[i]
                if not ok(cand):
                    continue
                sc = _order_score(rows, cand, rig)
                if sc < best - 1e-9:
                    order, best, improved = cand, sc, True
        for i in range(n):
            for j in range(n):
                if i == j:
                    continue
                cand = order[:]
                cand.insert(j, cand.pop(i))
                if not ok(cand):
                    continue
                sc = _order_score(rows, cand, rig)
                if sc < best - 1e-9:
                    order, best, improved = cand, sc, True
        if not improved:
            break
    return order


def _balanced_order(rows: list, rig: dict, auto: bool = True) -> list:
    """
    Row order nose→tail, the way a loader would do it. Zones by stop (last
    stop at the nose). In each zone the stacked rows form one block in the
    middle; with an odd number of single rows the extra one goes behind the
    block (e.g. 12 rows, 3 doubles: 4 singles, 3 doubles, 5 singles). Rows
    then trade places among singles / among doubles to even out the axles.
    If that isn't legal (with the rig slid when `auto`), the block moves the
    fewest rows needed; failing that, stacks may split up wherever the axles
    need them.
    """
    n = len(rows)
    if n < 2:
        return list(range(n))
    zones = [[i for i in range(n) if rows[i][3] == stop]
             for stop in sorted({r[3] for r in rows}, reverse=True)]
    singles = [sum(1 for i in z if not rows[i][4]) for z in zones]
    preferred = [s // 2 for s in singles]

    def build(fronts):
        order = []
        for z, f in zip(zones, fronts):
            order += _stack_block_zone(rows, z, f)
        return _improve_order(rows, order, rig, keep_pattern=True)

    fronts = list(preferred)
    order = build(fronts)
    over = _order_overload(rows, order, rig, auto)
    if over > 0:
        # Slide one zone's block at a time, nearest positions first (nose side on ties)
        improved = True
        while improved and over > 0:
            improved = False
            for zi in range(len(zones)):
                for f in sorted(range(singles[zi] + 1), key=lambda f: (abs(f - preferred[zi]), f)):
                    if f == fronts[zi]:
                        continue
                    trial = fronts[:zi] + [f] + fronts[zi + 1:]
                    cand = build(trial)
                    c_over = _order_overload(rows, cand, rig, auto)
                    if c_over < over - 0.5:
                        fronts, order, over, improved = trial, cand, c_over, True
                        break
    if over > 0:
        # Still overweight with the block anywhere: let the stacks split up
        cand = _improve_order(rows, order, rig, keep_pattern=False)
        if _order_overload(rows, cand, rig, auto) < over - 0.5:
            order = cand
    return order


def _relayout_rows(slots: list) -> float:
    """Recompute each row's start from the floor pieces' lengths; returns total length."""
    depth = {}
    for s in slots:
        if s.piece and s.tier == 0:
            depth[s.row] = max(depth.get(s.row, 0.0), s.piece.length_in)
    start, cursor = {}, 0.0
    for r in sorted(depth):
        start[r] = cursor
        cursor += depth[r]
    for s in slots:
        if s.row in start:
            s.row_start_in = start[s.row]
    return cursor


def plan_load(pieces: list, override: bool = False, excluded_ids: list = None,
              stop_order: list = None, rig: dict = None) -> LoadPlan:
    """
    Build a load plan that fits the trailer, stacks only as needed and keeps
    axle and side-to-side weight balanced. Pieces that won't fit are listed in
    LoadPlan.unplaced and called out in the violations. `rig` fixes the fifth
    wheel and tandem positions; None lets the planner place them.
    """
    excluded_ids = excluded_ids or []
    L, W, H = (TRAILER["interior_length_in"], TRAILER["interior_width_in"],
               TRAILER["interior_height_in"])
    active = [p for p in pieces if p.piece_id not in excluded_ids]
    if not active:
        return LoadPlan([], {}, [], 0.0, override, excluded_ids, [], [],
                        rig=dict(rig or default_rig()), rig_auto=rig is None)

    # Delivery order: the given route order, then any other stops by ID
    rank = {sid: i + 1 for i, sid in enumerate(stop_order or [])}
    for p in active:
        p.seq = rank.get(p.stop, len(rank) + p.stop)

    # Wide pieces can't pair: >70% of the width alternates against a wall, >50% centres
    for p in active:
        if p.width_in > W * 0.70:
            p.placement = "alt_center"
        elif p.width_in > W / 2:
            p.placement = "center"

    unplaced, loadable = [], []
    for p in active:
        if p.length_in > L:
            unplaced.append((p, f"longer than trailer ({p.length_in:.0f}\" > {L}\")"))
        elif p.width_in > W:
            unplaced.append((p, f"wider than trailer ({p.width_in:.0f}\" > {W}\")"))
        elif p.height_in > H:
            unplaced.append((p, f"taller than trailer ({p.height_in:.0f}\" > {H}\")"))
        else:
            loadable.append(p)

    columns, fits = _pack_columns(loadable)
    if not fits:
        # Load in delivery order, then list order, so what's left off comes from
        # the last-delivered stop — never a scattering of every stop's freight
        priority = sorted(loadable, key=lambda p: p.seq)   # stable: keeps list order
        lo, hi, best_cols = 0, len(priority), []
        while lo < hi:
            mid = (lo + hi + 1) // 2
            cols, ok = _pack_columns(priority[:mid])
            if ok:
                lo, best_cols = mid, cols
            else:
                hi = mid - 1
        columns = best_cols
        placed_ids = {c.bottom.piece_id for c in columns} | \
                     {c.top.piece_id for c in columns if c.top}
        for p in loadable:
            if p.piece_id not in placed_ids:
                unplaced.append((p, "no room left on the trailer"))

    rows = _floor_rows(columns)
    row_stats = [(max(c.length for c in r), sum(c.weight for c in r),
                  sum(c.moment_in_row for c in r), _row_stop(r),
                  any(c.top is not None for c in r)) for r in rows]
    order = _balanced_order(row_stats, rig or default_rig(), auto=rig is None)

    # Choose each row's walls to balance left/right weight, then build slots nose→tail
    ordered_rows = [rows[ri] for ri in order]
    sides = _balance_sides(ordered_rows)
    slots, cursor = [], 0.0
    for row_idx, (row, state) in enumerate(zip(ordered_rows, sides)):
        if len(row) == 1 and row[0].full_row:
            c = row[0]
            side = ""
            if c.bottom.placement == "alt_center":
                side = {LEFT: "left", RIGHT: "right", CENTER: "center"}[state]
            for tier, piece in ((0, c.bottom), (1, c.top)):
                if piece is None:
                    continue
                for col in (0, 1):
                    slots.append(PalletSlot(row=row_idx, col=col, piece=piece,
                                            row_start_in=cursor, is_center=True,
                                            tier=tier, side=side))
        else:
            left, right = _row_walls(row, state == RIGHT)
            for col, c in ((0, left), (1, right)):
                slots.append(PalletSlot(row=row_idx, col=col, piece=c.bottom if c else None,
                                        row_start_in=cursor, tier=0))
                if c and c.top:
                    slots.append(PalletSlot(row=row_idx, col=col, piece=c.top,
                                            row_start_in=cursor, tier=1))
        cursor += max(c.length for c in row)

    order = sorted({p.stop for p in active}, key=lambda sid: rank.get(sid, len(rank) + sid))
    plan = LoadPlan(slots, {}, [], cursor, override, excluded_ids, unplaced, [],
                    stop_order=order, rig=dict(rig) if rig else None, rig_auto=rig is None)
    _refresh_plan_metrics(plan)
    return plan


# Row side states: LEFT = heavier column / wall-hugging crate on the left wall,
# RIGHT = mirrored, CENTER = wall-hugging (alt-center) crate centred in the width.
LEFT, RIGHT, CENTER = 0, 1, 2


def _row_walls(row: list, flip) -> tuple:
    """(left column, right column) for a pair or lone-column row; flip swaps sides."""
    cols = sorted(row, key=lambda c: c.weight, reverse=True)
    left, right = (cols[0], cols[1] if len(cols) > 1 else None)
    return (right, left) if flip else (left, right)


def _is_wall_crate(row: list) -> bool:
    """A single crate too wide to pair, carried against one wall (alt-center)."""
    return len(row) == 1 and row[0].full_row and row[0].bottom.placement == "alt_center"


def _row_lateral(row: list, state: int) -> list:
    """[(seq, weight × offset right of centre)] for every piece in a row, as placed."""
    W = TRAILER["interior_width_in"]
    out = []
    if len(row) == 1 and row[0].full_row:
        c = row[0]
        if c.bottom.placement != "alt_center" or state == CENTER:
            return []                                   # centred: no lateral moment
        for p in (c.bottom, c.top):
            if p is not None:
                off = (W - min(p.width_in, W)) / 2
                out.append((p.seq, p.weight_lbs * (off if state == RIGHT else -off)))
        return out
    left, right = _row_walls(row, state == RIGHT)
    for col, c in ((0, left), (1, right)):
        if c is None:
            continue
        for p in (c.bottom, c.top):
            if p is not None:
                w = min(p.width_in, W)
                yc = w / 2 if col == 0 else W - w / 2   # each piece sits against its wall
                out.append((p.seq, p.weight_lbs * (yc - W / 2)))
    return out


LATERAL_AFTER_STOP_WEIGHT = 0.5   # balance after each delivery counts half as much as at departure
EXHAUSTIVE_SIDE_COMBOS = 4096     # side combinations tried outright before switching to local search


def _balance_sides(rows: list) -> list:
    """
    Side state per row (LEFT / RIGHT / CENTER) minimising left/right imbalance
    of the full load plus, at half weight, the load left after each delivery.

    Consecutive wall-hugging crates must alternate walls row by row — the
    correct way to load them — so a run only chooses its starting wall, and
    may centre its tail-most crate when that evens out the count. Side-by-side
    rows choose freely which column goes on which wall.
    """
    n = len(rows)
    wall = [_is_wall_crate(r) for r in rows]
    states = [(LEFT, RIGHT, CENTER) if wall[i] else (LEFT, RIGHT) for i in range(n)]
    lateral = [{st: _row_lateral(rows[i], st) for st in states[i]} for i in range(n)]
    seqs = sorted({sq for lat in lateral for moments in lat.values() for sq, _ in moments})
    thresholds = seqs[:-1]                       # deliveries after which freight remains
    # vec[i][state] = row i's moment for [full load, after delivery 1, after delivery 2, …]
    vec = [{st: [sum(m for _, m in moments)] +
                [sum(m for sq, m in moments if sq > th) for th in thresholds]
            for st, moments in lat.items()} for lat in lateral]
    K = len(thresholds) + 1

    def cost(flags):
        acc = [0.0] * K
        for v, f in zip(vec, flags):
            row_v = v[f]
            for k in range(K):
                acc[k] += row_v[k]
        return abs(acc[0]) + LATERAL_AFTER_STOP_WEIGHT * sum(abs(a) for a in acc[1:])

    # Decision groups: each run of wall crates (alternating pattern options) and
    # each side-by-side row whose side matters (left / right)
    groups = []                                   # [(row indices, [state tuples])]
    i = 0
    while i < n:
        if wall[i]:
            j = i
            while j + 1 < n and wall[j + 1]:
                j += 1
            run = list(range(i, j + 1))
            pats = [tuple((start + k) % 2 for k in range(len(run))) for start in (RIGHT, LEFT)]
            pats += [pat[:-1] + (CENTER,) for pat in pats]   # centring wins only if strictly better
            groups.append((run, pats))
            i = j + 1
        else:
            if vec[i][LEFT] != vec[i][RIGHT]:
                groups.append(([i], [(LEFT,), (RIGHT,)]))
            i += 1

    def apply(choice):
        flags = [LEFT] * n
        for (idx, pats), k in zip(groups, choice):
            for r, st in zip(idx, pats[k]):
                flags[r] = st
        return flags

    combos = 1
    for _, pats in groups:
        combos *= len(pats)
    if combos <= EXHAUSTIVE_SIDE_COMBOS:
        best, best_flags = None, apply([0] * len(groups))
        for choice in itertools.product(*[range(len(pats)) for _, pats in groups]):
            trial = apply(choice)
            c = cost(trial)
            if best is None or c < best - 1e-9:
                best, best_flags = c, trial
        return best_flags
    # Long loads: start from each group's first option, then change one or two
    # groups at a time while that improves balance
    choice = [0] * len(groups)
    best = cost(apply(choice))
    improved = True
    while improved:
        improved = False
        for g in range(len(groups)):
            for k in range(len(groups[g][1])):
                if k != choice[g]:
                    trial = choice[:]
                    trial[g] = k
                    c = cost(apply(trial))
                    if c < best - 1e-9:
                        best, choice, improved = c, trial, True
        for a in range(len(groups)):
            for b in range(a + 1, len(groups)):
                for ka in range(len(groups[a][1])):
                    for kb in range(len(groups[b][1])):
                        trial = choice[:]
                        trial[a], trial[b] = ka, kb
                        c = cost(apply(trial))
                        if c < best - 1e-9:
                            best, choice, improved = c, trial, True
    return apply(choice)


def _freight_loads(slots: list) -> list:
    """
    [(stop delivered or None, total lb, moment lb·in about the nose)]: the full
    load first, then the load left on the trailer after each delivery but the last.
    """
    x_of = {b[0].piece_id: b[1] for b in _slot_boxes(slots)}   # honours manual moves
    pieces = [s.piece for s in slots if s.piece and not (s.is_center and s.col == 1)]
    stop_of = {p.seq: p.stop for p in pieces}
    out = []
    for delivered in [0] + sorted(stop_of)[:-1]:
        rest = [p for p in pieces if p.seq > delivered]
        out.append((stop_of.get(delivered),
                    sum(p.weight_lbs for p in rest),
                    sum(p.weight_lbs * (x_of[p.piece_id] + p.length_in / 2) for p in rest)))
    return out


def _stop_axles(slots: list, rig: dict) -> list:
    """[(stop just delivered, steer, drive, trailer)] for the load left after each delivery."""
    return [(stop, *(round(v) for v in _axles_for(total, moment, rig)))
            for stop, total, moment in _freight_loads(slots)[1:]]


def plan_stops(plan: Optional[LoadPlan]) -> list:
    """Stop IDs present in a plan, in delivery order (first delivered first)."""
    if not plan or not plan.slots:
        return []
    present = {s.piece.stop for s in plan.slots if s.piece}
    order = [sid for sid in (plan.stop_order or []) if sid in present]
    return order + sorted(present - set(order))


def stop_label(plan: Optional[LoadPlan], stop: int, short: bool = False) -> str:
    """'S2' / 'Stop 2 — Coldwater, MI' using the plan's stop names."""
    if short:
        return f"S{stop}"
    names = (plan.stop_names if plan and plan.stop_names else None) or []
    name = names[stop - 1] if stop - 1 < len(names) else ""
    return f"Stop {stop}" + (f" — {name}" if name else "")


def _not_fit_message(unplaced: list) -> str:
    wt = sum(p.weight_lbs for p, _ in unplaced)
    names = [p.crate_label for p, _ in unplaced]
    labels = ", ".join(names[:10]) + (f" +{len(names) - 10} more" if len(names) > 10 else "")
    return (f"DOES NOT FIT: {len(unplaced)} piece(s), {wt:,.0f} lb not loaded — {labels}")


def _calc_axle_weights(slots: list, override: bool, rig: dict) -> tuple:
    """
    Axle weights of the loaded rig (see _axles_for) with the fifth wheel and
    tandem at `rig`, plus weight and stacked-height violations. Stacked pieces
    count at their own centre, like floor pieces.
    """
    trailer_h  = TRAILER["interior_height_in"]
    _stop, total_cargo, moment = _freight_loads(slots)[0]
    if total_cargo == 0:
        return {}, []
    steer, drive, trailer = _axles_for(total_cargo, moment, rig)
    aw = {
        "steer":         round(steer),
        "drive":         round(drive),
        "trailer":       round(trailer),
        "gross":         round(steer + drive + trailer),
        "freight_gross": round(total_cargo),
    }

    v = []
    if aw["steer"]   > TRAILER["steer_limit"]:
        v.append(f"STEER {aw['steer']:,} lb > {TRAILER['steer_limit']:,} lb limit")
    if aw["drive"]   > TRAILER["drive_limit"]:
        v.append(f"DRIVE {aw['drive']:,} lb > {TRAILER['drive_limit']:,} lb limit")
    if aw["trailer"] > TRAILER["trailer_limit"]:
        v.append(f"TRAILER TANDEM {aw['trailer']:,} lb > {TRAILER['trailer_limit']:,} lb limit")
    if aw["gross"]   > TRAILER["gross_limit"]:
        v.append(f"GROSS {aw['gross']:,} lb > {TRAILER['gross_limit']:,} lb limit")

    stack_map: dict = {}   # (row, col) -> {tier: piece}
    for slot in slots:
        if slot.piece and not (slot.is_center and slot.col == 1):
            stack_map.setdefault((slot.row, slot.col), {})[slot.tier] = slot.piece
    # Height clearance check for stacked positions
    for (row, col), tiers in stack_map.items():
        if 0 in tiers and 1 in tiers:
            combined = tiers[0].height_in + tiers[1].height_in
            if combined > trailer_h:
                v.append(
                    f"HEIGHT: {tiers[0].crate_label} + {tiers[1].crate_label} = "
                    f"{combined:.0f}\" > {trailer_h}\" interior limit"
                )

    return aw, v


# ── Stops, locations and trucks ──────────────────────────────────────────────
# A shipment has up to MAX_STOPS delivery stops (each with an address or GPS
# coordinates) and an origin. Stops are routed nearest-first from the origin;
# freight is split into trucks along that route, one stop per truck where
# possible. When a stop won't fit the user chooses: move the whole stop to a
# new truck, or load what fits and ship the rest another way.
NOMINATIM_URL = "https://nominatim.openstreetmap.org/search"
OSRM_TABLE_URL = "https://router.project-osrm.org/table/v1/driving/"
HTTP_USER_AGENT = "FreightLoader/1.0 (desktop load planner)"


def parse_coords(text: str) -> Optional[tuple]:
    """
    (lat, lon) from typed GPS text, or None if it isn't coordinates.
    Accepts "42.97, -85.67", "42.97 -85.67", "42.97°N 85.67°W" and Google Maps
    links containing "@42.97,-85.67".
    """
    if not text:
        return None
    t = text.strip()
    m = re.search(r"@(-?\d+(?:\.\d+)?),\s*(-?\d+(?:\.\d+)?)", t)
    if m:
        lat, lon = float(m.group(1)), float(m.group(2))
    else:
        m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*°?\s*([NSns])?\s*[,;\s]\s*"
                         r"(-?\d+(?:\.\d+)?)\s*°?\s*([EWew])?\s*", t)
        if not m:
            return None
        lat, lon = float(m.group(1)), float(m.group(3))
        if m.group(2) and m.group(2).upper() == "S":
            lat = -abs(lat)
        if m.group(4) and m.group(4).upper() == "W":
            lon = -abs(lon)
    if -90 <= lat <= 90 and -180 <= lon <= 180:
        return lat, lon
    return None


def geocode_address(address: str) -> Optional[tuple]:
    """(lat, lon, display_name) from OpenStreetMap Nominatim, or None. Needs internet."""
    import urllib.parse
    import urllib.request
    url = NOMINATIM_URL + "?" + urllib.parse.urlencode(
        {"q": address, "format": "json", "limit": 1, "countrycodes": "us,ca,mx"})
    req = urllib.request.Request(url, headers={"User-Agent": HTTP_USER_AGENT})
    with urllib.request.urlopen(req, timeout=15) as resp:
        results = json.loads(resp.read().decode("utf-8"))
    if not results:
        return None
    r = results[0]
    return float(r["lat"]), float(r["lon"]), r.get("display_name", "")


def road_distance_matrix(points: list) -> list:
    """Road miles between every pair of (lat, lon) points via the public OSRM server."""
    import urllib.request
    coords = ";".join(f"{lon:.6f},{lat:.6f}" for lat, lon in points)
    req = urllib.request.Request(OSRM_TABLE_URL + coords + "?annotations=distance",
                                 headers={"User-Agent": HTTP_USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    if data.get("code") != "Ok":
        raise RuntimeError(data.get("message", "routing server error"))
    return [[(d or 0.0) / 1609.344 for d in row] for row in data["distances"]]


def straight_miles(a: tuple, b: tuple) -> float:
    """Great-circle distance in miles between two (lat, lon) points."""
    lat1, lon1, lat2, lon2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = (math.sin((lat2 - lat1) / 2) ** 2 +
         math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2)
    return 3958.8 * 2 * math.asin(math.sqrt(h))


def route_order(stop_ids: list, coords: dict, dist=None, origin_key="origin") -> list:
    """
    Delivery order for `stop_ids`: nearest-neighbour from the origin, then
    2-opt to remove crossings. `coords` maps "origin" and stop IDs to (lat, lon);
    `dist(a_key, b_key)` gives miles (straight-line by default). Stops without
    a location keep their ID order after the located ones; with no origin the
    order is simply by stop ID.
    """
    if dist is None:
        def dist(a, b):
            return straight_miles(coords[a], coords[b])
    located = [sid for sid in stop_ids if sid in coords]
    unlocated = sorted(sid for sid in stop_ids if sid not in coords)
    if origin_key not in coords or len(located) < 2:
        return sorted(located) + unlocated if origin_key not in coords else \
            located + unlocated
    path, cur, left = [], origin_key, set(located)
    while left:
        nxt = min(left, key=lambda sid: (dist(cur, sid), sid))
        path.append(nxt)
        left.remove(nxt)
        cur = nxt

    def length(seq):
        pts = [origin_key] + seq
        return sum(dist(pts[k], pts[k + 1]) for k in range(len(pts) - 1))
    best = length(path)
    improved = True
    while improved:
        improved = False
        for i in range(len(path) - 1):
            for j in range(i + 1, len(path)):
                cand = path[:i] + path[i:j + 1][::-1] + path[j + 1:]
                cl = length(cand)
                if cl < best - 1e-9:
                    path, best, improved = cand, cl, True
    return path + unlocated


def _normalize_placement(p: FreightPiece):
    W = TRAILER["interior_width_in"]
    if p.width_in > W * 0.70:
        p.placement = "alt_center"
    elif p.width_in > W / 2:
        p.placement = "center"


def physically_loadable(p: FreightPiece) -> Optional[str]:
    """None if the piece can go in the trailer at all, else the reason it can't."""
    L, W, H = (TRAILER["interior_length_in"], TRAILER["interior_width_in"],
               TRAILER["interior_height_in"])
    if p.length_in > L:
        return f"longer than trailer ({p.length_in:.0f}\" > {L}\")"
    if p.width_in > W:
        return f"wider than trailer ({p.width_in:.0f}\" > {W}\")"
    if p.height_in > H:
        return f"taller than trailer ({p.height_in:.0f}\" > {H}\")"
    if p.weight_lbs > max_payload():
        return f"heavier than a truck can legally carry ({p.weight_lbs:,.0f} lb)"
    return None


def fits_one_truck(pieces: list, stop_rank: dict) -> bool:
    """True if `pieces` fit one trailer by floor space (with stacking) and payload."""
    if sum(p.weight_lbs for p in pieces) > max_payload():
        return False
    for p in pieces:
        _normalize_placement(p)
        p.seq = stop_rank.get(p.stop, len(stop_rank) + p.stop)
    return _pack_columns(pieces)[1]


def split_into_trucks(pieces: list, route: list, decide) -> tuple:
    """
    Assign pieces to trucks following `route` (stop IDs in global delivery order).
    Returns (trucks, ship_other, impossible): trucks is a list of piece-ID
    lists, ship_other the IDs to ship another way, impossible [(piece, reason)].

    decide(info) is called when a stop doesn't fully fit and returns one of
      "new_truck"   — move all of the stop's remaining pieces to a new truck
      "ship_other"  — load what fits, ship the rest another way
      "continue"    — load what fits, carry the rest on a new truck
    info: {"stop", "truck" (1-based), "truck_empty", "pieces" (remaining),
           "fit" (how many fit), "fit_lbs", "rest_lbs"}.
    """
    rank = {sid: i + 1 for i, sid in enumerate(route)}
    impossible = [(p, why) for p in pieces if (why := physically_loadable(p))]
    bad = {p.piece_id for p, _ in impossible}
    by_stop = {}
    for p in pieces:
        if p.piece_id not in bad:
            by_stop.setdefault(p.stop, []).append(p)
    stops = [sid for sid in route if sid in by_stop] + \
            sorted(sid for sid in by_stop if sid not in rank)

    trucks, other = [[]], []
    for sid in stops:
        remaining = list(by_stop[sid])
        while remaining:
            cur = trucks[-1]
            if fits_one_truck(cur + remaining, rank):
                cur.extend(remaining)
                break
            # Largest prefix of this stop (in list order) that still fits
            lo, hi = 0, len(remaining)
            while lo < hi:
                mid = (lo + hi + 1) // 2
                if fits_one_truck(cur + remaining[:mid], rank):
                    lo = mid
                else:
                    hi = mid - 1
            fit = lo
            choice = decide({
                "stop": sid, "truck": len(trucks), "truck_empty": not cur,
                "pieces": remaining, "fit": fit,
                "fit_lbs": sum(p.weight_lbs for p in remaining[:fit]),
                "rest_lbs": sum(p.weight_lbs for p in remaining[fit:]),
            })
            if choice == "new_truck" and cur:
                trucks.append([])
                continue
            cur.extend(remaining[:fit])
            rest = remaining[fit:]
            if choice == "continue" or (choice == "new_truck" and not cur):
                trucks.append([])
                remaining = rest
                if fit == 0 and not cur:   # can't make progress — ship the rest
                    other.extend(rest)
                    break
            else:
                other.extend(rest)
                break
    trucks = [t for t in trucks if t]
    return ([[p.piece_id for p in t] for t in trucks] or [[]],
            [p.piece_id for p in other], impossible)


# ── Save / Load plan files ────────────────────────────────────────────────────
def _safe_name(s: str) -> str:
    """Strip characters unsafe for filenames, collapse whitespace to underscores."""
    return re.sub(r'[^\w-]', '_', s).strip('_')

def make_base_filename(info: ShipmentInfo) -> str:
    """
    Build the standard base name: Customer_BOL_DATE
    Parts are omitted when blank. Spaces/special chars become underscores.
    """
    parts = [_safe_name(p) for p in [info.customer, info.bol_num, info.date] if p.strip()]
    return '_'.join(parts) if parts else 'load_plan'

def save_plan_json(pieces: list, info: ShipmentInfo, limits: dict,
                   excluded_ids: list, override: bool, path: str,
                   securement: list = None, shipment: dict = None):
    data = {
        "version": 1,
        "saved_at": datetime.datetime.now().isoformat(timespec='seconds'),
        "shipment": {
            "customer":      info.customer,
            "shipment_name": info.shipment_name,
            "trailer_num":   info.trailer_num,
            "bol_num":       info.bol_num,
            "date":          info.date,
            "stop_names":    list(info.stop_names),
        },
        "axle_limits": limits,
        "override":     override,
        "excluded_ids": excluded_ids,
        "freight": [
            {
                "piece_id":    p.piece_id,
                "crate_label": p.crate_label,
                "part_number": p.part_number,
                "length_in":   p.length_in,
                "width_in":    p.width_in,
                "height_in":   p.height_in,
                "weight_lbs":  p.weight_lbs,
                "pcs":         p.pcs,
                "notes":       p.notes,
                "placement":   p.placement,
                "stackable":   p.stackable,
                "stop":        p.stop,
            }
            for p in pieces
        ],
        "securement": securement or [],
        "shipment_plan": shipment or {},
    }
    with open(path, 'w') as f:
        json.dump(data, f, indent=2)

def load_plan_json(path: str):
    """Returns (pieces, info, limits, excluded_ids, override, securement, shipment_plan)."""
    with open(path) as f:
        data = json.load(f)
    s = data['shipment']
    info = ShipmentInfo(
        customer=s.get('customer', ''),
        shipment_name=s.get('shipment_name', ''),
        trailer_num=s.get('trailer_num', ''),
        bol_num=s.get('bol_num', ''),
        date=s.get('date', ''),
        stop_names=(list(s.get('stop_names', [])) + [""] * MAX_STOPS)[:MAX_STOPS],
    )
    pieces = [
        FreightPiece(
            piece_id=p['piece_id'],
            crate_label=p['crate_label'],
            part_number=p['part_number'],
            length_in=p['length_in'],
            width_in=p['width_in'],
            height_in=p['height_in'],
            weight_lbs=p['weight_lbs'],
            pcs=p.get('pcs', 0),
            notes=p.get('notes', ''),
            placement=p.get('placement', 'auto'),
            stackable=p.get('stackable', True),
            stop=max(1, min(MAX_STOPS, int(p.get('stop', 1) or 1))),
        )
        for p in data['freight']
    ]
    limits = data.get('axle_limits', {'steer': 12000, 'drive': 34000, 'trailer': 34000})
    excluded_ids = data.get('excluded_ids', [])
    override = data.get('override', False)
    securement = data.get('securement', [])
    shipment = data.get('shipment_plan', {}) or {}
    return pieces, info, limits, excluded_ids, override, securement, shipment


# ── PDF helpers ───────────────────────────────────────────────────────────────
def _pdf_writer(path: str):
    w = QPdfWriter(path)
    w.setPageSize(QPageSize(QPageSize.PageSizeId.Letter))
    w.setPageOrientation(QPageLayout.Orientation.Portrait)
    w.setResolution(150)
    return w


def securement_summary(items: list) -> str:
    """'1 load bar, 2 straps, 4 airbags' (empty string when there are none)."""
    counts = {k: 0 for k in SECUREMENT_KINDS}
    for item in items or []:
        counts[item["kind"]] = counts.get(item["kind"], 0) + 1
    return ", ".join(f"{n} {SECUREMENT_KINDS[k].lower()}{'s' if n != 1 else ''}"
                     for k, n in counts.items() if n)


# ── Shared to-scale renderers (screen tabs + PDF) ─────────────────────────────
def _draw_securement_footprint(painter: QPainter, r: QRectF, kind: str, color: QColor,
                               u: float, translucent: bool = False):
    """Draw one securement part already mapped to a 2D rect (top or side view)."""
    min_t = 1.6 * u   # keep thin parts (straps) visible
    if r.width() < min_t:
        r = QRectF(r.center().x() - min_t / 2, r.y(), min_t, r.height())
    if r.height() < min_t:
        r = QRectF(r.x(), r.center().y() - min_t / 2, r.width(), min_t)
    if kind == "airbag":
        if translucent:
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.setPen(QPen(QColor("#5d6d7e"), 0.9 * u, Qt.PenStyle.DashLine))
        else:
            painter.setBrush(QBrush(color))
            painter.setPen(QPen(QColor("#7f8c8d"), 0.8 * u))
        rad = min(r.width(), r.height()) * 0.4
        painter.drawRoundedRect(r, rad, rad)
    else:
        painter.setBrush(QBrush(color))
        painter.setPen(QPen(color.darker(160), 0.5 * u))
        painter.drawRect(r)


def _draw_top_plan(painter: QPainter, rect: QRectF, plan: Optional[LoadPlan],
                   securement: list = None, u: float = 1.0, geom: dict = None) -> dict:
    """
    To-scale top view, nose at top, left wall (Position A) on the left — the
    same geometry as the 3D view. Row notes fill the space to the right.
    Shared by the Top View tab and the floor-plan PDF; `u` scales line widths
    and spacing (1.0 on screen). Returns {(row, col): QRectF} for hit-testing.
    """
    L, W = TRAILER["interior_length_in"], TRAILER["interior_width_in"]
    multi = len(plan_stops(plan)) > 1
    LBL, END, LEGEND = 30 * u, 18 * u, (32 if multi else 16) * u
    s = min((rect.height() - 2 * END - LEGEND) / L, rect.width() * 0.40 / W)
    tw, th = W * s, L * s
    tx, ty = rect.x() + LBL, rect.y() + END
    note_x = tx + tw + 12 * u
    note_w = max(10.0, rect.right() - note_x)
    center = Qt.AlignmentFlag.AlignCenter
    left_mid = Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter

    def to_rect(x, y, l, w):
        return QRectF(tx + y * s, ty + x * s, w * s, l * s)

    if geom is not None:   # lets the Top View map mouse pixels to trailer inches
        geom.update(tx=tx, ty=ty, s=s)

    # Trailer outline and end labels
    painter.setPen(QPen(QColor("#444"), 2 * u))
    painter.setBrush(QBrush(QColor("#f5f5f0")))
    painter.drawRect(QRectF(tx, ty, tw, th))
    painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
    painter.setPen(QColor("#333"))
    painter.drawText(QRectF(tx - 40 * u, rect.y(), tw + 80 * u, END), center, "▲ NOSE")
    painter.drawText(QRectF(tx - 40 * u, ty + th, tw + 80 * u, END), center, "DOORS")
    painter.setFont(QFont("Arial", 6))
    painter.setPen(QColor("#777"))
    painter.drawText(QRectF(tx, ty + th + 1 * u, 20 * u, END), Qt.AlignmentFlag.AlignLeft, "L")
    painter.drawText(QRectF(tx + tw - 20 * u, ty + th + 1 * u, 20 * u, END),
                     Qt.AlignmentFlag.AlignRight, "R")

    # Axle reference lines (labels right-aligned in the notes column)
    for depth, label, color, kind in axle_marks(plan):
        if kind == "kingpin":
            continue   # 1 ft from the drive line; the side view shows it
        ay = ty + depth * s
        painter.setPen(QPen(QColor(color), 1.2 * u, Qt.PenStyle.DashLine))
        painter.drawLine(QPointF(tx, ay), QPointF(tx + tw, ay))
        painter.setFont(QFont("Arial", 6))
        painter.setPen(QColor(color))
        painter.drawText(QRectF(note_x, ay - 11 * u, note_w, 11 * u),
                         Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignBottom, label)

    if not plan or not plan.slots:
        return {}

    items = securement or []
    boxes = _crate_boxes(plan)
    box_of = {b[0].piece_id: b for b in boxes}
    tops = {(sl.row, sl.col): sl.piece for sl in plan.slots
            if sl.tier == 1 and sl.piece and not (sl.is_center and sl.col == 1)}
    hits = {}
    deferred = []   # labels/badges drawn last so straps don't cover them

    # ── Crates (and open positions) ──────────────────────────────────────
    for sl in plan.slots:
        if sl.tier != 0 or (sl.is_center and sl.col == 1):
            continue
        if sl.piece is None:
            other = next((o for o in plan.slots if o.row == sl.row and o.tier == 0
                          and o.col != sl.col and o.piece), None)
            depth = other.piece.length_in if other else 48
            ow = min(other.piece.width_in, W) if other else W / 2
            y0 = ow if sl.col == 1 else 0
            r = to_rect(sl.row_start_in, y0, depth, W - ow).adjusted(u, u, -u, -u)
            painter.setBrush(QBrush(VOID_BG))
            painter.setPen(QPen(GRAY_LINE, u, Qt.PenStyle.DashLine))
            painter.drawRect(r)
            painter.setFont(QFont("Arial", 6))
            painter.setPen(QColor("#999"))
            painter.drawText(r, center, "open")
            hits[(sl.row, sl.col)] = r
            continue

        p, x, y, z, l, w, h, _t = box_of[sl.piece.piece_id]
        r = to_rect(x, y, l, w).adjusted(0.5 * u, 0.5 * u, -0.5 * u, -0.5 * u)
        color = QColor(CRATE_COLORS[(p.piece_id - 1) % len(CRATE_COLORS)])
        top = tops.get((sl.row, sl.col))
        painter.setPen(QPen(QColor("#222"), 0.6 * u))
        if top:
            # 2-high: bottom crate peeks out as an offset "card" behind the top crate
            painter.setBrush(QBrush(color.darker(135)))
            painter.drawRect(r.translated(3 * u, 3 * u))
            tb = box_of[top.piece_id]
            tr = to_rect(tb[1], tb[2], tb[4], tb[5]).adjusted(0.5 * u, 0.5 * u, -0.5 * u, -0.5 * u)
            painter.setBrush(QBrush(QColor(CRATE_COLORS[(top.piece_id - 1) % len(CRATE_COLORS)])))
            painter.drawRect(tr)
            deferred.append((tr.adjusted(2 * u, 10 * u, -2 * u, -1 * u),
                             f"{top.crate_label} ▲\non {p.crate_label}",
                             QRectF(tr.x() + 2 * u, tr.y() + 2 * u, 34 * u, 11 * u),
                             p.stop if multi else None))
        else:
            painter.setBrush(QBrush(color))
            painter.drawRect(r)
            deferred.append((r.adjusted(2 * u, 1 * u, -2 * u, -1 * u),
                             f"{p.crate_label}\n{p.weight_lbs:,.0f} lb", None,
                             p.stop if multi else None))
        hits[(sl.row, sl.col)] = r

    # ── Securement footprints ────────────────────────────────────────────
    order = {"airbag": 0, "load_bar": 1, "strap": 2}
    for item in sorted(items, key=lambda it: order.get(it["kind"], 9)):
        parts, err = securement_parts(item, boxes, items)
        for x, y, z, l, w, h, color, _shape in parts:
            _draw_securement_footprint(painter, to_rect(x, y, l, w), item["kind"], color, u)

    for text_rect, text, badge, stop in deferred:
        if stop is not None:
            sb = QRectF(text_rect.right() - 18 * u, text_rect.bottom() - 12 * u, 18 * u, 11 * u)
            painter.setBrush(QBrush(QColor(STOP_COLORS[(stop - 1) % len(STOP_COLORS)])))
            painter.setPen(QPen(Qt.GlobalColor.white, 0.6 * u))
            painter.drawRoundedRect(sb, 5 * u, 5 * u)
            painter.setFont(QFont("Arial", 5, QFont.Weight.Bold))
            painter.setPen(Qt.GlobalColor.white)
            painter.drawText(sb, center, f"S{stop}")
        if badge is not None:
            painter.setBrush(QBrush(TAN))
            painter.setPen(QPen(QColor("#8B6914"), 0.6 * u))
            painter.drawRoundedRect(badge, 2 * u, 2 * u)
            painter.setFont(QFont("Arial", 5, QFont.Weight.Bold))
            painter.setPen(Qt.GlobalColor.white)
            painter.drawText(badge, center, "2-HIGH")
        # Shrink the label to fit narrow crates
        size = 7
        f = QFont("Arial", size, QFont.Weight.Bold)
        while size > 4 and max(QFontMetrics(f, painter.device()).horizontalAdvance(line)
                               for line in text.split("\n")) > text_rect.width():
            size -= 0.5
            f.setPointSizeF(size)
        painter.setFont(f)
        painter.setPen(Qt.GlobalColor.white)
        painter.drawText(text_rect, center, text)

    # ── Row labels + notes ───────────────────────────────────────────────
    fm_font = QFont("Arial", 7)
    for r_pos, row in enumerate(sorted(set(sl.row for sl in plan.slots))):
        floor = [sl for sl in plan.slots if sl.row == row and sl.tier == 0 and sl.piece
                 and not (sl.is_center and sl.col == 1)]
        if not floor:
            continue
        xs = [(box_of[sl.piece.piece_id][1], box_of[sl.piece.piece_id][4]) for sl in floor]
        start = min(x for x, _ in xs)
        depth = max(x + l for x, l in xs) - start
        cy = ty + (start + depth / 2) * s
        painter.setFont(QFont("Arial", 7, QFont.Weight.Bold))
        painter.setPen(QColor("#444"))
        painter.drawText(QRectF(rect.x(), cy - 8 * u, LBL - 4 * u, 16 * u),
                         Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, f"R{row + 1}")

        bits, weight, pids = [], 0.0, set()
        for sl in sorted(floor, key=lambda q: q.col):
            top = tops.get((sl.row, sl.col))
            weight += sl.piece.weight_lbs + (top.weight_lbs if top else 0)
            pids.add(sl.piece.piece_id)
            if top:
                pids.add(top.piece_id)
                bits.append(f"{sl.piece.crate_label} + {top.crate_label} (2-high)")
            else:
                bits.append(sl.piece.crate_label)
            if sl.is_center and sl.piece.placement == "alt_center":
                y_off = box_of[sl.piece.piece_id][2]
                bits[-1] += (" ← LEFT" if y_off < 0.5 else
                             " RIGHT →" if y_off + box_of[sl.piece.piece_id][5] > W - 0.5
                             else " CENTER")
        sec = securement_summary([it for it in items if it["piece_id"] in pids])
        note = f"{' | '.join(bits)}  ·  {weight:,.0f} lb" + (f"  ·  {sec}" if sec else "")
        if multi:
            row_stops = [pc.stop for pc in sorted({sl.piece.seq: sl.piece for sl in floor}.values(),
                                                    key=lambda pc: -pc.seq)]
            note = f"[{'/'.join(f'S{st}' for st in row_stops)}]  " + note
        painter.setFont(fm_font)
        painter.setPen(QColor("#333"))
        note = QFontMetrics(fm_font, painter.device()).elidedText(
            note, Qt.TextElideMode.ElideRight, int(note_w))
        painter.drawText(QRectF(note_x, cy - 8 * u, note_w, 16 * u), left_mid, note)

    # ── Stop zone dividers ───────────────────────────────────────────────
    if multi:
        zone_rows = {}
        for sl in plan.slots:
            if sl.piece and sl.tier == 0:
                zone_rows.setdefault(sl.row, [sl.row_start_in, set()])[1].add(sl.piece.seq)
        ordered = [zone_rows[r] for r in sorted(zone_rows)]
        # Divide between rows whose stops don't overlap; a row shared by two
        # stops is itself the boundary, so it gets no divider
        for (start_a, stops_a), (start_b, stops_b) in zip(ordered, ordered[1:]):
            stop_a, stop_b = min(stops_a), max(stops_b)
            seq_to_id = {sl.piece.seq: sl.piece.stop for sl in plan.slots if sl.piece}
            if stop_a > stop_b:
                stop_a, stop_b = seq_to_id[stop_a], seq_to_id[stop_b]
                zy = ty + start_b * s
                painter.setPen(QPen(QColor("#333"), 2 * u, Qt.PenStyle.DashDotLine))
                painter.drawLine(QPointF(tx - 6 * u, zy), QPointF(tx + tw + 6 * u, zy))
                painter.setFont(QFont("Arial", 6, QFont.Weight.Bold))
                painter.setPen(QColor("#333"))
                painter.drawText(QRectF(note_x, zy - 6 * u, note_w, 12 * u),
                                 Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                                 f"▲ S{stop_a}   ─ ─   S{stop_b} ▼")

    # ── Legend ───────────────────────────────────────────────────────────
    lx, ly = note_x, rect.bottom() - LEGEND + 3 * u
    if multi:
        sx = note_x
        painter.setFont(QFont("Arial", 6))
        for st in plan_stops(plan):
            sb = QRectF(sx, ly + 2 * u, 18 * u, 9 * u)
            painter.setBrush(QBrush(QColor(STOP_COLORS[(st - 1) % len(STOP_COLORS)])))
            painter.setPen(Qt.PenStyle.NoPen)
            painter.drawRoundedRect(sb, 4 * u, 4 * u)
            painter.setFont(QFont("Arial", 5, QFont.Weight.Bold))
            painter.setPen(Qt.GlobalColor.white)
            painter.drawText(sb, center, f"S{st}")
            painter.setFont(QFont("Arial", 6))
            painter.setPen(QColor("#555"))
            text = stop_label(plan, st)[len(f"Stop {st}"):].lstrip(" —") or f"Stop {st}"
            if st == plan_stops(plan)[0]:
                text += " (1st, doors)"
            painter.drawText(QRectF(sx + 21 * u, ly, 200 * u, 12 * u), left_mid, text)
            sx += 21 * u + QFontMetrics(painter.font(), painter.device()).horizontalAdvance(text) + 12 * u
        ly += 16 * u
    painter.setFont(QFont("Arial", 6))
    for kind, label in [("stack", "2-high stack"), ("load_bar", "Load bar"),
                        ("strap", "Strap"), ("airbag", "Airbag")]:
        sw = QRectF(lx, ly + 2 * u, 14 * u, 8 * u)
        if kind == "stack":
            painter.setBrush(QBrush(TAN))
            painter.setPen(QPen(QColor("#8B6914"), 0.6 * u))
            painter.drawRoundedRect(sw, 2 * u, 2 * u)
        else:
            _draw_securement_footprint(painter, sw, kind, SECUREMENT_COLORS[kind], u)
        painter.setPen(QColor("#555"))
        painter.drawText(QRectF(lx + 17 * u, ly, 70 * u, 12 * u), left_mid, label)
        lx += 17 * u + QFontMetrics(painter.font(), painter.device()).horizontalAdvance(label) + 12 * u
    return hits


def _draw_side_plan(painter: QPainter, rect: QRectF, plan: Optional[LoadPlan],
                    securement: list = None, u: float = 1.0, geom: dict = None):
    """
    To-scale side elevation seen from the left wall (nose at left), same
    geometry as the 3D view's Side preset, with the kingpin, drive and trailer
    axles under the floor. Shared by the Side View tab and the PDF; `geom`
    receives the scale so the view can map mouse drags to inches.
    """
    L, H = TRAILER["interior_length_in"], TRAILER["interior_height_in"]
    ML, MT, MB = 38 * u, 22 * u, 52 * u
    s = min((rect.width() - ML - 10 * u) / L, (rect.height() - MT - MB) / H)
    ox = rect.x() + ML
    spare = rect.height() - MT - MB - H * s
    floor_y = rect.y() + MT + spare / 2 + H * s
    top_y = floor_y - H * s
    center = Qt.AlignmentFlag.AlignCenter

    painter.setPen(QPen(QColor("#444"), 2 * u))
    painter.setBrush(QBrush(QColor("#f5f5f0")))
    painter.drawRect(QRectF(ox, top_y, L * s, H * s))

    painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
    painter.setPen(QColor("#333"))
    painter.drawText(QRectF(ox, top_y - 16 * u, 80 * u, 14 * u),
                     Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, "◀ NOSE")
    painter.drawText(QRectF(ox + L * s - 80 * u, top_y - 16 * u, 80 * u, 14 * u),
                     Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, "DOORS ▶")
    painter.setFont(QFont("Arial", 6))
    painter.setPen(QColor("#666"))
    painter.drawText(QRectF(ox, top_y - 16 * u, L * s, 14 * u), center,
                     f"Interior {L}\" L × {H}\" H  (viewed from left side)")

    # Height ruler (every 12")
    for hm in range(0, H + 1, 12):
        yy = floor_y - hm * s
        painter.setPen(QPen(QColor("#bbb"), u))
        painter.drawLine(QPointF(ox - 5 * u, yy), QPointF(ox, yy))
        painter.setPen(QColor("#666"))
        painter.drawText(QRectF(rect.x(), yy - 6 * u, ML - 7 * u, 12 * u),
                         Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, f"{hm}\"")

    # Length ruler (every 5 ft)
    for xm in range(0, L + 1, 60):
        xx = ox + xm * s
        painter.setPen(QPen(QColor("#bbb"), u))
        painter.drawLine(QPointF(xx, floor_y + 31 * u), QPointF(xx, floor_y + 36 * u))
        painter.setPen(QColor("#666"))
        painter.drawText(QRectF(xx - 20 * u, floor_y + 36 * u, 40 * u, 12 * u), center,
                         f"{xm // 12}'")

    # Kingpin, drive and trailer axles: dashed line, wheels under the floor, weight
    if geom is not None:
        geom.update(ox=ox, s=s, top_y=top_y, floor_y=floor_y)
    aw = plan.axle_weights if plan else {}
    wheel_y, wheel_r = floor_y + 22 * u, 6.5 * u
    for depth, label, color, kind in axle_marks(plan):
        xx = ox + depth * s
        if kind == "kingpin":
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QBrush(QColor(color)))
            painter.drawPolygon(QPolygonF([QPointF(xx, floor_y + 2 * u), QPointF(xx - 4 * u, floor_y + 9 * u),
                                           QPointF(xx + 4 * u, floor_y + 9 * u)]))
            continue
        painter.setPen(QPen(QColor(color), 1.0 * u, Qt.PenStyle.DashLine))
        painter.drawLine(QPointF(xx, top_y), QPointF(xx, floor_y))
        spread = RIG["tandem_spread_in"] if kind == "tandem" else 52   # tractor tandem ≈ 52"
        painter.setPen(QPen(QColor("#222"), 1.0 * u))
        painter.setBrush(QBrush(QColor(color)))
        for dx in (-spread / 2, spread / 2):
            painter.drawEllipse(QPointF(xx + dx * s, wheel_y), wheel_r, wheel_r)
        weight = aw.get("drive" if kind == "drive" else "trailer")
        over = weight is not None and weight > TRAILER["drive_limit" if kind == "drive" else "trailer_limit"]
        painter.setFont(QFont("Arial", 6, QFont.Weight.Bold if over else QFont.Weight.Normal))
        painter.setPen(QColor("#c0392b" if over else color))
        painter.drawText(QRectF(xx + 2 * u, top_y + 1 * u, 90 * u, 20 * u),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignTop,
                         label + (f"\n{weight:,} lb" if weight is not None else ""))
    if aw:
        over = aw["steer"] > TRAILER["steer_limit"]
        painter.setFont(QFont("Arial", 6, QFont.Weight.Bold if over else QFont.Weight.Normal))
        painter.setPen(QColor("#c0392b" if over else "#2a7d6e"))
        painter.drawText(QRectF(rect.x(), wheel_y - 12 * u, ML - 2 * u, 24 * u),
                         Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                         f"◀ Steer\n{aw['steer']:,}")

    if not plan or not plan.slots:
        return

    items = securement or []
    boxes = _crate_boxes(plan)
    multi = len(plan_stops(plan)) > 1
    dividers = []   # (x, label) between delivery-stop zones, drawn last
    if multi:
        zone_rows = {}
        for sl in plan.slots:
            if sl.piece and sl.tier == 0:
                zone_rows.setdefault(sl.row, [sl.row_start_in, set()])[1].add(sl.piece.seq)
        ordered = [zone_rows[r] for r in sorted(zone_rows)]
        # Divide between rows whose stops don't overlap; a row shared by two
        # stops is itself the boundary, so it gets no divider
        for (start_a, stops_a), (start_b, stops_b) in zip(ordered, ordered[1:]):
            stop_a, stop_b = min(stops_a), max(stops_b)
            seq_to_id = {sl.piece.seq: sl.piece.stop for sl in plan.slots if sl.piece}
            if stop_a > stop_b:
                stop_a, stop_b = seq_to_id[stop_a], seq_to_id[stop_b]
                dividers.append((ox + start_b * s, f"S{stop_a} | S{stop_b}"))
    # Far side (right wall) first so crates nearer the viewer paint over it
    labels = []   # drawn after securement so straps don't cover them
    for p, x, y, z, l, w, h, tier in sorted(boxes, key=lambda b: -b[2]):
        r = QRectF(ox + x * s, floor_y - (z + h) * s, l * s, h * s).adjusted(
            0.5 * u, 0.5 * u, -0.5 * u, 0)
        painter.setBrush(QBrush(QColor(CRATE_COLORS[(p.piece_id - 1) % len(CRATE_COLORS)])))
        painter.setPen(QPen(QColor("#222"), 0.6 * u))
        painter.drawRect(r)
        # Skip the label when a nearer crate hides most of this one
        hidden = any(o[2] < y and
                     (min(x + l, o[1] + o[4]) - max(x, o[1])) *
                     (min(z + h, o[3] + o[6]) - max(z, o[3])) > 0.5 * l * h
                     for o in boxes
                     if min(x + l, o[1] + o[4]) > max(x, o[1]) and min(z + h, o[3] + o[6]) > max(z, o[3]))
        if not hidden:
            labels.append((r.adjusted(u, u, -u, -u),
                           f"{p.crate_label}{' ▲' if tier == 1 else ''}"
                           f"{f' S{p.stop}' if multi else ''}\n{h:.0f}\""))
        if z + h > H:
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.setPen(QPen(QColor("#c0392b"), 2 * u))
            painter.drawRect(QRectF(ox + x * s, floor_y - (z + h) * s, l * s, (z + h) * s))

    # Securement (airbags as dashed outlines — they may sit in front of a crate)
    for item in items:
        parts, err = securement_parts(item, boxes, items)
        for x, y, z, l, w, h, color, _shape in parts:
            r = QRectF(ox + x * s, floor_y - (z + h) * s, l * s, h * s)
            _draw_securement_footprint(painter, r, item["kind"], color, u, translucent=True)

    painter.setFont(QFont("Arial", 6, QFont.Weight.Bold))
    painter.setPen(Qt.GlobalColor.white)
    for r, text in labels:
        painter.drawText(r, center, text)

    for xx, text in dividers:
        painter.setPen(QPen(QColor("#333"), 1.6 * u, Qt.PenStyle.DashDotLine))
        painter.drawLine(QPointF(xx, top_y - 4 * u), QPointF(xx, floor_y + 4 * u))
        tag = QRectF(xx - 22 * u, top_y + 3 * u, 44 * u, 11 * u)
        painter.setBrush(QBrush(QColor("#ffffff")))
        painter.setPen(QPen(QColor("#333"), 0.8 * u))
        painter.drawRoundedRect(tag, 3 * u, 3 * u)
        painter.setFont(QFont("Arial", 6, QFont.Weight.Bold))
        painter.setPen(QColor("#333"))
        painter.drawText(tag, center, text)

    # Row labels under the floor
    painter.setFont(QFont("Arial", 6))
    painter.setPen(QColor("#555"))
    for row in sorted(set(sl.row for sl in plan.slots)):
        floor = [sl for sl in plan.slots if sl.row == row and sl.tier == 0 and sl.piece]
        if not floor:
            continue
        start = floor[0].row_start_in
        depth = max(sl.piece.length_in for sl in floor)
        painter.drawText(QRectF(ox + start * s, floor_y + 2 * u, depth * s, 12 * u),
                         center, f"R{row + 1}")


# ── PDF: Load Plan Floor Plan ─────────────────────────────────────────────────
def export_load_plan_pdf(plan: LoadPlan, info: ShipmentInfo, path: str,
                         securement: list = None, include_3d: bool = False):
    """Floor plan PDF for one truck (same pages as a one-truck shipment)."""
    export_shipment_pdf([(plan, info)], path, securement, include_3d)


PDF_3D_ANGLES = [                                   # (caption, yaw°, pitch°)
    ("3/4 view from rear-left (doors)", -40.0, 28.0),
    ("3/4 view from front-right (nose)", 140.0, 28.0),
    ("Overhead from left side", -70.0, 55.0),
    ("Rear — looking in through the doors", 0.0, 10.0),
]


def _draw_3d_page(painter: QPainter, PW: int, PH: int, plan: LoadPlan,
                  info: ShipmentInfo, securement: list = None):
    """One landscape page: the loaded trailer from four 3D angles in a 2×2 grid."""
    M = 72
    painter.setPen(QColor("#111"))
    painter.setFont(QFont("Arial", 20, QFont.Weight.Bold))
    painter.drawText(QRectF(M, M, PW - 2 * M, 44),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, "3D Views" + (f"  —  {info.truck_label}" if info.truck_label else ""))
    hdr = "  |  ".join(part for part in [info.customer or info.shipment_name, info.date,
                                         securement_summary(securement)] if part)
    if hdr:
        painter.setFont(QFont("Arial", 10))
        painter.setPen(QColor("#555"))
        painter.drawText(QRectF(M, M + 44, PW - 2 * M, 24),
                         Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, hdr)

    top, gap, cap = M + 78, 24, 26
    cell_w = (PW - 2 * M - gap) / 2
    cell_h = (PH - top - M - gap) / 2
    for i, (caption, yaw, pitch) in enumerate(PDF_3D_ANGLES):
        cx = M + (i % 2) * (cell_w + gap)
        cy = top + (i // 2) * (cell_h + gap)
        frame = QRectF(cx, cy, cell_w, cell_h)
        painter.fillRect(frame, QColor("#fafafa"))
        painter.setPen(QPen(QColor("#ccc"), 1))
        painter.setBrush(Qt.BrushStyle.NoBrush)
        painter.drawRect(frame)
        painter.setFont(QFont("Arial", 10, QFont.Weight.Bold))
        painter.setPen(QColor("#333"))
        painter.drawText(QRectF(cx + 10, cy + 4, cell_w - 20, cap),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, caption)
        _draw_3d_snapshot(painter, QRectF(cx + 4, cy + cap, cell_w - 8, cell_h - cap - 4),
                          plan, securement, yaw, pitch)


def _draw_floor_plan(painter: QPainter, PW: int, PH: int,
                     plan: LoadPlan, info: ShipmentInfo, securement: list = None):
    M = 72       # margin ~0.48"
    u = 150 / 96   # PDF px per screen px (writer runs at 150 dpi)

    # ── Title ──────────────────────────────────────────────────────────
    painter.setPen(QColor("#111"))
    painter.setFont(QFont("Arial", 20, QFont.Weight.Bold))
    painter.drawText(QRectF(M, M, PW - 2 * M, 48),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                     "Trailer Loading Floor Plan" + (f"  —  {info.truck_label}" if info.truck_label else ""))

    painter.setFont(QFont("Arial", 10))
    # Deduplicate center pieces (two slots share one piece)
    seen = set()
    total_w, n_crates = 0, 0
    for s in plan.slots:
        if s.piece and s.piece.piece_id not in seen:
            seen.add(s.piece.piece_id)
            total_w += s.piece.weight_lbs
            n_crates += 1
    excl = (f"Crates {', '.join(str(x) for x in plan.excluded_ids)} excluded  |  "
            if plan.excluded_ids else "")
    sub = "  |  ".join(part for part in [
        info.customer or info.shipment_name, TRAILER['name'],
        f"{excl}Total: {total_w:,.0f} lb / {n_crates} crates"] if part)
    painter.drawText(QRectF(M, M + 50, PW - 2 * M, 26),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, sub)

    detail_parts = []
    if info.date:        detail_parts.append(f"Date: {info.date}")
    if info.trailer_num: detail_parts.append(f"Trailer: {info.trailer_num}")
    if info.bol_num:     detail_parts.append(f"BOL: {info.bol_num}")
    if detail_parts:
        painter.drawText(QRectF(M, M + 76, PW - 2 * M, 22),
                         Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                         "  |  ".join(detail_parts))

    # ── To-scale plan ───────────────────────────────────────────────────
    not_loaded, NOT_LOADED_H = "", 0             # wrapped "does not fit" block
    if plan.unplaced:
        not_loaded = (f"NOT LOADED — does not fit ({len(plan.unplaced)} pcs, "
                      f"{sum(p.weight_lbs for p, _ in plan.unplaced):,.0f} lb): "
                      + ", ".join(p.crate_label for p, _ in plan.unplaced))
        painter.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        NOT_LOADED_H = painter.fontMetrics().boundingRect(
            0, 0, int(PW - 2 * M), 10_000, int(Qt.TextFlag.TextWordWrap.value), not_loaded).height() + 12
    TAIL_RESERVE = 110 + NOT_LOADED_H            # space for notes at bottom
    plan_top = M + 108
    _draw_top_plan(painter, QRectF(M, plan_top, PW - 2 * M, PH - M - TAIL_RESERVE - plan_top),
                   plan, securement, u)

    # ── Footer notes ───────────────────────────────────────────────────
    rows = sorted(set(s.row for s in plan.slots))
    n_rows = len(rows)
    tail_y = PH - M - TAIL_RESERVE + 8
    if not_loaded:
        painter.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        painter.setPen(QColor("#c0392b"))
        painter.drawText(QRectF(M, tail_y, PW - 2 * M, NOT_LOADED_H),
                         Qt.AlignmentFlag.AlignHCenter | Qt.TextFlag.TextWordWrap, not_loaded)
        tail_y += NOT_LOADED_H
    clear_in = TRAILER["interior_length_in"] - plan.total_length_used_in
    painter.setFont(QFont("Arial", 8))
    painter.setPen(QColor("#555"))
    painter.drawText(QRectF(M, tail_y, PW - 2 * M, 20), Qt.AlignmentFlag.AlignHCenter,
                     f"~{clear_in / 12.0:.1f} ft ({clear_in:.0f}\") clear space behind row {n_rows}")
    sec = securement_summary(securement)
    if sec:
        painter.setPen(QColor("#333"))
        painter.drawText(QRectF(M, tail_y + 20, PW - 2 * M, 20), Qt.AlignmentFlag.AlignHCenter,
                         f"Securement: {sec}")
    painter.setPen(QColor("#555"))
    note = (f"Load strictly R1 to R{n_rows} in order.  "
            "Weights are estimates — weigh on a certified scale before transport.")
    painter.drawText(QRectF(M, tail_y + 40, PW - 2 * M, 34),
                     Qt.AlignmentFlag.AlignHCenter | Qt.TextFlag.TextWordWrap, note)

    # Axle weight summary
    aw = plan.axle_weights
    if aw:
        aw_str = (f"Est. Axle Weights  —  "
                  f"Steer: {aw.get('steer', 0):,} lb  |  "
                  f"Drive: {aw.get('drive', 0):,} lb  |  "
                  f"Trailer: {aw.get('trailer', 0):,} lb  |  "
                  f"Gross: {aw.get('gross', 0):,} lb  (freight {aw.get('freight_gross', 0):,})")
        painter.setPen(TEAL_DARK)
        painter.drawText(QRectF(M, PH - M - 16, PW - 2 * M, 16),
                         Qt.AlignmentFlag.AlignHCenter, aw_str)
        if plan.rig:
            painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
            painter.drawText(QRectF(M, PH - M - 34, PW - 2 * M, 16),
                             Qt.AlignmentFlag.AlignHCenter, f"Set {rig_text(plan.rig)}")


def _draw_side_elevation(painter: QPainter, PW: int, PH: int,
                         plan: LoadPlan, info: ShipmentInfo, securement: list = None):
    """Side elevation page: to-scale profile of every crate, stacks and securement."""
    M = 72
    u = 150 / 96   # PDF px per screen px (writer runs at 150 dpi)

    painter.setPen(QColor("#111"))
    painter.setFont(QFont("Arial", 20, QFont.Weight.Bold))
    painter.drawText(QRectF(M, M, PW - 2 * M, 48),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                     "Side Elevation — Stack Height Profile" + (f"  —  {info.truck_label}" if info.truck_label else ""))

    painter.setFont(QFont("Arial", 10))
    painter.setPen(QColor("#555"))
    hdr = info.customer or info.shipment_name or ""
    if info.date:
        hdr = (hdr + "  |  " if hdr else "") + info.date
    if hdr:
        painter.drawText(QRectF(M, M + 50, PW - 2 * M, 26),
                         Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, hdr)

    # A 53' van is ~6:1, so the to-scale profile is a wide strip near the top
    side_h = (PW - 2 * M) * 0.26
    _draw_side_plan(painter, QRectF(M, M + 84, PW - 2 * M, side_h), plan, securement, u)
    if securement:
        _draw_securement_table(painter, QRectF(M, M + 100 + side_h, PW - 2 * M,
                                               PH - 2 * M - 100 - side_h), plan, securement)


def securement_rows(plan: LoadPlan, securement: list) -> list:
    """One loader-checklist row per securement item, sorted nose to tail."""
    row_of, label_of, start_of = {}, {}, {}
    for sl in plan.slots:
        if sl.piece:
            row_of[sl.piece.piece_id] = sl.row
            label_of[sl.piece.piece_id] = sl.piece.crate_label
            start_of[sl.piece.piece_id] = sl.row_start_in
    boxes = _crate_boxes(plan)
    face_text = {"nose": "nose face", "tail": "tail face", "left": "left side",
                 "right": "right side", "top": "over the top"}
    rows = []
    for item in securement:
        pid = item["piece_id"]
        if pid not in row_of:
            continue
        parts, err = securement_parts(item, boxes, securement)
        where = face_text.get(item["face"], item["face"])
        if item["kind"] == "strap" and item["face"] != "top":
            same = [it for it in securement if it["kind"] == "strap"
                    and it["piece_id"] == pid and it["face"] == item["face"]]
            if len(same) > 1:
                where += ", upper" if item.get("n", 0) == min(it.get("n", 0) for it in same) else ", lower"
        detail = ""
        if parts and not err:
            x, y, z, l, w, h = parts[0][:6]
            if item["kind"] == "airbag":
                gap = l if item["face"] in ("nose", "tail") else w
                detail = f"fills {gap + 0.6:.0f}\" void"
            elif item["kind"] == "load_bar":
                detail = f"{z + h / 2:.0f}\" above floor, {x + l / 2:.0f}\" from nose"
            elif item["face"] != "top":
                detail = f"{z + h / 2:.0f}\" above floor"
        elif err:
            detail = "does not fit current plan"
        rows.append((row_of[pid], start_of[pid], f"R{row_of[pid] + 1}", label_of[pid],
                     SECUREMENT_KINDS[item["kind"]], where, detail))
    rows.sort(key=lambda r: (r[1], r[0]))
    return [r[2:] for r in rows]


def _draw_securement_table(painter: QPainter, rect: QRectF, plan: LoadPlan, securement: list):
    """Loader checklist: one line per load bar / strap / airbag, nose to tail."""
    rows = securement_rows(plan, securement)
    painter.setFont(QFont("Arial", 12, QFont.Weight.Bold))
    painter.setPen(QColor("#111"))
    painter.drawText(QRectF(rect.x(), rect.y(), rect.width(), 26),
                     Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                     f"Securement Checklist  —  {securement_summary(securement)}")
    cols = [("✓", 0.04), ("Row", 0.07), ("Crate", 0.14), ("Item", 0.12),
            ("Location", 0.25), ("Detail", 0.38)]
    line_h = 26
    y = rect.y() + 32
    # Two side-by-side columns of rows if they won't fit in one
    max_rows = max(1, int((rect.bottom() - y) / line_h) - 1)
    n_cols = 1 if len(rows) <= max_rows else 2
    col_w = rect.width() / n_cols
    for c in range(n_cols):
        chunk = rows[c * max_rows:(c + 1) * max_rows]
        if not chunk:
            break
        x0 = rect.x() + c * col_w
        painter.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        painter.setPen(TEAL_DARK)
        cx = x0
        for name, frac in cols:
            painter.drawText(QRectF(cx, y, col_w * frac, line_h),
                             Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, name)
            cx += col_w * frac
        painter.setPen(QPen(TEAL_DARK, 1))
        painter.drawLine(QPointF(x0, y + line_h), QPointF(x0 + col_w - 20, y + line_h))
        painter.setFont(QFont("Arial", 9))
        for i, row in enumerate(chunk):
            ry = y + line_h * (i + 1)
            if i % 2:
                painter.fillRect(QRectF(x0, ry, col_w - 20, line_h), QColor("#f2f9f8"))
            painter.setPen(QPen(QColor("#555"), 1))
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.drawRect(QRectF(x0 + 4, ry + 7, 12, 12))
            cx = x0 + col_w * cols[0][1]
            painter.setPen(QColor("#222"))
            for val, (_name, frac) in zip(row, cols[1:]):
                painter.drawText(QRectF(cx, ry, col_w * frac - 4, line_h),
                                 Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, val)
                cx += col_w * frac


# ── PDF: Crate Labels ─────────────────────────────────────────────────────────
def export_crate_labels_pdf(plan: LoadPlan, info: ShipmentInfo, path: str):
    export_shipment_labels([(plan, info)], path)


def _draw_crate_label(painter: QPainter, PW: int, PH: int,
                      slot: PalletSlot, plan: LoadPlan, info: ShipmentInfo):
    PAD = 18    # outer border inset
    INNER_X = PAD + 22
    INNER_W = PW - 2 * (PAD + 22)

    # Outer border (tan/gold)
    painter.setPen(QPen(TAN, 1.5))
    painter.setBrush(Qt.BrushStyle.NoBrush)
    painter.drawRect(QRectF(PAD, PAD, PW - 2 * PAD, PH - 2 * PAD))

    # ── Top band: shipment name ─────────────────────────────────────────
    sep1_y = PH * 0.16

    font = QFont("Arial", 11)
    painter.setFont(font)
    painter.setPen(QColor("#555"))
    hdr = info.customer or info.shipment_name or "Freight Shipment"
    if info.date:
        hdr += f"   |   {info.date}"
    if info.truck_label:
        hdr += f"   |   {info.truck_label}"
    painter.drawText(QRectF(INNER_X, PAD + 10, INNER_W, sep1_y - PAD - 14),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, hdr)

    painter.setPen(QPen(GRAY_LINE, 1))
    painter.drawLine(QPointF(INNER_X, sep1_y), QPointF(INNER_X + INNER_W, sep1_y))

    # ── Teal main box ──────────────────────────────────────────────────
    box_top = sep1_y + 18
    box_h   = PH * 0.30
    box_x   = INNER_X + 10
    box_w   = INNER_W - 20
    box     = QRectF(box_x, box_top, box_w, box_h)

    painter.setBrush(QBrush(TEAL_LIGHT))
    painter.setPen(QPen(TEAL_DARK, 2.5))
    painter.drawRoundedRect(box, 16, 16)

    if len(plan_stops(plan)) > 1:   # big delivery-stop badge for multi-stop loads
        stop = slot.piece.stop
        badge = QRectF(box_x + box_w - 190, box_top - 14, 180, 46)
        painter.setBrush(QBrush(QColor(STOP_COLORS[(stop - 1) % len(STOP_COLORS)])))
        painter.setPen(QPen(Qt.GlobalColor.white, 2))
        painter.drawRoundedRect(badge, 10, 10)
        painter.setFont(QFont("Arial", 16, QFont.Weight.Bold))
        painter.setPen(Qt.GlobalColor.white)
        painter.drawText(badge, Qt.AlignmentFlag.AlignCenter, f"STOP {stop}")
        name = stop_label(plan, stop)[len(f"Stop {stop}"):].lstrip(" —")
        if name:
            painter.setFont(QFont("Arial", 10, QFont.Weight.Bold))
            painter.setPen(QColor(STOP_COLORS[(stop - 1) % len(STOP_COLORS)]))
            painter.drawText(QRectF(box_x + box_w - 400, box_top + 34, 390, 20),
                             Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, name)

    # Layout zones — shift to make room for notes when present
    has_notes = bool(slot.piece.notes)
    crate_h_frac  = 0.46 if has_notes else 0.58
    notes_top_frac = 0.47
    notes_h_frac   = 0.20
    pos_top_frac  = 0.68 if has_notes else 0.62
    pos_h_frac    = 0.29 if has_notes else 0.34

    # Crate number — very large
    font = QFont("Arial", 54, QFont.Weight.Bold)
    painter.setFont(font)
    painter.setPen(TEAL_DARK)
    painter.drawText(QRectF(box_x + 10, box_top + 16, box_w - 20, box_h * crate_h_frac),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                     slot.piece.crate_label.upper())

    # Notes — smaller, italic, below crate number
    if has_notes:
        font = QFont("Arial", 13)
        font.setItalic(True)
        painter.setFont(font)
        painter.setPen(QColor("#1a5c52"))
        painter.drawText(
            QRectF(box_x + 14, box_top + box_h * notes_top_frac, box_w - 28, box_h * notes_h_frac),
            Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter
            | Qt.TextFlag.TextWordWrap,
            slot.piece.notes)

    # Position label
    row_num = slot.row + 1
    if slot.is_center:
        if slot.piece.placement == "alt_center":
            rows_list = sorted(set(s.row for s in plan.slots))
            r_pos = rows_list.index(slot.row) if slot.row in rows_list else slot.row
            direction = "← Load Left" if r_pos % 2 == 0 else "Load Right →"
            pos_str = f"{direction}   ·   Row {row_num}"
        else:
            pos_str = f"Full Width   ·   Row {row_num}"
    else:
        side_letter = "A" if slot.col == 0 else "B"
        side_name   = "Left" if slot.col == 0 else "Right"
        pos_str = f"Position {side_letter}  ({side_name})   ·   Row {row_num}"
    if slot.tier == 1:
        t0_map = {(s.row, s.col): s for s in plan.slots if s.tier == 0}
        bot = t0_map.get((slot.row, slot.col))
        if bot and bot.piece:
            pos_str += f"   ·   Stacked on {bot.piece.crate_label}"
    font = QFont("Arial", 16)
    font.setItalic(False)
    painter.setFont(font)
    painter.setPen(TEAL_DARK)
    painter.drawText(QRectF(box_x + 10, box_top + box_h * pos_top_frac, box_w - 20, box_h * pos_h_frac),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter, pos_str)

    # ── Middle separator ───────────────────────────────────────────────
    sep2_y = box_top + box_h + 22
    painter.setPen(QPen(GRAY_LINE, 1))
    painter.drawLine(QPointF(INNER_X, sep2_y), QPointF(INNER_X + INNER_W, sep2_y))

    # ── Detail section ─────────────────────────────────────────────────
    detail_lines = []
    if slot.piece.part_number:
        detail_lines.append(("Part Number:", slot.piece.part_number))
    detail_lines.append(("Dimensions:",
                          f'{slot.piece.length_in:.0f}" × {slot.piece.width_in:.0f}" × {slot.piece.height_in:.0f}"  (L × W × H)'))
    detail_lines.append(("Weight:", f"{slot.piece.weight_lbs:,.0f} lbs"))
    if slot.piece.pcs:
        detail_lines.append(("Pieces:", f"{slot.piece.pcs:,}"))
    if slot.piece.notes:
        detail_lines.append(("Notes:", slot.piece.notes))

    LBL_W = 145
    ROW_H = 40
    det_y = sep2_y + 14
    font_lbl = QFont("Arial", 11)
    font_val = QFont("Arial", 13, QFont.Weight.Bold)

    for label, val in detail_lines:
        painter.setFont(font_lbl)
        painter.setPen(QColor("#666"))
        painter.drawText(QRectF(INNER_X + 20, det_y, LBL_W, ROW_H - 4),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, label)
        painter.setFont(font_val)
        painter.setPen(QColor("#111"))
        painter.drawText(QRectF(INNER_X + LBL_W + 10, det_y, INNER_W - LBL_W - 20, ROW_H - 4),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, val)
        det_y += ROW_H

    # ── Lower separator ────────────────────────────────────────────────
    sep3_y = det_y + 10
    painter.setPen(QPen(GRAY_LINE, 1))
    painter.drawLine(QPointF(INNER_X, sep3_y), QPointF(INNER_X + INNER_W, sep3_y))

    # ── Shipment footer ────────────────────────────────────────────────
    footer_lines = []
    if info.customer:        footer_lines.append(("Customer:", info.customer))
    if info.shipment_name:   footer_lines.append(("Shipment:", info.shipment_name))
    if info.bol_num:         footer_lines.append(("BOL #:", info.bol_num))
    if info.trailer_num:     footer_lines.append(("Trailer #:", info.trailer_num))

    fot_y = sep3_y + 14
    font_lbl2 = QFont("Arial", 11)
    font_val2 = QFont("Arial", 11)

    for label, val in footer_lines:
        painter.setFont(font_lbl2)
        painter.setPen(QColor("#666"))
        painter.drawText(QRectF(INNER_X + 20, fot_y, LBL_W, 30),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, label)
        painter.setFont(font_val2)
        painter.setPen(QColor("#222"))
        painter.drawText(QRectF(INNER_X + LBL_W + 10, fot_y, INNER_W - LBL_W - 20, 30),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, val)
        fot_y += 32


# ── Live Trailer View widget ──────────────────────────────────────────────────
SNAP_IN = 2.0      # a dragged crate snaps to a wall or neighbouring edge within this distance


def _slide_footprint(fp: tuple, target: tuple, obstacles: list) -> tuple:
    """
    Move footprint fp = (x, y, l, w) toward target (x, y) without leaving the
    trailer or passing through obstacles [(x, y, l, w)]: it slides along each
    axis until it touches something. Tries length-then-width and width-then-
    length and keeps whichever ends closer to the target.
    """
    L, W = TRAILER["interior_length_in"], TRAILER["interior_width_in"]
    x0, y0, l, w = fp
    tx, ty = target

    def sweep(x, y, t, axis):
        size, lim_hi = (l, L) if axis == 0 else (w, W)
        pos = x if axis == 0 else y
        hi, lo = lim_hi - size, 0.0
        for ox, oy, ol, ow in obstacles:
            # Only obstacles beside this crate's path along the other axis matter
            if axis == 0 and min(y + w, oy + ow) - max(y, oy) <= 0.01:
                continue
            if axis == 1 and min(x + l, ox + ol) - max(x, ox) <= 0.01:
                continue
            o_lo, o_size = (ox, ol) if axis == 0 else (oy, ow)
            if o_lo >= pos + size - 0.01:
                hi = min(hi, o_lo - size)
            elif o_lo + o_size <= pos + 0.01:
                lo = max(lo, o_lo + o_size)
        return max(lo, min(hi, t))

    xa = sweep(x0, y0, tx, 0)
    ya = sweep(xa, y0, ty, 1)
    yb = sweep(x0, y0, ty, 1)
    xb = sweep(x0, yb, tx, 0)
    da = (xa - tx) ** 2 + (ya - ty) ** 2
    db = (xb - tx) ** 2 + (yb - ty) ** 2
    return (xa, ya) if da <= db else (xb, yb)


def _snap_footprint(fp: tuple, pos: tuple, obstacles: list) -> tuple:
    """Pull a position onto nearby walls / neighbouring crate edges, if still clear."""
    L, W = TRAILER["interior_length_in"], TRAILER["interior_width_in"]
    _x0, _y0, l, w = fp
    x, y = pos
    xs = [0.0, L - l] + [v for ox, oy, ol, ow in obstacles for v in (ox + ol, ox - l, ox)]
    ys = [0.0, W - w] + [v for ox, oy, ol, ow in obstacles for v in (oy + ow, oy - w, oy)]

    def clear(cx, cy):
        return (0 <= cx <= L - l + 0.01 and 0 <= cy <= W - w + 0.01 and
                all(min(cx + l, ox + ol) - max(cx, ox) <= 0.01 or
                    min(cy + w, oy + ow) - max(cy, oy) <= 0.01 for ox, oy, ol, ow in obstacles))
    best_x = min(xs, key=lambda v: abs(v - x))
    if abs(best_x - x) <= SNAP_IN and clear(best_x, y):
        x = best_x
    best_y = min(ys, key=lambda v: abs(v - y))
    if abs(best_y - y) <= SNAP_IN and clear(x, best_y):
        y = best_y
    return x, y


class TrailerView(QWidget):
    """
    To-scale top view. Drag a crate to shift it (it slides until it touches a
    crate or wall, snapping to nearby edges); Ctrl+drag onto another position
    swaps two side-by-side crates.
    """
    plan_changed = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.plan: Optional[LoadPlan] = None
        self.securement: list = []
        self.setMinimumSize(320, 480)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self._drag_src: Optional[tuple] = None   # Ctrl+drag swap source (row, col)
        self._move = None                        # active shift: dict, see _start_move
        self._hover_slot: Optional[tuple] = None
        self._hits: dict = {}        # (row, col) -> QRectF, from the last paint
        self._geom: dict = {}        # tx, ty, s from the last paint (pixels ↔ inches)
        self.zoom_focus_x: Optional[float] = None
        self.setMouseTracking(True)

    def set_plan(self, plan: LoadPlan):
        self.plan = plan
        self._drag_src = None
        self._move = None
        self._hover_slot = None
        self.update()

    def set_securement(self, items: list):
        self.securement = items
        self.update()

    # ── Hit-test: which slot is under a mouse position ───────────────────────
    def _slot_at(self, pos) -> Optional[tuple]:
        if not self.plan or not self.plan.slots:
            return None
        for key, r in self._hits.items():
            if r.contains(QPointF(pos)):
                return key
        return None

    def _floor_slot(self, key) -> Optional[PalletSlot]:
        if not key or not self.plan:
            return None
        return next((s for s in self.plan.slots
                     if (s.row, s.col) == key and s.tier == 0 and s.piece), None)

    # ── Shift a crate (and whatever is stacked on it) ────────────────────────
    def _start_move(self, slot: PalletSlot, pos):
        boxes = _crate_boxes(self.plan)
        me = next(b for b in boxes if b[0].piece_id == slot.piece.piece_id)
        moving_ids = {slot.piece.piece_id} | {
            s.piece.piece_id for s in self.plan.slots
            if s.tier == 1 and s.piece and s.row == slot.row
            and (slot.is_center or s.col == slot.col)}
        self._move = {
            "slot": slot, "press": QPointF(pos), "fp": (me[1], me[2], me[4], me[5]),
            "pos": (me[1], me[2]),
            "obstacles": [(b[1], b[2], b[4], b[5]) for b in boxes
                          if b[7] == 0 and b[0].piece_id not in moving_ids],
        }

    def _update_move(self, pos):
        m, g = self._move, self._geom
        if not g:
            return
        dx_px, dy_px = pos.x() - m["press"].x(), pos.y() - m["press"].y()
        x0, y0, l, w = m["fp"]
        # Screen down = toward the doors (+x); screen right = toward the right wall (+y)
        target = (round(x0 + dy_px / g["s"]), round(y0 + dx_px / g["s"]))
        slid = _slide_footprint(m["fp"], target, m["obstacles"])
        m["pos"] = _snap_footprint(m["fp"], slid, m["obstacles"])
        self.update()

    def _finish_move(self):
        m = self._move
        self._move = None
        x0, y0, _l, _w = m["fp"]
        nx, ny = m["pos"]
        if abs(nx - x0) < 0.01 and abs(ny - y0) < 0.01:
            self.update()
            return
        slot = m["slot"]
        slot.x_in, slot.y_in = nx, ny
        _refresh_plan_metrics(self.plan)
        self.plan_changed.emit()
        self.update()

    # ── Swap two pieces between slots (Ctrl+drag) ────────────────────────────
    def _swap_pieces(self, src: tuple, dst: tuple):
        slot_map = {(s.row, s.col): s for s in self.plan.slots if s.tier == 0}
        s1, s2 = slot_map.get(src), slot_map.get(dst)
        if not s1 or s2 is None or s1.is_center or s2.is_center:
            return
        # Stacked tops travel with their bottom crate
        tops = {(s.row, s.col): s for s in self.plan.slots if s.tier == 1}
        t1, t2 = tops.get(src), tops.get(dst)
        # Crates only trade places with crates for the same stop (or move into an
        # open position in a row of their own stop)
        if s2.piece is not None:
            same_stop = s2.piece.stop == s1.piece.stop
        else:
            same_stop = all(s.piece.stop == s1.piece.stop for s in self.plan.slots
                            if s.row == s2.row and s.piece)
        if not same_stop:
            QToolTip.showText(QCursor.pos(), "Swap not allowed — crates stay in their "
                              "delivery stop's zone.", self)
            return
        saved = [(s1.x_in, s1.y_in), (s2.x_in, s2.y_in)]

        def exchange():
            s1.piece, s2.piece = s2.piece, s1.piece
            for t in (t1, t2):
                if t:
                    t.row, t.col = (dst if (t.row, t.col) == src else src)
            return _relayout_rows(self.plan.slots)

        # Swapped crates take their new position's normal spot (manual moves cleared)
        s1.x_in = s1.y_in = s2.x_in = s2.y_in = None
        new_len = exchange()
        problems = _layout_problems(self.plan.slots)
        if new_len > TRAILER["interior_length_in"] + FIT_TOLERANCE_IN or problems:
            exchange()
            (s1.x_in, s1.y_in), (s2.x_in, s2.y_in) = saved
            why = (f"the load would be {new_len:.0f}\" long, more than the "
                   f"{TRAILER['interior_length_in']}\" trailer" if not problems
                   else problems[0] + " (a crate you moved is in the way)")
            QToolTip.showText(QCursor.pos(), f"Swap not allowed — {why}.", self)
            self.update()
            return
        _refresh_plan_metrics(self.plan)
        self.plan_changed.emit()
        self.update()

    # ── Mouse events ─────────────────────────────────────────────────────────
    def _draggable(self, key) -> bool:
        """Side-by-side crates can be swapped (Ctrl+drag)."""
        s = self._floor_slot(key)
        return bool(s and not s.is_center)

    def mousePressEvent(self, event):
        if event.button() != Qt.MouseButton.LeftButton:
            return
        key = self._slot_at(event.pos())
        slot = self._floor_slot(key)
        if slot is None:
            return
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            if self._draggable(key):
                self._drag_src = key
        else:
            self._start_move(slot, event.pos())
        self.setCursor(Qt.CursorShape.ClosedHandCursor)

    def mouseMoveEvent(self, event):
        if self._move:
            self._update_move(event.pos())
            return
        key = self._slot_at(event.pos())
        if key != self._hover_slot:
            self._hover_slot = key
            self.update()
        if self._drag_src:
            tgt = next((s for s in self.plan.slots
                        if (s.row, s.col) == key and s.tier == 0), None) if key else None
            can_drop = (key and key != self._drag_src and
                        tgt is not None and not tgt.is_center)
            self.setCursor(Qt.CursorShape.DragMoveCursor if can_drop
                           else Qt.CursorShape.ClosedHandCursor)
        elif self._floor_slot(key):
            self.setCursor(Qt.CursorShape.OpenHandCursor)
            slot = self._floor_slot(key)
            self.setToolTip(f"{slot.piece.crate_label}: drag to shift"
                            + ("  ·  Ctrl+drag to swap" if not slot.is_center else ""))
        else:
            self.setCursor(Qt.CursorShape.ArrowCursor)
            self.setToolTip("")

    def mouseReleaseEvent(self, event):
        if event.button() != Qt.MouseButton.LeftButton:
            return
        if self._move:
            self._finish_move()
        elif self._drag_src:
            dst = self._slot_at(event.pos())
            if dst and dst != self._drag_src:
                self._swap_pieces(self._drag_src, dst)
            self._drag_src = None
        self.setCursor(Qt.CursorShape.ArrowCursor)
        self.update()

    def leaveEvent(self, event):
        if self._move:          # keep a drag alive if the mouse briefly leaves
            return
        self._drag_src = None
        self._hover_slot = None
        self.setCursor(Qt.CursorShape.ArrowCursor)
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        self._geom = {}
        self._hits = _draw_top_plan(painter, QRectF(self.rect()).adjusted(8, 4, -8, -4),
                                    self.plan, self.securement, geom=self._geom)
        if self._hits:   # zoom buttons centre on the trailer, not the notes column
            self.zoom_focus_x = (min(r.left() for r in self._hits.values()) +
                                 max(r.right() for r in self._hits.values())) / 2
        if not self.plan or not self.plan.slots:
            painter.setPen(QColor("#555"))
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter,
                             "No load plan — press Calculate.")
            return
        painter.setBrush(Qt.BrushStyle.NoBrush)
        if self._move and self._geom:
            # Ghost of where the crate will land, with its position readout
            g, m = self._geom, self._move
            x, y = m["pos"]
            _x0, _y0, l, w = m["fp"]
            ghost = QRectF(g["tx"] + y * g["s"], g["ty"] + x * g["s"], w * g["s"], l * g["s"])
            painter.setBrush(QBrush(QColor(243, 156, 18, 70)))
            painter.setPen(QPen(QColor("#f39c12"), 2.5, Qt.PenStyle.DashLine))
            painter.drawRect(ghost)
            W_in = TRAILER["interior_width_in"]
            label = (f"{m['slot'].piece.crate_label}: {x:.0f}\" from nose · "
                     f"{y:.0f}\" from left · {W_in - y - w:.0f}\" from right")
            painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
            fm = painter.fontMetrics()
            tag = QRectF(ghost.right() + 6, ghost.center().y() - 10,
                         fm.horizontalAdvance(label) + 12, 20)
            painter.setBrush(QBrush(QColor(255, 255, 255, 235)))
            painter.setPen(QPen(QColor("#f39c12"), 1))
            painter.drawRoundedRect(tag, 4, 4)
            painter.setPen(QColor("#333"))
            painter.drawText(tag, Qt.AlignmentFlag.AlignCenter, label)
        elif self._drag_src in self._hits:
            painter.setPen(QPen(QColor("#f39c12"), 2.5))
            painter.drawRect(self._hits[self._drag_src])
            if self._hover_slot in self._hits and self._hover_slot != self._drag_src:
                painter.setPen(QPen(QColor("#27ae60"), 2.5))
                painter.drawRect(self._hits[self._hover_slot])


# ── Side Elevation View widget ────────────────────────────────────────────────
class SideElevationView(QWidget):
    """
    Live to-scale side elevation, nose at left, viewed from the left wall.
    Drag the drive or trailer axles to slide the fifth wheel / tandem; they
    snap to the notches and holes in RIG and emit rig_moved while dragging.
    """
    rig_moved = pyqtSignal(dict)   # {"fw_in", "tandem_in"}

    def __init__(self, parent=None):
        super().__init__(parent)
        self.plan: Optional[LoadPlan] = None
        self.securement: list = []
        self._geom = {}
        self._drag = None          # "drive" / "tandem" while an axle is being dragged
        self.setMouseTracking(True)
        self.setMinimumSize(320, 200)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

    def _axle_at(self, pos) -> Optional[str]:
        g = self._geom
        if not g or not self.plan or not self.plan.slots:
            return None
        if not g["top_y"] <= pos.y() <= g["floor_y"] + 32:
            return None
        for depth, _label, _color, kind in axle_marks(self.plan):
            if kind != "kingpin" and abs(pos.x() - (g["ox"] + depth * g["s"])) <= 8:
                return kind
        return None

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self._drag = self._axle_at(event.position())
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        pos = event.position()
        if self._drag:
            inches = (pos.x() - self._geom["ox"]) / self._geom["s"] - RIG["kingpin_in"]
            fw_opts, holes = rig_positions(legal_only=False)
            rig = dict(self.plan.rig)
            if self._drag == "drive":
                rig["fw_in"] = min(fw_opts, key=lambda f: abs(f - inches))
            else:
                rig["tandem_in"] = min(holes, key=lambda t: abs(t - inches))
            if rig != self.plan.rig:
                self.rig_moved.emit(rig)
            QToolTip.showText(event.globalPosition().toPoint(), rig_text(rig), self)
            return
        kind = self._axle_at(pos)
        self.setCursor(Qt.CursorShape.SizeHorCursor if kind else Qt.CursorShape.ArrowCursor)
        if kind:
            QToolTip.showText(event.globalPosition().toPoint(),
                              "Drag to slide the " + ("fifth wheel" if kind == "drive" else "trailer tandem"),
                              self)
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self._drag = None
        super().mouseReleaseEvent(event)

    def set_plan(self, plan: Optional[LoadPlan]):
        self.plan = plan
        self.update()

    def set_securement(self, items: list):
        self.securement = items
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        _draw_side_plan(painter, QRectF(self.rect()).adjusted(8, 8, -8, -8),
                        self.plan, self.securement, geom=self._geom)
        if not self.plan or not self.plan.slots:
            painter.setPen(QColor("#555"))
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter,
                             "No load plan — press Calculate.")


# ── 3D View widget ────────────────────────────────────────────────────────────
def _crate_boxes(plan: LoadPlan) -> list:
    """World-space boxes for every placed piece (see _slot_boxes)."""
    return _slot_boxes(plan.slots)


def _slot_boxes(slots: list) -> list:
    """
    World-space boxes for every placed piece: (piece, x, y, z, l, w, h, tier).
    x = inches from nose, y = inches from left wall (Position A), z = inches up.
    Lateral placement mirrors the top view: alt_center alternates left/right by
    visual row, center is centered, side-by-side pieces sit against their wall.
    A floor crate moved by hand (x_in / y_in) carries its stacked crate with it.
    """
    interior_w = TRAILER["interior_width_in"]
    rows = sorted(set(s.row for s in slots))
    r_pos_of = {r: i for i, r in enumerate(rows)}
    bottoms = {(s.row, s.col): s.piece for s in slots if s.tier == 0 and s.piece}

    placed = []   # (slot, piece, x, y, z, l, w, h)
    for s in slots:
        if s.piece is None or (s.is_center and s.col == 1):
            continue
        p = s.piece
        w = min(p.width_in, interior_w)
        if s.is_center:
            gap = interior_w - w
            if p.placement == "alt_center":
                if s.side:
                    y = {"left": 0.0, "center": gap / 2}.get(s.side, gap)
                else:   # plans built before walls were chosen: alternate by row
                    y = 0.0 if r_pos_of[s.row] % 2 == 0 else gap
            else:
                y = gap / 2
        else:
            y = 0.0 if s.col == 0 else interior_w - w
        z = 0.0
        if s.tier == 1:
            base = bottoms.get((s.row, s.col))
            z = base.height_in if base else 0.0
        placed.append((s, p, s.row_start_in, y, z, p.length_in, w, p.height_in))

    # Manual moves: offset each moved floor crate and whatever sits on it
    shift = {}
    for s, p, x, y, *_ in placed:
        if s.tier == 0 and s.x_in is not None:
            shift[(s.row, 0 if s.is_center else s.col)] = (s.x_in - x, s.y_in - y)
    boxes = []
    for s, p, x, y, z, l, w, h in placed:
        dx, dy = shift.get((s.row, 0 if s.is_center else s.col), (0.0, 0.0))
        boxes.append((p, x + dx, y + dy, z, l, w, h, s.tier))
    return boxes


def _layout_problems(slots: list) -> list:
    """Floor crates outside the trailer or overlapping each other (after manual moves)."""
    L, W = TRAILER["interior_length_in"], TRAILER["interior_width_in"]
    floor = [b for b in _slot_boxes(slots) if b[7] == 0]
    out = []
    for p, x, y, z, l, w, h, _t in floor:
        if x < -0.01 or y < -0.01 or x + l > L + 0.01 or y + w > W + 0.01:
            out.append(f"{p.crate_label} is outside the trailer")
    for i, a in enumerate(floor):
        for b2 in floor[i + 1:]:
            if (min(a[1] + a[4], b2[1] + b2[4]) - max(a[1], b2[1]) > 0.01 and
                    min(a[2] + a[5], b2[2] + b2[5]) - max(a[2], b2[2]) > 0.01):
                out.append(f"{a[0].crate_label} overlaps {b2[0].crate_label}")
    return out


def _refresh_plan_metrics(plan: LoadPlan):
    """Recompute length, axle weights and violations from where crates actually sit."""
    boxes = _crate_boxes(plan)
    plan.total_length_used_in = max((b[1] + b[4] for b in boxes if b[7] == 0), default=0.0)
    if plan.rig_auto or not plan.rig:
        loads = [(t, m) for _s, t, m in _freight_loads(plan.slots)]
        plan.rig = auto_rig(loads) if loads[0][0] else default_rig()
    plan.axle_weights, violations = _calc_axle_weights(plan.slots, plan.override, plan.rig)
    if not _rear_axle_ok(plan.rig["tandem_in"]):
        rear, lim = plan.rig["tandem_in"] + RIG["tandem_spread_in"] / 2, RIG["kpra_max_in"]
        violations.append(f"KINGPIN TO REAR AXLE {int(rear // 12)}'{rear % 12:.0f}\" > "
                          f"{int(lim // 12)}'{lim % 12:.0f}\" limit — slide the tandem forward")
    plan.stop_axles = _stop_axles(plan.slots, plan.rig)
    for delivered, steer, drive, trailer in plan.stop_axles:
        for name, val, lim in (("STEER", steer, "steer_limit"), ("DRIVE", drive, "drive_limit"),
                               ("TRAILER TANDEM", trailer, "trailer_limit")):
            if val > TRAILER[lim]:
                violations.append(f"AFTER STOP {delivered} DELIVERY: {name} {val:,.0f} lb > "
                                  f"{TRAILER[lim]:,} lb limit")
    if plan.unplaced:
        violations.insert(0, _not_fit_message(plan.unplaced))
    plan.violations = violations


# ── Securement (load bars, straps, airbags) ──────────────────────────────────
# Each item is anchored to a crate face: {"kind": ..., "piece_id": ..., "face": ...}
# face ∈ nose / tail / left / right / top. Geometry is rebuilt from the current
# plan on every draw, so items follow their crate when the plan is recalculated.
SECUREMENT_KINDS = {"load_bar": "Load Bar", "strap": "Strap", "airbag": "Airbag"}
SECUREMENT_COLORS = {
    "load_bar": QColor("#9aa3ab"),
    "strap":    QColor("#f1c40f"),
    "airbag":   QColor(226, 233, 242, 225),
}
RATCHET_COLOR = QColor("#555555")
FACE_OF_NORMAL = {(-1, 0, 0): "nose", (1, 0, 0): "tail", (0, -1, 0): "left",
                  (0, 1, 0): "right", (0, 0, 1): "top", (0, 0, -1): "bottom"}
AIRBAG_MIN_GAP = 2.0     # inches — smaller voids don't need a bag
AIRBAG_MAX_GAP = 36.0    # inches — larger voids need a load bar or filler


def _free_gap(face: str, plane: float, ranges: dict, boxes: list, skip_pid: int) -> float:
    """
    Clear distance from a crate face (at `plane`) to the next crate or trailer
    wall in the face's direction. `ranges` maps the other two axes (0=x, 1=y,
    2=z) to the (lo, hi) span that an obstacle must overlap to count.
    """
    axis = 0 if face in ("nose", "tail") else 1
    positive = face in ("tail", "right")
    limit = (TRAILER["interior_length_in"] if axis == 0 else TRAILER["interior_width_in"])
    best = limit if positive else 0.0
    for p, x, y, z, l, w, h, _t in boxes:
        if p.piece_id == skip_pid:
            continue
        lo, size = (x, y, z), (l, w, h)
        if any(min(b, lo[ax] + size[ax]) - max(a, lo[ax]) <= 0.5
               for ax, (a, b) in ranges.items()):
            continue
        if positive and lo[axis] >= plane - 0.5:
            best = min(best, lo[axis])
        elif not positive and lo[axis] + size[axis] <= plane + 0.5:
            best = max(best, lo[axis] + size[axis])
    return abs(best - plane)


def _boxes_overlap(a: tuple, b: tuple) -> bool:
    """True if two (x, y, z, l, w, h, ...) parts intersect by more than a hair."""
    return all(min(a[i] + a[i + 3], b[i] + b[i + 3]) - max(a[i], b[i]) > 0.05
               for i in range(3))


MAX_STRAPS_PER_FACE = 2


def securement_parts(item: dict, boxes: list, items: list = None) -> tuple:
    """
    Solid parts for one securement item:
    ([(x, y, z, l, w, h, QColor, shape), ...], error) with shape "box" or "pillow".
    `items` (all securement) lets two straps on one face spread out.
    Returns ([], None) when the anchor crate isn't in the plan (e.g. excluded).
    """
    parts, err = _securement_boxes(item, boxes, items or [item])
    shape = "pillow" if item["kind"] == "airbag" else "box"
    return [pt + (shape,) for pt in parts], err


def _securement_boxes(item: dict, boxes: list, items: list) -> tuple:
    box = next((b for b in boxes if b[0].piece_id == item["piece_id"]), None)
    if box is None:
        return [], None
    p, x, y, z, l, w, h, _t = box
    kind, face = item["kind"], item["face"]
    W = TRAILER["interior_width_in"]
    color = SECUREMENT_COLORS[kind]
    tail = face == "tail"

    if kind == "load_bar":
        zc = z + h * 0.6
        plane = x + l if tail else x
        gap = _free_gap(face, plane, {1: (0, W), 2: (zc - 1.5, zc + 1.5)}, boxes, p.piece_id)
        if gap < 3.0:
            return [], f"No room for a load bar ({gap:.0f}\" clear behind {p.crate_label})"
        x0 = plane + 0.3 if tail else plane - 3.3
        return [(x0, 0, zc - 1.25, 3, W, 2.5, color),
                (x0 - 0.5, 0, zc - 3, 4, 1.5, 6, color),          # wall feet
                (x0 - 0.5, W - 1.5, zc - 3, 4, 1.5, 6, color)], None

    if kind == "strap":
        # Rank among straps sharing this crate face (0 or 1)
        siblings = sorted(it.get("n", 0) for it in items if it["kind"] == "strap"
                          and it["piece_id"] == item["piece_id"] and it["face"] == face)
        rank = siblings.index(item.get("n", 0)) if item.get("n", 0) in siblings else 0
        if face in ("nose", "tail"):
            # Upper strap clears a load bar (0.6h) on the same face; second strap goes low
            zc = z + h * (0.8 if rank == 0 else 0.3)
            plane = x + l if tail else x
            gap = _free_gap(face, plane, {1: (0, W), 2: (zc - 1, zc + 1)}, boxes, p.piece_id)
            if gap < 1.0:
                return [], (f"No room for a strap on the {face} of {p.crate_label} — "
                            f"the next crate is right against it")
            x0 = x + l + 0.1 if tail else x - 0.5
            return [(x0, 0, zc - 1, 0.4, W, 2, color),
                    (x0 - 0.6 if not tail else x0, W * 0.65, zc - 1.5, 1.0, 4, 3, RATCHET_COLOR)], None
        # Over the top, down both sides of the crate; two straps sit at 1/4 and 3/4
        xc = x + l * ((0.5,) if len(siblings) < 2 else (0.25, 0.75))[min(rank, 1)]
        return [(xc - 1, y, z + h, 2, w, 0.4, color),
                (xc - 1, y - 0.4, z, 2, 0.4, h + 0.4, color),
                (xc - 1, y + w, z, 2, 0.4, h + 0.4, color),
                (xc - 1.5, y + w + 0.4, z + h * 0.5, 3, 1.0, 4, RATCHET_COLOR)], None

    if kind == "airbag":
        if face in ("nose", "tail"):
            plane = x + l if tail else x
            gap = _free_gap(face, plane, {1: (y, y + w), 2: (z, z + h)}, boxes, p.piece_id)
        else:
            def side_gap(sd):
                pl = y + w if sd == "right" else y
                return pl, _free_gap(sd, pl, {0: (x, x + l), 2: (z, z + h)}, boxes, p.piece_id)
            plane, gap = side_gap(face)
            other = "right" if face == "left" else "left"
            # An airbag beside a crate follows the gap: if a re-plan moved the crate to
            # the other wall, use the opposite side (unless another airbag is already there)
            if gap < AIRBAG_MIN_GAP and not any(
                    it is not item and it["kind"] == "airbag" and
                    it["piece_id"] == item["piece_id"] and it["face"] == other for it in items):
                o_plane, o_gap = side_gap(other)
                if o_gap >= AIRBAG_MIN_GAP:
                    face, plane, gap = other, o_plane, o_gap
        if gap < AIRBAG_MIN_GAP:
            return [], f"No void to fill on the {face} of {p.crate_label} ({gap:.1f}\")"
        if gap > AIRBAG_MAX_GAP:
            return [], (f"Void on the {face} of {p.crate_label} is {gap:.0f}\" — "
                        f"too large for an airbag. Use a load bar or filler.")
        t = gap - 0.6
        zb, hb = z + h * 0.1, h * 0.8
        if face in ("nose", "tail"):
            x0 = plane + 0.3 if tail else plane - 0.3 - t
            return [(x0, y + w * 0.1, zb, t, w * 0.8, hb, color)], None
        y0 = plane + 0.3 if face == "right" else plane - 0.3 - t
        return [(x + l * 0.1, y0, zb, l * 0.8, t, hb, color)], None

    return [], f"Unknown securement type: {kind}"


class Trailer3DView(QWidget):
    """
    Rotatable 3D view of the loaded trailer (software-rendered with QPainter).
    In a securement tool mode, clicking a crate face adds a load bar / strap /
    airbag against that face; Remove mode deletes the clicked item.
    """
    securement_changed = pyqtSignal()

    PRESETS = {                      # (yaw°, pitch°)
        "3/4 View": (-40.0, 28.0),
        "Top":      (0.0, 89.5),
        "Side":     (-90.0, 0.0),
        "Rear":     (0.0, 6.0),
    }
    TOOL_HINTS = {
        "rotate":   "Drag: rotate  ·  Right-drag: pan  ·  Wheel: zoom  ·  Double-click: reset",
        "load_bar": "Load Bar: click the nose or tail face of a crate  ·  drag still rotates",
        "strap":    "Strap: click a crate's nose/tail face (across) or top (over the top)",
        "airbag":   "Airbag: click a crate face next to a gap (side gap to wall, or between rows)",
        "remove":   "Remove: click a load bar, strap or airbag to delete it",
    }
    CLICK_SLOP = 4   # px of mouse movement still treated as a click

    def __init__(self, parent=None):
        super().__init__(parent)
        self.plan: Optional[LoadPlan] = None
        self.securement: list = []
        self.tool = "rotate"
        self.setMinimumSize(320, 300)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.setMouseTracking(True)
        self._last_mouse = None
        self._press_pos = None
        self._dragged = False
        self._drag_button = None
        self._fit_focal = None
        self._pick: list = []        # [(QPolygonF, tag)] in paint order
        self._hover = None           # (QPolygonF, tag) under the cursor
        self._message = ("", None)   # (text, QColor) feedback line
        self.print_mode = False      # True when rendering into a PDF: no hints/overlay
        self.reset_view()

    def set_plan(self, plan: Optional[LoadPlan]):
        self.plan = plan
        self.update()

    def set_securement(self, items: list):
        self.securement = items
        self.update()

    def set_tool(self, tool: str):
        self.tool = tool
        self._hover = None
        self._message = ("", None)
        self.setCursor(Qt.CursorShape.ArrowCursor if tool == "rotate"
                       else Qt.CursorShape.CrossCursor)
        self.update()

    def set_preset(self, name: str):
        self.yaw, self.pitch = self.PRESETS[name]
        self.zoom = 1.0
        self.target = [TRAILER["interior_length_in"] / 2,
                       TRAILER["interior_width_in"] / 2,
                       TRAILER["interior_height_in"] / 2]
        self._fit_focal = None   # re-fit the trailer to the window for this angle
        self.update()

    def reset_view(self):
        self.set_preset("3/4 View")

    def resizeEvent(self, event):
        self._fit_focal = None
        super().resizeEvent(event)

    # ── Securement editing ───────────────────────────────────────────────────
    def add_securement(self, kind: str, piece_id: int, face: str) -> bool:
        if face == "bottom":
            face = "top"
        if kind == "load_bar" and face not in ("nose", "tail"):
            face = "tail"
        elif kind == "strap" and face in ("left", "right"):
            face = "top"
        elif kind == "airbag" and face == "top":
            self._message = ("Airbags go in gaps — click a side, nose or tail face.",
                             QColor("#c0392b"))
            return False

        same = [it for it in self.securement if it["kind"] == kind
                and it["piece_id"] == piece_id and it["face"] == face]
        limit = MAX_STRAPS_PER_FACE if kind == "strap" else 1
        if len(same) >= limit:
            self._message = (f"Already {limit} straps on that face." if limit > 1
                             else "That item is already there.", QColor("#c0392b"))
            return False
        used = {it.get("n", 0) for it in same}
        item = {"kind": kind, "piece_id": piece_id, "face": face,
                "n": min(set(range(limit)) - used)}
        boxes = _crate_boxes(self.plan) if self.plan else []
        trial = self.securement + [item]
        parts, err = securement_parts(item, boxes, trial)
        if err:
            self._message = (err, QColor("#c0392b"))
            return False
        for other in self.securement:
            other_parts, _ = securement_parts(other, boxes, trial)
            if any(_boxes_overlap(a, b) for a in parts for b in other_parts):
                self._message = (f"That would overlap an existing "
                                 f"{SECUREMENT_KINDS[other['kind']].lower()}.",
                                 QColor("#c0392b"))
                return False

        self.securement.append(item)
        label = next((b[0].crate_label for b in boxes if b[0].piece_id == piece_id), "")
        nth = " (2nd strap on this face)" if kind == "strap" and same else ""
        self._message = (f"Added {SECUREMENT_KINDS[kind].lower()} — {label} {face}{nth}",
                         QColor("#27ae60"))
        self.securement_changed.emit()
        return True

    def unfit_indices(self) -> list:
        """Indices of this truck's securement items that don't fit the current layout."""
        if not self.plan:
            return []
        boxes = _crate_boxes(self.plan)
        ids = {b[0].piece_id for b in boxes}
        return [i for i, it in enumerate(self.securement)
                if it["piece_id"] in ids and securement_parts(it, boxes, self.securement)[1]]

    def remove_unfit(self) -> int:
        bad = self.unfit_indices()
        for i in reversed(bad):
            self.securement.pop(i)
        if bad:
            self._message = (f"Removed {len(bad)} hidden item(s)", QColor("#27ae60"))
            self.securement_changed.emit()
        self.update()
        return len(bad)

    def remove_securement(self, index: int):
        if 0 <= index < len(self.securement):
            item = self.securement.pop(index)
            self._message = (f"Removed {SECUREMENT_KINDS[item['kind']].lower()}",
                             QColor("#27ae60"))
            self.securement_changed.emit()

    def _tag_at(self, pos):
        """Topmost picked face under `pos`: (QPolygonF, tag) or None."""
        for poly, tag in reversed(self._pick):
            if poly.containsPoint(pos, Qt.FillRule.OddEvenFill):
                return poly, tag
        return None

    def _hover_target(self, pos):
        """What a click at `pos` would act on in the current tool mode."""
        if self.tool == "rotate":
            return None
        hit = self._tag_at(pos)
        if not hit:
            return None
        want = "item" if self.tool == "remove" else "crate"
        return hit if hit[1][0] == want else None

    def _click(self, pos):
        target = self._hover_target(pos)
        if target is None:
            if self.tool == "remove":
                self._message = ("Click directly on a load bar, strap or airbag.", QColor("#777"))
            else:
                self._message = ("Click on a crate face.", QColor("#777"))
        elif self.tool == "remove":
            self.remove_securement(target[1][1])
        else:
            _, pid, face = target[1]
            self.add_securement(self.tool, pid, face)
        self.update()
        self._hover = self._hover_target(pos)

    # ── Camera ───────────────────────────────────────────────────────────────
    def _camera(self):
        L, W, H = (TRAILER["interior_length_in"], TRAILER["interior_width_in"],
                   TRAILER["interior_height_in"])
        diag = math.sqrt(L * L + W * W + H * H)
        base_d = diag * 2.2
        dist = base_d / self.zoom
        yaw, pitch = math.radians(self.yaw), math.radians(self.pitch)
        off = (math.cos(pitch) * math.cos(yaw),
               math.cos(pitch) * math.sin(yaw),
               math.sin(pitch))
        cam = [self.target[i] + dist * off[i] for i in range(3)]
        fwd = [-o for o in off]
        # right = fwd × up(Z); up = right × fwd
        right = [fwd[1], -fwd[0], 0.0]
        rn = math.hypot(right[0], right[1]) or 1e-9
        right = [right[0] / rn, right[1] / rn, 0.0]
        up = [right[1] * fwd[2] - right[2] * fwd[1],
              right[2] * fwd[0] - right[0] * fwd[2],
              right[0] * fwd[1] - right[1] * fwd[0]]
        if self._fit_focal is None:
            self._fit_focal = self._compute_fit(cam, fwd, right, up)
        return cam, fwd, right, up, self._fit_focal

    def _compute_fit(self, cam, fwd, right, up) -> float:
        """Focal length that makes the whole trailer fill ~90% of the widget."""
        L, W, H = (TRAILER["interior_length_in"], TRAILER["interior_width_in"],
                   TRAILER["interior_height_in"])
        xs, ys = [], []
        for pt in [(x, y, z) for x in (0, L) for y in (0, W) for z in (0, H)]:
            d = [pt[i] - cam[i] for i in range(3)]
            zc = sum(d[i] * fwd[i] for i in range(3))
            if zc < 1.0:
                continue
            xs.append(sum(d[i] * right[i] for i in range(3)) / zc)
            ys.append(sum(d[i] * up[i] for i in range(3)) / zc)
        if not xs:
            return float(min(self.width(), self.height()))
        # The view is centred on the trailer, so fit the larger half-extent each way
        half_w = max(abs(min(xs)), abs(max(xs))) or 1e-6
        half_h = max(abs(min(ys)), abs(max(ys))) or 1e-6
        return 0.45 * min(self.width() / half_w, (self.height() - 40) / half_h)

    def _project(self, pt, cam_data):
        cam, fwd, right, up, focal = cam_data
        d = [pt[0] - cam[0], pt[1] - cam[1], pt[2] - cam[2]]
        zc = d[0] * fwd[0] + d[1] * fwd[1] + d[2] * fwd[2]
        if zc < 1.0:
            return None
        xc = d[0] * right[0] + d[1] * right[1] + d[2] * right[2]
        yc = d[0] * up[0] + d[1] * up[1] + d[2] * up[2]
        return QPointF(self.width() / 2 + focal * xc / zc,
                       self.height() / 2 - focal * yc / zc)

    def _poly(self, pts, cam_data):
        out = [self._project(p, cam_data) for p in pts]
        return None if any(q is None for q in out) else out

    # ── Mouse: left-drag rotate, right/middle-drag pan, wheel zoom, click = tool ─
    def mousePressEvent(self, event):
        self._last_mouse = event.position()
        self._press_pos = event.position()
        self._dragged = False
        self._drag_button = event.button()

    def mouseMoveEvent(self, event):
        pos = event.position()
        if self._last_mouse is None:
            # No button held — just track what a click would hit
            hover = self._hover_target(pos)
            if (hover and hover[1]) != (self._hover and self._hover[1]):
                self._hover = hover
                self.update()
            return
        if not self._dragged:
            if (abs(pos.x() - self._press_pos.x()) + abs(pos.y() - self._press_pos.y())
                    < self.CLICK_SLOP):
                return
            self._dragged = True
            self._hover = None
            self.setCursor(Qt.CursorShape.ClosedHandCursor)
        dx, dy = pos.x() - self._last_mouse.x(), pos.y() - self._last_mouse.y()
        self._last_mouse = pos
        if self._drag_button == Qt.MouseButton.LeftButton:
            self.yaw = (self.yaw - dx * 0.4) % 360
            self.pitch = max(-10.0, min(89.5, self.pitch + dy * 0.4))
        else:
            cam, fwd, right, up, focal = self._camera()
            dist = math.dist(cam, self.target)
            scale = dist / focal
            for i in range(3):
                self.target[i] += (-dx * right[i] + dy * up[i]) * scale
        self.update()

    def mouseReleaseEvent(self, event):
        was_click = (not self._dragged and self._drag_button == Qt.MouseButton.LeftButton)
        self._last_mouse = None
        self._press_pos = None
        self._drag_button = None
        self._dragged = False
        self.setCursor(Qt.CursorShape.ArrowCursor if self.tool == "rotate"
                       else Qt.CursorShape.CrossCursor)
        if was_click and self.tool != "rotate":
            self._click(event.position())

    def mouseDoubleClickEvent(self, event):
        if self.tool == "rotate":
            self.reset_view()

    def leaveEvent(self, event):
        if self._hover:
            self._hover = None
            self.update()

    def wheelEvent(self, event):
        steps = event.angleDelta().y() / 120
        self.zoom = max(0.5, min(2.5, self.zoom * (1.15 ** steps)))
        self.update()

    # ── Drawing ──────────────────────────────────────────────────────────────
    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.fillRect(self.rect(), QColor("#fafafa"))
        self._draw(painter)

    def _draw_solid(self, painter, cd, x0, y0, z0, x1, y1, z1, base: QColor, tag):
        """Draw the camera-facing faces of an axis-aligned box and record them for picking."""
        cam = cd[0]
        light = (0.35, -0.45, 0.82)
        faces = [
            ((-1, 0, 0), [(x0, y0, z0), (x0, y1, z0), (x0, y1, z1), (x0, y0, z1)]),
            ((1, 0, 0),  [(x1, y0, z0), (x1, y0, z1), (x1, y1, z1), (x1, y1, z0)]),
            ((0, -1, 0), [(x0, y0, z0), (x0, y0, z1), (x1, y0, z1), (x1, y0, z0)]),
            ((0, 1, 0),  [(x0, y1, z0), (x1, y1, z0), (x1, y1, z1), (x0, y1, z1)]),
            ((0, 0, -1), [(x0, y0, z0), (x1, y0, z0), (x1, y1, z0), (x0, y1, z0)]),
            ((0, 0, 1),  [(x0, y0, z1), (x0, y1, z1), (x1, y1, z1), (x1, y0, z1)]),
        ]
        for n, pts in faces:
            fc = pts[0]
            to_cam = (cam[0] - fc[0], cam[1] - fc[1], cam[2] - fc[2])
            if n[0] * to_cam[0] + n[1] * to_cam[1] + n[2] * to_cam[2] <= 0:
                continue  # back face
            poly = self._poly(pts, cd)
            if not poly:
                continue
            shade = 0.55 + 0.45 * max(0.0, n[0] * light[0] + n[1] * light[1] + n[2] * light[2])
            painter.setBrush(QBrush(QColor(int(base.red() * shade), int(base.green() * shade),
                                           int(base.blue() * shade), base.alpha())))
            painter.setPen(QPen(QColor("#222"), 0.8))
            painter.drawPolygon(poly)
            if tag is not None:
                face_tag = tag + (FACE_OF_NORMAL[n],) if tag[0] == "crate" else tag
                self._pick.append((QPolygonF(poly), face_tag))

    _PILLOW_MESH = None

    @classmethod
    def _pillow_mesh(cls):
        """Unit superellipsoid quads (flat faces, rounded edges). It is convex,
        so back-face culling alone gives the correct drawing order."""
        if cls._PILLOW_MESH is None:
            e, NU, NV = 0.45, 8, 16

            def sp(t):
                return math.copysign(abs(t) ** e, t)

            def pt(i, j):
                uu = -math.pi / 2 + math.pi * i / NU
                vv = -math.pi + 2 * math.pi * j / NV
                return (sp(math.cos(uu)) * sp(math.cos(vv)),
                        sp(math.cos(uu)) * sp(math.sin(vv)),
                        sp(math.sin(uu)))
            cls._PILLOW_MESH = [(pt(i, j), pt(i, j + 1), pt(i + 1, j + 1), pt(i + 1, j))
                                for i in range(NU) for j in range(NV)]
        return cls._PILLOW_MESH

    def _draw_pillow(self, painter, cd, x0, y0, z0, x1, y1, z1, base: QColor, tag):
        """Airbag: a puffy rounded slab filling the box (x0..x1, y0..y1, z0..z1)."""
        cam = cd[0]
        c = ((x0 + x1) / 2, (y0 + y1) / 2, (z0 + z1) / 2)
        half = ((x1 - x0) / 2, (y1 - y0) / 2, (z1 - z0) / 2)
        light = (0.35, -0.45, 0.82)
        for quad in self._pillow_mesh():
            pts = [tuple(c[k] + half[k] * q[k] for k in range(3)) for q in quad]
            a, b, d = pts[0], pts[1], pts[3]
            if math.dist(a, b) < 1e-6:     # degenerate edge at the poles
                b = pts[2]
            e1 = [b[k] - a[k] for k in range(3)]
            e2 = [d[k] - a[k] for k in range(3)]
            n = [e1[1] * e2[2] - e1[2] * e2[1], e1[2] * e2[0] - e1[0] * e2[2],
                 e1[0] * e2[1] - e1[1] * e2[0]]
            qc = [sum(p[k] for p in pts) / 4 for k in range(3)]
            if sum(n[k] * (qc[k] - c[k]) for k in range(3)) < 0:
                n = [-v for v in n]           # make it point outward
            nl = math.sqrt(sum(v * v for v in n)) or 1e-9
            n = [v / nl for v in n]
            if sum(n[k] * (cam[k] - qc[k]) for k in range(3)) <= 0:
                continue  # back face
            poly = self._poly(pts, cd)
            if not poly:
                continue
            shade = 0.6 + 0.4 * max(0.0, sum(n[k] * light[k] for k in range(3)))
            col = QColor(int(base.red() * shade), int(base.green() * shade),
                         int(base.blue() * shade), base.alpha())
            painter.setBrush(QBrush(col))
            painter.setPen(QPen(col, 0.6))   # same-colour pen hides mesh seams
            painter.drawPolygon(poly)
            if tag is not None:
                self._pick.append((QPolygonF(poly), tag))

    def _draw(self, painter: QPainter):
        L, W, H = (TRAILER["interior_length_in"], TRAILER["interior_width_in"],
                   TRAILER["interior_height_in"])
        cd = self._camera()
        cam = cd[0]
        self._pick = []

        # Floor
        floor = self._poly([(0, 0, 0), (L, 0, 0), (L, W, 0), (0, W, 0)], cd)
        if floor:
            painter.setPen(QPen(QColor("#888"), 1))
            painter.setBrush(QBrush(QColor("#ebe6d8")))
            painter.drawPolygon(floor)

        # Nose wall (light fill so the front of the trailer reads clearly)
        nose = self._poly([(0, 0, 0), (0, W, 0), (0, W, H), (0, 0, H)], cd)
        if nose:
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QBrush(QColor(68, 68, 68, 28)))
            painter.drawPolygon(nose)

        # Axle reference lines on the floor
        for depth, label, color, kind in axle_marks(self.plan):
            if kind == "kingpin":
                continue
            seg = self._poly([(depth, 0, 0.2), (depth, W, 0.2)], cd)
            if seg:
                painter.setPen(QPen(QColor(color), 1.5, Qt.PenStyle.DashLine))
                painter.drawLine(seg[0], seg[1])
                painter.setFont(QFont("Arial", 7))
                painter.setPen(QColor(color))
                painter.drawText(seg[1] + QPointF(4, 0), label)

        # Trailer frame edges
        corners = [(x, y, z) for x in (0, L) for y in (0, W) for z in (0, H)]
        painter.setPen(QPen(QColor("#666"), 1.2))
        for i, a in enumerate(corners):
            for b in corners[i + 1:]:
                if sum(1 for k in range(3) if a[k] != b[k]) == 1:
                    seg = self._poly([a, b], cd)
                    if seg:
                        painter.drawLine(seg[0], seg[1])

        for pt, text in [((0, W / 2, H + 8), "NOSE"), ((L, W / 2, H + 8), "DOORS")]:
            q = self._project(pt, cd)
            if q:
                painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
                painter.setPen(QColor("#333"))
                painter.drawText(QRectF(q.x() - 40, q.y() - 14, 80, 14),
                                 Qt.AlignmentFlag.AlignCenter, text)

        if not self.plan or not self.plan.slots:
            painter.setPen(QColor("#555"))
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter,
                             "No load plan — press Calculate.")
            self._draw_overlay(painter, 0)
            return

        # Solids to draw: crates (with labels) and securement parts
        boxes = _crate_boxes(self.plan)
        multi = len(plan_stops(self.plan)) > 1
        solids = []   # (centre, (x0,y0,z0,x1,y1,z1), color, tag, label)
        GAP = 0.6     # small inset so neighbouring crates show a seam
        for piece, x, y, z, l, w, h, tier in boxes:
            solids.append(((x + l / 2, y + w / 2, z + h / 2),
                           (x + GAP, y + GAP, z + GAP, x + l - GAP, y + w - GAP, z + h - GAP),
                           QColor(CRATE_COLORS[(piece.piece_id - 1) % len(CRATE_COLORS)]),
                           ("crate", piece.piece_id),
                           piece.crate_label + (" ▲" if tier == 1 else "")
                           + (f" · S{piece.stop}" if multi else "")))
        labels = {sd[3]: sd[4] for sd in solids}
        solids = [sd[:4] + ("box",) for sd in solids]
        n_unfit = 0
        for idx, item in enumerate(self.securement):
            parts, err = securement_parts(item, boxes, self.securement)
            if err:
                n_unfit += 1
            for x, y, z, l, w, h, color, shape in parts:
                solids.append(((x + l / 2, y + w / 2, z + h / 2),
                               (x, y, z, x + l, y + w, z + h), color, ("item", idx), shape))

        # Painter's algorithm: farthest solid first
        solids.sort(key=lambda s: -math.dist(cam, s[0]))
        for centre, (x0, y0, z0, x1, y1, z1), color, tag, shape in solids:
            if shape == "pillow":
                self._draw_pillow(painter, cd, x0, y0, z0, x1, y1, z1, color, tag)
            else:
                self._draw_solid(painter, cd, x0, y0, z0, x1, y1, z1, color, tag)
            label = labels.get(tag) if tag[0] == "crate" else None
            if label:   # drawn now so nearer solids cover it
                c = self._project(centre, cd)
                if c:
                    painter.setFont(QFont("Arial", 7, QFont.Weight.Bold))
                    painter.setPen(Qt.GlobalColor.white)
                    painter.drawText(QRectF(c.x() - 45, c.y() - 9, 90, 18),
                                     Qt.AlignmentFlag.AlignCenter, label)

        # Hover highlight for the current tool
        if self._hover:
            painter.setBrush(QBrush(QColor(255, 255, 255, 60)))
            painter.setPen(QPen(QColor("#c0392b") if self.tool == "remove"
                                else QColor("#f1c40f"), 3))
            painter.drawPolygon(self._hover[0])

        self._draw_overlay(painter, n_unfit)

    def _draw_overlay(self, painter: QPainter, n_unfit: int):
        if self.print_mode:
            return
        painter.setFont(QFont("Arial", 7))
        painter.setPen(QColor("#777"))
        painter.drawText(QRectF(8, self.height() - 18, self.width() - 16, 16),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                         self.TOOL_HINTS[self.tool])
        lines = []
        if self._message[0]:
            lines.append(self._message)
        if n_unfit:
            lines.append((f"⚠ {n_unfit} securement item(s) don't fit the current plan "
                          f"and are hidden — click Remove Hidden to clear them.", QColor("#c0392b")))
        painter.setFont(QFont("Arial", 8, QFont.Weight.Bold))
        for i, (text, color) in enumerate(lines):
            painter.setPen(color)
            painter.drawText(QRectF(8, 6 + i * 16, self.width() - 16, 16),
                             Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, text)


def _draw_3d_snapshot(painter: QPainter, rect: QRectF, plan: LoadPlan,
                      securement: list, yaw: float, pitch: float):
    """Render the 3D view at a fixed angle into `rect` (vector output, e.g. a PDF)."""
    view = Trailer3DView()
    view.print_mode = True
    view.resize(int(rect.width()), int(rect.height()))
    view.set_plan(plan)
    view.set_securement(securement or [])
    view.yaw, view.pitch = yaw, pitch
    view._fit_focal = None
    painter.save()
    painter.translate(rect.topLeft())
    painter.setClipRect(QRectF(0, 0, rect.width(), rect.height()))
    view._draw(painter)
    painter.restore()
    view.deleteLater()


# ── Zoomable container for the Top / Side views ───────────────────────────────
class ZoomScrollArea(QScrollArea):
    """
    Hosts a view that draws to fit its own size. Zooming enlarges the view
    beyond the viewport and the scroll bars take over. Ctrl+wheel zooms about
    the cursor; right-drag pans.
    """
    zoom_changed = pyqtSignal(float)
    MIN_ZOOM, MAX_ZOOM = 1.0, 6.0

    def __init__(self, view: QWidget, parent=None):
        super().__init__(parent)
        self.setWidgetResizable(True)
        self.setWidget(view)
        self.zoom = 1.0
        self._pan_last = None
        view.installEventFilter(self)

    def set_zoom(self, zoom: float, anchor: Optional[QPointF] = None):
        zoom = max(self.MIN_ZOOM, min(self.MAX_ZOOM, zoom))
        view = self.widget()
        hb, vb = self.horizontalScrollBar(), self.verticalScrollBar()
        centre_x = False
        if anchor is None:
            anchor = QPointF(self.viewport().width() / 2, self.viewport().height() / 2)
            # Views can name the x they care about (e.g. the trailer, not the notes);
            # button zooms then keep that x centred in the window
            focus_x = getattr(view, "zoom_focus_x", None)
            if focus_x is not None:
                anchor.setX(focus_x - hb.value())
                centre_x = True
        # Keep the point under the anchor fixed while the view grows/shrinks
        fx = (hb.value() + anchor.x()) / max(1, view.width())
        fy = (vb.value() + anchor.y()) / max(1, view.height())
        self.zoom = zoom
        self._apply_zoom()
        # Scroll ranges update lazily; set them now so the anchor can be honoured
        hb.setRange(0, max(0, view.width() - self.viewport().width()))
        vb.setRange(0, max(0, view.height() - self.viewport().height()))
        target_x = self.viewport().width() / 2 if centre_x else anchor.x()
        hb.setValue(int(fx * view.width() - target_x))
        vb.setValue(int(fy * view.height() - anchor.y()))
        self.zoom_changed.emit(zoom)

    def _apply_zoom(self):
        vp = self.maximumViewportSize()
        size = QSize(int(vp.width() * self.zoom), int(vp.height() * self.zoom))
        self.widget().setMinimumSize(size)
        self.widget().resize(size)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply_zoom()

    def wheelEvent(self, event):
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            steps = event.angleDelta().y() / 120
            self.set_zoom(self.zoom * (1.2 ** steps), event.position())
            event.accept()
        else:
            super().wheelEvent(event)

    def eventFilter(self, obj, event):
        et = event.type()
        if et == QEvent.Type.MouseButtonPress and event.button() == Qt.MouseButton.RightButton:
            self._pan_last = event.globalPosition()
            obj.setCursor(Qt.CursorShape.ClosedHandCursor)
            return True
        if et == QEvent.Type.MouseMove and self._pan_last is not None:
            pos = event.globalPosition()
            self.horizontalScrollBar().setValue(
                int(self.horizontalScrollBar().value() - (pos.x() - self._pan_last.x())))
            self.verticalScrollBar().setValue(
                int(self.verticalScrollBar().value() - (pos.y() - self._pan_last.y())))
            self._pan_last = pos
            return True
        if et == QEvent.Type.MouseButtonRelease and event.button() == Qt.MouseButton.RightButton:
            self._pan_last = None
            obj.setCursor(Qt.CursorShape.ArrowCursor)
            return True
        return super().eventFilter(obj, event)


# ── Delivery stops editor ─────────────────────────────────────────────────────
class StopsDialog(QDialog):
    """
    Origin + up to MAX_STOPS delivery stops. Each takes a street address
    (located online via OpenStreetMap) or GPS coordinates (used as typed —
    for job sites without an address). Shows the resulting delivery route.
    """

    def __init__(self, origin: dict, stops: list, use_road: bool, road: Optional[dict],
                 parent=None):
        super().__init__(parent)
        self.setWindowTitle("Delivery Stops & Route")
        self.resize(900, 560)
        self.origin = dict(origin)
        self.stops = [dict(st) for st in stops]
        self.road = road
        self._loading = True

        lay = QVBoxLayout(self)
        hint = QLabel("Enter a street address <b>or</b> GPS coordinates (e.g. <i>42.9634, -85.6681</i> "
                      "or a Google Maps link) for the origin and each stop. Coordinates work offline; "
                      "addresses are looked up online with OpenStreetMap when you click "
                      "<b>Locate</b>. Delivery order is planned nearest-first from the origin.")
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#444;")
        lay.addWidget(hint)

        self.table = QTableWidget(MAX_STOPS + 1, 4)
        self.table.setHorizontalHeaderLabels(["", "Name", "Address or GPS coordinates", "Location"])
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.setColumnWidth(0, 120)
        self.table.setColumnWidth(1, 170)
        self.table.setColumnWidth(3, 200)
        self.table.verticalHeader().setVisible(False)
        for r in range(MAX_STOPS + 1):
            d = self.origin if r == 0 else self.stops[r - 1]
            tag = QTableWidgetItem("Origin (ship from)" if r == 0 else f"Stop S{r}")
            tag.setFlags(Qt.ItemFlag.ItemIsEnabled)
            if r:
                tag.setForeground(QColor(STOP_COLORS[(r - 1) % len(STOP_COLORS)]))
            self.table.setItem(r, 0, tag)
            self.table.setItem(r, 1, QTableWidgetItem(d.get("name", "")))
            self.table.setItem(r, 2, QTableWidgetItem(d.get("location", "")))
            loc = QTableWidgetItem("")
            loc.setFlags(Qt.ItemFlag.ItemIsEnabled)
            self.table.setItem(r, 3, loc)
            self._show_location(r)
        self.table.itemChanged.connect(self._cell_changed)
        lay.addWidget(self.table)

        row = QHBoxLayout()
        btn_locate = QPushButton("Locate")
        btn_locate.setToolTip("Read GPS coordinates and look up addresses online (OpenStreetMap)")
        btn_locate.clicked.connect(self._locate_all)
        self.chk_road = QCheckBox("Use road distances (online)")
        self.chk_road.setChecked(use_road)
        self.chk_road.toggled.connect(self._road_toggled)
        btn_road = QPushButton("Fetch Road Distances")
        btn_road.clicked.connect(self._fetch_road)
        row.addWidget(btn_locate)
        row.addSpacing(16)
        row.addWidget(self.chk_road)
        row.addWidget(btn_road)
        row.addStretch()
        lay.addLayout(row)

        self.lbl_route = QLabel()
        self.lbl_route.setWordWrap(True)
        self.lbl_route.setStyleSheet("background:#f4f8f7; border:1px solid #cfe3df; padding:6px;")
        lay.addWidget(self.lbl_route)

        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok |
                                QDialogButtonBox.StandardButton.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        lay.addWidget(btns)
        self._loading = False
        self._update_route()

    # ── helpers ──────────────────────────────────────────────────────────
    def _entry(self, r: int) -> dict:
        return self.origin if r == 0 else self.stops[r - 1]

    def _show_location(self, r: int, status: str = ""):
        d = self._entry(r)
        item = self.table.item(r, 3)
        if d.get("lat") is not None:
            item.setText(f"✓ {d['lat']:.5f}, {d['lon']:.5f}")
            item.setForeground(QColor("#1e7b34"))
        elif d.get("location"):
            item.setText(status or "not located — click Locate")
            item.setForeground(QColor("#b03a2e"))
        else:
            item.setText("")
        item.setToolTip(d.get("display", ""))

    def _cell_changed(self, item):
        if self._loading or item.column() not in (1, 2):
            return
        r = item.row()
        d = self._entry(r)
        if item.column() == 1:
            d["name"] = item.text().strip()
        else:
            d["location"] = item.text().strip()
            c = parse_coords(d["location"])
            d["lat"], d["lon"] = (c if c else (None, None))
            d["display"] = "GPS coordinates" if c else ""
            self.road = None                      # locations changed — road table is stale
            self._loading = True
            self._show_location(r)
            self._loading = False
        self._update_route()

    def _coords(self) -> dict:
        out = {}
        if self.origin.get("lat") is not None:
            out["origin"] = (self.origin["lat"], self.origin["lon"])
        for i, st in enumerate(self.stops, start=1):
            if st.get("lat") is not None:
                out[i] = (st["lat"], st["lon"])
        return out

    def _dist(self):
        return shipment_distance_fn(self._coords(), self.chk_road.isChecked(), self.road)

    def _update_route(self):
        coords = self._coords()
        used = [i for i, st in enumerate(self.stops, start=1)
                if st.get("name") or st.get("location")]
        if not used:
            self.lbl_route.setText("<i>No stops entered yet.</i>")
            return
        dist = self._dist()
        order = route_order(used, coords, dist)
        parts, total, prev = [], 0.0, "origin" if "origin" in coords else None
        for sid in order:
            nm = self.stops[sid - 1].get("name") or f"Stop S{sid}"
            leg = ""
            if prev is not None and sid in coords and prev in coords:
                miles = dist(prev, sid)
                total += miles
                leg = f" <span style='color:#777'>({miles:,.0f} mi)</span>"
            parts.append(f"<b>S{sid}</b> {nm}{leg}")
            prev = sid if sid in coords else prev
        basis = ("road miles" if self.chk_road.isChecked() and self.road else "straight-line miles")
        head = "Origin → " if "origin" in coords else "<i>(no origin located — ordered by stop #)</i> "
        missing = [f"S{sid}" for sid in used if sid not in coords]
        tail = (f"<br><span style='color:#b03a2e'>Not located: {', '.join(missing)} — "
                f"they're delivered last, in stop-number order.</span>" if missing else "")
        self.lbl_route.setText(f"<b>Delivery route</b> ({basis}, ≈{total:,.0f} mi):<br>"
                               f"{head}{' → '.join(parts)}{tail}")

    # ── actions ──────────────────────────────────────────────────────────
    def _locate_all(self):
        todo = [r for r in range(MAX_STOPS + 1)
                if self._entry(r).get("location") and self._entry(r).get("lat") is None]
        if not todo:
            self._update_route()
            return
        failures = []
        self.setCursor(Qt.CursorShape.WaitCursor)
        try:
            for k, r in enumerate(todo):
                d = self._entry(r)
                self._loading = True
                self.table.item(r, 3).setText("looking up…")
                self._loading = False
                QApplication.processEvents()
                try:
                    if k:
                        time.sleep(1.1)           # OpenStreetMap usage policy: ≤ 1 request/sec
                    hit = geocode_address(d["location"])
                except Exception as e:
                    hit, err = None, str(e)
                else:
                    err = "address not found"
                self._loading = True
                if hit:
                    d["lat"], d["lon"], d["display"] = hit
                    self._show_location(r)
                else:
                    self._show_location(r, f"⚠ {err}")
                    failures.append(f"{self.table.item(r, 0).text()}: {err}")
                self._loading = False
        finally:
            self.unsetCursor()
        self.road = None
        self._update_route()
        if failures:
            QMessageBox.warning(self, "Locate", "Some locations couldn't be found:\n\n" +
                                "\n".join(failures) +
                                "\n\nCheck the address, or enter GPS coordinates instead.")

    def _road_toggled(self, on: bool):
        if on and not self.road:
            self._fetch_road()
        self._update_route()

    def _fetch_road(self):
        coords = self._coords()
        keys = list(coords)
        if len(keys) < 2:
            QMessageBox.information(self, "Road Distances", "Locate the origin and at least one stop first.")
            self.chk_road.setChecked(False)
            return
        self.setCursor(Qt.CursorShape.WaitCursor)
        try:
            matrix = road_distance_matrix([coords[k] for k in keys])
        except Exception as e:
            self.unsetCursor()
            QMessageBox.warning(self, "Road Distances",
                                f"Couldn't get road distances (are you online?):\n{e}\n\n"
                                "Straight-line distances will be used.")
            self.chk_road.setChecked(False)
            return
        self.unsetCursor()
        self.road = {"keys": [str(k) for k in keys], "miles": matrix,
                     "points": [list(coords[k]) for k in keys]}
        self.chk_road.setChecked(True)
        self._update_route()


def shipment_distance_fn(coords: dict, use_road: bool, road: Optional[dict]):
    """dist(a, b) in miles: road table when enabled and current, else straight-line."""
    table = None
    if use_road and road:
        keys = road.get("keys", [])
        # Only trust the table if every point still matches what was fetched
        if all(str(k) in keys and
               list(coords[k]) == road["points"][keys.index(str(k))] for k in coords):
            table = {(keys[i], keys[j]): road["miles"][i][j]
                     for i in range(len(keys)) for j in range(len(keys))}

    def dist(a, b):
        if table is not None and (str(a), str(b)) in table:
            return table[(str(a), str(b))]
        return straight_miles(coords[a], coords[b])
    return dist


# ── Ship-separately list ──────────────────────────────────────────────────────
class ShipSeparatelyDialog(QDialog):
    """Pieces that ship another way (UPS, hot-shot, …); return any to a truck."""

    def __init__(self, pieces: list, reasons: dict, n_trucks: int, stop_name, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Ship Separately")
        self.resize(760, 420)
        self.returned: list = []      # [(piece_id, truck_index)]
        lay = QVBoxLayout(self)
        lay.addWidget(QLabel("These pieces aren't on any truck — ship them another way "
                             "(UPS, hot-shot, …), or select some and return them to a truck."))
        self.table = QTableWidget(len(pieces), 6)
        self.table.setHorizontalHeaderLabels(["Crate", "Stop", "L×W×H", "Weight", "Notes", "Why"])
        self.table.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._ids = []
        for r, p in enumerate(pieces):
            self._ids.append(p.piece_id)
            for c, v in enumerate([p.crate_label, stop_name(p.stop),
                                   f"{p.length_in:.0f}×{p.width_in:.0f}×{p.height_in:.0f}\"",
                                   f"{p.weight_lbs:,.0f} lb", p.notes,
                                   reasons.get(p.piece_id, "didn't fit — your choice")]):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        lay.addWidget(self.table)
        row = QHBoxLayout()
        row.addWidget(QLabel(f"Total: {len(pieces)} pcs, "
                             f"{sum(p.weight_lbs for p in pieces):,.0f} lb"))
        row.addStretch()
        self.cmb = QComboBox()
        self.cmb.addItems([f"Truck {i + 1}" for i in range(n_trucks)] + ["New truck"])
        btn = QPushButton("Return Selected to")
        btn.clicked.connect(self._return)
        row.addWidget(btn)
        row.addWidget(self.cmb)
        close = QPushButton("Close")
        close.clicked.connect(self.accept)
        row.addWidget(close)
        lay.addLayout(row)

    def _return(self):
        rows = sorted({i.row() for i in self.table.selectedItems()})
        if not rows:
            return
        self.returned = [(self._ids[r], self.cmb.currentIndex()) for r in rows]
        self.accept()


# ── PDF: whole shipment ───────────────────────────────────────────────────────
def export_shipment_pdf(entries: list, path: str, securement: list = None,
                        include_3d: bool = False, ship_other: list = None):
    """
    entries: [(plan, info)] one per truck. Each truck gets its floor plan, a side
    elevation page when stacked/secured, and optionally the 3D page; a final
    page lists pieces shipping separately.
    """
    writer = _pdf_writer(path)
    painter = QPainter(writer)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    first = True

    def new_page(orientation):
        nonlocal first
        writer.setPageOrientation(orientation)
        if not first:
            writer.newPage()
        first = False

    for plan, info in entries:
        ids = {s.piece.piece_id for s in plan.slots if s.piece}
        sec = [it for it in (securement or []) if it["piece_id"] in ids]
        boxes = _crate_boxes(plan)   # leave out items that don't fit this layout (hidden on screen)
        sec = [it for it in sec if not securement_parts(it, boxes, sec)[1]]
        new_page(QPageLayout.Orientation.Portrait)
        _draw_floor_plan(painter, writer.width(), writer.height(), plan, info, sec)
        if sec or any(s.tier == 1 for s in plan.slots if s.piece):
            new_page(QPageLayout.Orientation.Landscape)
            _draw_side_elevation(painter, writer.width(), writer.height(), plan, info, sec)
        if include_3d:
            new_page(QPageLayout.Orientation.Landscape)
            _draw_3d_page(painter, writer.width(), writer.height(), plan, info, sec)
    if ship_other:
        new_page(QPageLayout.Orientation.Portrait)
        _draw_ship_other_page(painter, writer.width(), writer.height(), ship_other,
                              entries[0][1] if entries else ShipmentInfo())
    painter.end()


def _draw_ship_other_page(painter: QPainter, PW: int, PH: int, rows: list, info: ShipmentInfo):
    """rows: [(piece, stop_text, reason)]"""
    M = 72
    painter.setPen(QColor("#111"))
    painter.setFont(QFont("Arial", 20, QFont.Weight.Bold))
    painter.drawText(QRectF(M, M, PW - 2 * M, 44),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                     "Ship Separately (UPS / Hot-Shot)")
    painter.setFont(QFont("Arial", 10))
    painter.setPen(QColor("#555"))
    wt = sum(p.weight_lbs for p, _, _ in rows)
    painter.drawText(QRectF(M, M + 44, PW - 2 * M, 24),
                     Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignVCenter,
                     "  |  ".join(x for x in [info.customer or info.shipment_name, info.date,
                                               f"{len(rows)} pcs, {wt:,.0f} lb"] if x))
    cols = [("✓", 0.04), ("Crate", 0.13), ("Stop", 0.22), ("L×W×H", 0.14),
            ("Weight", 0.11), ("Notes", 0.18), ("Why", 0.18)]
    y, line_h, x0, width = M + 90, 26, M, PW - 2 * M
    painter.setFont(QFont("Arial", 9, QFont.Weight.Bold))
    painter.setPen(TEAL_DARK)
    cx = x0
    for name, frac in cols:
        painter.drawText(QRectF(cx, y, width * frac, line_h),
                         Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, name)
        cx += width * frac
    painter.drawLine(QPointF(x0, y + line_h), QPointF(x0 + width, y + line_h))
    painter.setFont(QFont("Arial", 9))
    for i, (p, stop_text, why) in enumerate(rows):
        ry = y + line_h * (i + 1)
        if ry > PH - M - line_h:
            break
        if i % 2:
            painter.fillRect(QRectF(x0, ry, width, line_h), QColor("#f2f9f8"))
        painter.setPen(QPen(QColor("#555"), 1))
        painter.setBrush(Qt.BrushStyle.NoBrush)
        painter.drawRect(QRectF(x0 + 4, ry + 7, 12, 12))
        painter.setPen(QColor("#222"))
        cx = x0 + width * cols[0][1]
        for val, (_n, frac) in zip([p.crate_label, stop_text,
                                    f"{p.length_in:.0f}×{p.width_in:.0f}×{p.height_in:.0f}\"",
                                    f"{p.weight_lbs:,.0f} lb", p.notes, why], cols[1:]):
            painter.drawText(QRectF(cx, ry, width * frac - 4, line_h),
                             Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter,
                             QFontMetrics(painter.font(), painter.device()).elidedText(
                                 str(val), Qt.TextElideMode.ElideRight, int(width * frac - 6)))
            cx += width * frac


def export_shipment_labels(entries: list, path: str):
    """Crate labels for every truck: [(plan, info)]."""
    writer = _pdf_writer(path)
    painter = QPainter(writer)
    first = True
    for plan, info in entries:
        seen: set = set()
        for s in sorted(plan.slots, key=lambda s: (s.row, s.tier, s.col)):
            if s.piece and s.piece.piece_id not in seen:
                seen.add(s.piece.piece_id)
                if not first:
                    writer.newPage()
                first = False
                _draw_crate_label(painter, writer.width(), writer.height(), s, plan, info)
    painter.end()


# ── PDF preview ───────────────────────────────────────────────────────────────
class PdfPreviewDialog(QDialog):
    """Shows a generated PDF before it's saved; Save PDF… asks where to keep it."""

    def __init__(self, pdf_path: str, title: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"Preview — {title}")
        self.setWindowFlags(Qt.WindowType.Window | Qt.WindowType.WindowMinMaxButtonsHint |
                            Qt.WindowType.WindowCloseButtonHint)
        self.resize(900, 1000)
        self.save_requested = False
        lay = QVBoxLayout(self)

        # Load from memory so the viewer never holds the temp file open
        with open(pdf_path, "rb") as f:
            self._buffer = QBuffer(self)
            self._buffer.setData(QByteArray(f.read()))
        self._buffer.open(QIODevice.OpenModeFlag.ReadOnly)
        self.doc = QPdfDocument(self)
        self.doc.load(self._buffer)
        self.view = QPdfView(self)
        self.view.setDocument(self.doc)
        self.view.setPageMode(QPdfView.PageMode.MultiPage)
        self.view.setZoomMode(QPdfView.ZoomMode.FitToWidth)

        bar = QHBoxLayout()
        pages = self.doc.pageCount()
        bar.addWidget(QLabel(f"<b>{title}</b>  ·  {pages} page{'s' if pages != 1 else ''}"))
        bar.addStretch()
        for text, mode in (("Fit Width", QPdfView.ZoomMode.FitToWidth),
                           ("Fit Page", QPdfView.ZoomMode.FitInView)):
            b = QPushButton(text)
            b.clicked.connect(lambda _=False, m=mode: self.view.setZoomMode(m))
            bar.addWidget(b)
        for text, factor in (("−", 1 / 1.25), ("+", 1.25)):
            b = QPushButton(text)
            b.setFixedWidth(36)
            b.clicked.connect(lambda _=False, f=factor: self._zoom(f))
            bar.addWidget(b)
        lay.addLayout(bar)
        lay.addWidget(self.view)

        btns = QHBoxLayout()
        btns.addStretch()
        btn_save = QPushButton("Save PDF…")
        btn_save.setDefault(True)
        btn_save.clicked.connect(self._save)
        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.reject)
        btns.addWidget(btn_save)
        btns.addWidget(btn_close)
        lay.addLayout(btns)

    def _zoom(self, factor: float):
        self.view.setZoomMode(QPdfView.ZoomMode.Custom)
        self.view.setZoomFactor(self.view.zoomFactor() * factor)

    def _save(self):
        self.save_requested = True
        self.accept()

    def done(self, result):
        self.view.setDocument(None)
        self.doc.close()
        super().done(result)


# ── Excel Template ────────────────────────────────────────────────────────────
def open_excel_template(parent) -> Optional[str]:
    """Copy the import template to a new packing list file and open it in Excel.

    Returns the new file's path, or None if cancelled / failed.
    """
    if not os.path.exists(TEMPLATE_PATH):
        try:
            from make_template import build_template
            build_template(TEMPLATE_PATH)
        except Exception as e:
            QMessageBox.critical(parent, "Template Missing",
                                 f"Could not create the Excel template:\n{e}")
            return None

    os.makedirs(PACKING_LIST_DIR, exist_ok=True)
    default_name = f"Packing List {datetime.date.today():%Y-%m-%d}.xlsx"
    path, _ = QFileDialog.getSaveFileName(
        parent, "Save New Packing List As",
        os.path.join(PACKING_LIST_DIR, default_name), "Excel Files (*.xlsx)")
    if not path:
        return None
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"

    try:
        shutil.copyfile(TEMPLATE_PATH, path)
        if sys.platform == "win32":
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except Exception as e:
        QMessageBox.critical(parent, "Open Template Error", str(e))
        return None

    QMessageBox.information(
        parent, "Template Opened",
        "Your new packing list is open in Excel.\n\n"
        "Replace the sample rows with your pieces (keep row 1 headers),\n"
        "save and close the file, then click \"Import from Excel…\".")
    return path


# ── Excel Import Dialog ───────────────────────────────────────────────────────
class ImportDialog(QDialog):
    """Preview packing list from Excel, check/uncheck rows to include/exclude."""

    def __init__(self, parent=None, initial_path: str = ""):
        super().__init__(parent)
        self.setWindowTitle("Import Packing List from Excel")
        self.setMinimumSize(820, 520)
        self.pieces: list[FreightPiece] = []
        self.excluded_ids: list[int] = []
        self._raw_rows = []
        self.template_path: Optional[str] = None

        layout = QVBoxLayout(self)

        # File picker
        file_row = QHBoxLayout()
        self.path_edit = QLineEdit()
        self.path_edit.setPlaceholderText("Select Excel packing list (.xlsx)…")
        self.path_edit.setReadOnly(True)
        btn_browse = QPushButton("Browse…")
        btn_browse.clicked.connect(self._browse)
        btn_reload = QPushButton("Reload")
        btn_reload.setToolTip("Re-read the file after saving changes in Excel")
        btn_reload.clicked.connect(lambda: self.path_edit.text() and self._load(self.path_edit.text()))
        btn_template = QPushButton("Open Template…")
        btn_template.setToolTip("Create a new packing list from the Excel template and open it")
        btn_template.clicked.connect(self._open_template)
        file_row.addWidget(self.path_edit)
        file_row.addWidget(btn_browse)
        file_row.addWidget(btn_reload)
        file_row.addWidget(btn_template)
        layout.addLayout(file_row)

        lbl = QLabel("Check rows to include. Uncheck to exclude from the load plan.")
        lbl.setStyleSheet("color:#555; font-size:11px;")
        layout.addWidget(lbl)

        # Preview table
        self.table = QTableWidget(0, 11)
        self.table.setHorizontalHeaderLabels(
            ["✓ Include", "Crate #", "Part Number", "L\"", "W\"", "H\"", "Weight (lbs)", "PCS", "Notes",
             "Stack?", "Stop"])
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(8, QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

        btns = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        btns.accepted.connect(self._accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

        if initial_path and os.path.exists(initial_path):
            self.path_edit.setText(initial_path)
            self._load(initial_path)

    def _browse(self):
        start_dir = PACKING_LIST_DIR if os.path.isdir(PACKING_LIST_DIR) else ""
        path, _ = QFileDialog.getOpenFileName(
            self, "Open Packing List", start_dir, "Excel Files (*.xlsx *.xls)")
        if path:
            self.path_edit.setText(path)
            self._load(path)

    def _open_template(self):
        path = open_excel_template(self)
        if path:
            self.template_path = path
            self.path_edit.setText(path)
            self.table.setRowCount(0)
            self._raw_rows = []

    def _load(self, path: str):
        if not HAS_OPENPYXL:
            QMessageBox.critical(self, "Missing Library",
                                 "openpyxl is required for Excel import.\n"
                                 "Run: pip3 install openpyxl")
            return
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            self.table.setRowCount(0)
            self._raw_rows = []

            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                try:
                    crate_id = int(row[0])
                except (TypeError, ValueError):
                    continue

                part      = str(row[1] or "")
                l_in      = float(row[3] or 0)
                w_in      = float(row[4] or 0)
                h_in      = float(row[5] or 0)
                weight    = float(row[6] or 0)
                pcs       = int(row[7] or 0)
                notes     = str(row[9] or "")
                raw_stk   = row[10] if len(row) > 10 else None
                stackable = str(raw_stk).strip().upper() not in ("N", "NO", "0", "FALSE") if raw_stk is not None else True
                try:
                    stop = int(row[11]) if len(row) > 11 and row[11] not in (None, "") else 1
                except (TypeError, ValueError):
                    stop = 1
                stop = max(1, min(MAX_STOPS, stop))

                self._raw_rows.append((crate_id, part, l_in, w_in, h_in, weight, pcs, notes,
                                       stackable, stop))

                r = self.table.rowCount()
                self.table.insertRow(r)

                chk = QCheckBox()
                chk.setChecked(True)
                chk_widget = QWidget()
                chk_layout = QHBoxLayout(chk_widget)
                chk_layout.addWidget(chk)
                chk_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
                chk_layout.setContentsMargins(0, 0, 0, 0)
                self.table.setCellWidget(r, 0, chk_widget)

                for col, val in enumerate([crate_id, part, l_in, w_in, h_in,
                                           f"{weight:,.0f}", pcs, notes,
                                           "Yes" if stackable else "No", stop], start=1):
                    item = QTableWidgetItem(str(val))
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                    self.table.setItem(r, col, item)

        except Exception as e:
            QMessageBox.critical(self, "Import Error", str(e))

    def _accept(self):
        self.pieces = []
        self.excluded_ids = []

        for r, raw in enumerate(self._raw_rows):
            crate_id, part, l_in, w_in, h_in, weight, pcs, notes, stackable, stop = raw
            chk_widget = self.table.cellWidget(r, 0)
            chk = chk_widget.findChild(QCheckBox)
            if chk and chk.isChecked():
                self.pieces.append(FreightPiece(
                    piece_id=crate_id,
                    crate_label=f"Crate {crate_id}",
                    part_number=part,
                    length_in=l_in,
                    width_in=w_in,
                    height_in=h_in,
                    weight_lbs=weight,
                    pcs=pcs,
                    notes=notes,
                    stackable=stackable,
                    stop=stop,
                ))
            else:
                self.excluded_ids.append(crate_id)

        self.accept()


# ── Freight Pop-Out Window ────────────────────────────────────────────────────
class FreightPopout(QDialog):
    """Resizable, non-modal window showing all freight pieces with include/exclude toggles."""

    def __init__(self, pieces, excluded_ids, on_change, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Freight Pieces — Full View")
        self.setWindowFlags(Qt.WindowType.Window |
                            Qt.WindowType.WindowMinMaxButtonsHint |
                            Qt.WindowType.WindowCloseButtonHint)
        self.resize(980, 520)
        self._on_change = on_change
        self._blocking  = False

        layout = QVBoxLayout(self)

        lbl = QLabel("Check = include in load plan.  Uncheck = exclude.  Changes apply immediately.")
        lbl.setStyleSheet("color:#555; font-size:11px; padding:2px 0;")
        layout.addWidget(lbl)

        self.table = QTableWidget(0, 11)
        self.table.setHorizontalHeaderLabels(
            ["✓ Include", "Crate #", "Label", "Part #", 'L"', 'W"', 'H"',
             "Weight (lbs)", "PCS", "Stop", "Notes"])
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(10, QHeaderView.ResizeMode.Stretch)
        for col, w in [(0, 80), (1, 62), (4, 46), (5, 46), (6, 46), (7, 100), (8, 50), (9, 150)]:
            self.table.setColumnWidth(col, w)
        layout.addWidget(self.table)

        btn_row = QHBoxLayout()
        btn_all  = QPushButton("Select All")
        btn_none = QPushButton("Select None")
        btn_all.clicked.connect(lambda: self._set_all(True))
        btn_none.clicked.connect(lambda: self._set_all(False))
        btn_row.addWidget(btn_all)
        btn_row.addWidget(btn_none)
        btn_row.addStretch()
        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.close)
        btn_row.addWidget(btn_close)
        layout.addLayout(btn_row)

        self.refresh(pieces, excluded_ids)

    def refresh(self, pieces, excluded_ids):
        self._blocking = True
        self.table.setRowCount(0)
        for p in pieces:
            r = self.table.rowCount()
            self.table.insertRow(r)

            chk = QCheckBox()
            chk.setChecked(p.piece_id not in excluded_ids)
            chk.stateChanged.connect(self._emit_change)
            chk_widget = QWidget()
            chk_lay = QHBoxLayout(chk_widget)
            chk_lay.addWidget(chk)
            chk_lay.setAlignment(Qt.AlignmentFlag.AlignCenter)
            chk_lay.setContentsMargins(0, 0, 0, 0)
            self.table.setCellWidget(r, 0, chk_widget)

            stop_text = getattr(self.parent(), "_stop_text", lambda sid: f"S{sid}")(p.stop)
            for col, val in enumerate([
                p.piece_id, p.crate_label, p.part_number,
                f"{p.length_in:.0f}", f"{p.width_in:.0f}", f"{p.height_in:.0f}",
                f"{p.weight_lbs:,.0f}", str(p.pcs), stop_text, p.notes or ""
            ], start=1):
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if col == 9:
                    item.setForeground(QColor(STOP_COLORS[(p.stop - 1) % len(STOP_COLORS)]))
                self.table.setItem(r, col, item)

        self.table.resizeRowsToContents()
        self._blocking = False

    def _emit_change(self):
        if not self._blocking:
            self._on_change(self._get_excluded())

    def _get_excluded(self):
        excluded = []
        for r in range(self.table.rowCount()):
            chk_w = self.table.cellWidget(r, 0)
            chk = chk_w.findChild(QCheckBox)
            if chk and not chk.isChecked():
                item = self.table.item(r, 1)
                if item:
                    try:
                        excluded.append(int(item.text()))
                    except ValueError:
                        pass
        return excluded

    def _set_all(self, checked: bool):
        self._blocking = True
        for r in range(self.table.rowCount()):
            chk_w = self.table.cellWidget(r, 0)
            chk = chk_w.findChild(QCheckBox)
            if chk:
                chk.setChecked(checked)
        self._blocking = False
        self._emit_change()


# ── Main Window ───────────────────────────────────────────────────────────────
# ── Truck & Trailer settings ─────────────────────────────────────────────────
RIG_FIELDS = [   # (key, label, unit, max, step, tooltip)
    ("steer_tare", "Tractor steer axle, empty", "lb", 20000, 100,
     "Steer axle weight of the tractor alone (bobtail scale ticket)"),
    ("drive_tare", "Tractor drive axles, empty", "lb", 30000, 100,
     "Drive tandem weight of the tractor alone"),
    ("wheelbase_in", "Tractor wheelbase", "in", 400, 1,
     "Steer axle to the centre of the drive tandem"),
    ("fw_min_in", "5th wheel slide: most rearward", "in", 60, 1,
     "Fifth wheel position ahead of the drive tandem centre"),
    ("fw_max_in", "5th wheel slide: most forward", "in", 60, 1, ""),
    ("fw_step_in", "5th wheel notch spacing", "in", 12, 1, ""),
    ("fw_home_in", "5th wheel usual position", "in", 60, 1,
     "Where the planner leaves it unless moving it keeps an axle legal"),
    ("trailer_tare", "Trailer empty weight", "lb", 30000, 100, ""),
    ("slider_tare", "  of which sliding tandem assembly", "lb", 10000, 100,
     "Axles, wheels, brakes and slider box — this weight moves with the tandem"),
    ("trailer_cg_in", "Empty trailer body balance point", "in", 636, 1,
     "Centre of gravity of the empty body (without the slider), measured from the nose"),
    ("kingpin_in", "Kingpin behind the nose", "in", 120, 1, ""),
    ("tandem_min_in", "Tandem slide: most forward", "in", 636, 1,
     "Trailer tandem centre behind the kingpin"),
    ("tandem_max_in", "Tandem slide: most rearward", "in", 636, 1, ""),
    ("tandem_step_in", "Tandem hole spacing", "in", 24, 1, ""),
    ("tandem_spread_in", "Trailer axle spread", "in", 120, 1, "Front trailer axle to rear trailer axle"),
    ("kpra_max_in", "Kingpin to rear axle limit", "in", 720, 1,
     "Legal maximum, kingpin to the centre of the rear trailer axle: 480\" (40') in "
     "California, 492\" (41') in many states. 0 = no limit."),
]


def load_rig_settings():
    """Truck & Trailer settings saved by the dialog (RIG_DEFAULTS if none)."""
    try:
        saved = json.loads(QSettings("FreightLoader", "FreightLoader").value("rig", "") or "{}")
    except ValueError:
        saved = {}
    RIG.clear()
    RIG.update(RIG_DEFAULTS)
    RIG.update({k: v for k, v in saved.items() if k in RIG_DEFAULTS})


class RigDialog(QDialog):
    """Edit the tractor / trailer weights, dimensions and slide ranges in RIG."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Truck & Trailer")
        form = QFormLayout(self)
        intro = QLabel("Empty weights and slide ranges used for axle weights. Defaults are "
                       "typical for a sleeper tractor and a 53' dry van — enter your own scale "
                       "tickets and measurements for better estimates.")
        intro.setWordWrap(True)
        intro.setStyleSheet("color:#555;")
        form.addRow(intro)
        self._spins = {}
        for key, label, unit, top, step, tip in RIG_FIELDS:
            spin = QSpinBox()
            spin.setRange(0, top)
            spin.setSingleStep(step)
            spin.setSuffix(f" {unit}")
            spin.setGroupSeparatorShown(True)
            spin.setValue(int(RIG[key]))
            spin.setToolTip(tip)
            self._spins[key] = spin
            form.addRow(label + ":", spin)
        self._lbl_payload = QLabel()
        form.addRow("Legal payload at gross limit:", self._lbl_payload)
        for key in ("steer_tare", "drive_tare", "trailer_tare"):
            self._spins[key].valueChanged.connect(self._update_payload)
        self._update_payload()
        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok |
                                QDialogButtonBox.StandardButton.Cancel |
                                QDialogButtonBox.StandardButton.RestoreDefaults)
        btns.accepted.connect(self._accept)
        btns.rejected.connect(self.reject)
        btns.button(QDialogButtonBox.StandardButton.RestoreDefaults).clicked.connect(
            lambda: [self._spins[k].setValue(RIG_DEFAULTS[k]) for k in self._spins])
        form.addRow(btns)

    def _update_payload(self):
        tare = sum(self._spins[k].value() for k in ("steer_tare", "drive_tare", "trailer_tare"))
        self._lbl_payload.setText(f"{TRAILER['gross_limit'] - tare:,} lb  (empty rig {tare:,} lb)")

    def _accept(self):
        v = {k: sp.value() for k, sp in self._spins.items()}
        problems = []
        if v["fw_min_in"] > v["fw_max_in"] or v["tandem_min_in"] > v["tandem_max_in"]:
            problems.append("A slide range's first value is past its second.")
        if not v["fw_step_in"] or not v["tandem_step_in"]:
            problems.append("Notch and hole spacing must be at least 1\".")
        if v["slider_tare"] > v["trailer_tare"]:
            problems.append("The sliding tandem can't weigh more than the whole trailer.")
        if v["fw_max_in"] >= v["wheelbase_in"]:
            problems.append("The fifth wheel must sit behind the steer axle.")
        if not v["tandem_min_in"]:
            problems.append("The trailer tandem must sit behind the kingpin.")
        if problems:
            QMessageBox.warning(self, "Truck & Trailer", "\n".join(problems))
            return
        RIG.update(v)
        QSettings("FreightLoader", "FreightLoader").setValue("rig", json.dumps(RIG))
        self.accept()


class FreightLoaderApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"Freight Loader v{APP_VERSION} — 53' Dry Van")
        self.setMinimumSize(1150, 760)
        self._next_id = 1
        self._plan: Optional[LoadPlan] = None
        self._excluded_ids: list = []
        self._securement: list = []     # load bars / straps / airbags, shared with 3D view
        # Shipment: stops with locations, trucks, pieces shipping separately
        self._origin = {"name": "", "location": "", "lat": None, "lon": None}
        self._stops = [{"name": "", "location": "", "lat": None, "lon": None}
                       for _ in range(MAX_STOPS)]
        self._use_road = False
        self._road = None
        self._trucks: list = []          # [{"trailer_num": str, "piece_ids": [..]}]
        self._ship_other: list = []      # piece IDs shipping another way
        self._ship_reasons: dict = {}    # piece_id -> why it ships separately
        self._split_key = None           # {piece_id: stop} when trucks were last split
        self._plans: list = []           # LoadPlan per truck
        self._manual: dict = {}          # piece_id -> (x, y) for floor crates moved in the Top View
        self._cur_truck = 0
        self._current_file: Optional[str] = None
        self._popout: Optional[FreightPopout] = None
        self._build_ui()
        self._sync_rig_controls(None)

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setSpacing(2)

        lbl_version = QLabel(f"Freight Loader  v{APP_VERSION}")
        lbl_version.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_version.setStyleSheet("color:#2a7d6e; font-weight:bold; font-size:12px;")
        root.addWidget(lbl_version)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        root.addWidget(splitter)

        # ══ LEFT PANEL ══════════════════════════════════════════════════
        left = QWidget()
        ll = QVBoxLayout(left)
        ll.setSpacing(6)

        # ── File actions ───────────────────────────────────────────────
        file_row = QHBoxLayout()
        btn_new   = QPushButton("New Plan")
        btn_open  = QPushButton("Open Plan…")
        btn_save  = QPushButton("Save Plan")
        btn_saveas = QPushButton("Save As…")
        for b in [btn_new, btn_open, btn_save, btn_saveas]:
            b.setFixedHeight(30)
            file_row.addWidget(b)
        btn_new.clicked.connect(self._new_plan)
        btn_open.clicked.connect(self._open_plan)
        btn_save.clicked.connect(self._save_plan)
        btn_saveas.clicked.connect(self._save_plan_as)
        ll.addLayout(file_row)

        # ── Shipment Info ──────────────────────────────────────────────
        grp_info = QGroupBox("Shipment Info")
        info_form = QFormLayout(grp_info)
        info_form.setRowWrapPolicy(QFormLayout.RowWrapPolicy.DontWrapRows)

        self.f_customer      = QLineEdit(); self.f_customer.setPlaceholderText("Customer name")
        self.f_shipment_name = QLineEdit(); self.f_shipment_name.setPlaceholderText("Shipment / load name")
        self.f_trailer_num   = QLineEdit(); self.f_trailer_num.setPlaceholderText("Optional (this truck)")
        self.f_trailer_num.textEdited.connect(self._trailer_edited)
        self.f_bol           = QLineEdit(); self.f_bol.setPlaceholderText("Optional")
        self.f_date          = QLineEdit(datetime.date.today().strftime("%Y-%m-%d"))

        info_form.addRow("Customer:",      self.f_customer)
        info_form.addRow("Shipment Name:", self.f_shipment_name)
        info_form.addRow("Trailer #:",     self.f_trailer_num)
        info_form.addRow("BOL #:",         self.f_bol)
        info_form.addRow("Date:",          self.f_date)
        self._stop_names = [""] * MAX_STOPS
        stops_row = QHBoxLayout()
        btn_stops = QPushButton("Edit Stops…")
        btn_stops.setToolTip("Names and addresses / GPS coordinates for the origin and stops; "
                             "plans the delivery route")
        btn_stops.clicked.connect(self._edit_stops)
        self.lbl_stops = QLabel("")
        self.lbl_stops.setStyleSheet("color:#555;")
        stops_row.addWidget(btn_stops)
        stops_row.addWidget(self.lbl_stops, 1)
        info_form.addRow("Delivery Stops:", stops_row)
        ll.addWidget(grp_info)

        # ── Freight table ──────────────────────────────────────────────
        grp_freight = QGroupBox("Freight Pieces")
        gl = QVBoxLayout(grp_freight)

        self.table = QTableWidget(0, 12)
        self.table.setHorizontalHeaderLabels(
            ["Crate #", "Label", "Part #", 'L"', 'W"', 'H"', "Weight (lbs)", "PCS", "Pos.", "Stack?",
             "Notes", "Stop"])
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        for col, w in [(0,58),(3,46),(4,46),(5,46),(6,88),(7,50),(8,72),(9,56)]:
            self.table.setColumnWidth(col, w)
        self.table.setColumnHidden(10, True)
        self.table.setColumnWidth(11, 52)
        self.table.horizontalHeaderItem(11).setToolTip(
            "Delivery stop ID — set names and addresses / GPS under Edit Stops. "
            "Delivery order follows the route (by stop # if no locations).")
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        gl.addWidget(self.table)

        btn_row = QHBoxLayout()
        btn_template = QPushButton("Open Excel Template…")
        btn_template.setToolTip("Create a new packing list from the Excel template, fill it in, then import")
        btn_template.clicked.connect(self._open_template)
        btn_import = QPushButton("Import from Excel…")
        btn_import.clicked.connect(self._import_excel)
        btn_add    = QPushButton("+ Add Row")
        btn_add.clicked.connect(self._add_row)
        btn_del    = QPushButton("Remove Selected")
        btn_del.clicked.connect(self._remove_row)
        btn_clear  = QPushButton("Clear All")
        btn_clear.clicked.connect(self._clear_table)
        btn_popout = QPushButton("↗ Pop Out")
        btn_popout.setToolTip("Open freight list in a larger window to manage include/exclude")
        btn_popout.clicked.connect(self._open_popout)
        for b in [btn_template, btn_import, btn_add, btn_del, btn_clear, btn_popout]:
            btn_row.addWidget(b)
        gl.addLayout(btn_row)
        ll.addWidget(grp_freight)

        # ── Axle limits ────────────────────────────────────────────────
        grp_limits = QGroupBox("DOT Axle Weight Limits (lbs)")
        lim_row = QHBoxLayout(grp_limits)
        self._limit_fields = {}
        for key, label, default in [
            ("steer",   "Steer",   12000),
            ("drive",   "Drive",   34000),
            ("trailer", "Trailer", 34000),
            ("gross",   "Gross",   80000),
        ]:
            vb = QVBoxLayout()
            vb.addWidget(QLabel(label, alignment=Qt.AlignmentFlag.AlignCenter))
            spin = QSpinBox()
            spin.setRange(1000, 200000)
            spin.setValue(default)
            spin.setSingleStep(500)
            spin.setGroupSeparatorShown(True)
            self._limit_fields[key] = spin
            vb.addWidget(spin)
            lim_row.addLayout(vb)
        self.chk_override = QCheckBox("Allow override")
        lim_row.addWidget(self.chk_override)
        ll.addWidget(grp_limits)

        # ── Calculate ─────────────────────────────────────────────────
        btn_calc = QPushButton("⚡  Calculate Load Plan")
        btn_calc.setFixedHeight(44)
        font = btn_calc.font(); font.setPointSize(13); font.setBold(True)
        btn_calc.setFont(font)
        btn_calc.clicked.connect(self._calculate)
        ll.addWidget(btn_calc)

        # ── Summary / violations ───────────────────────────────────────
        grp_sum = QGroupBox("Axle Weight Summary")
        sl = QVBoxLayout(grp_sum)
        sl.setContentsMargins(4, 4, 4, 4)
        sum_scroll = QScrollArea()
        sum_scroll.setWidgetResizable(True)
        sum_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        sum_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        sum_scroll.setFixedHeight(110)
        self.lbl_summary = QLabel("—")
        self.lbl_summary.setTextFormat(Qt.TextFormat.RichText)
        self.lbl_summary.setWordWrap(True)
        self.lbl_summary.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignLeft)
        self.lbl_summary.setContentsMargins(4, 2, 4, 2)
        sum_scroll.setWidget(self.lbl_summary)
        sl.addWidget(sum_scroll)
        ll.addWidget(grp_sum)

        grp_status = QGroupBox("Status / Violations")
        stl = QVBoxLayout(grp_status)
        self.lbl_status = QLabel("No plan calculated.")
        self.lbl_status.setTextFormat(Qt.TextFormat.RichText)
        self.lbl_status.setWordWrap(True)
        stl.addWidget(self.lbl_status)
        ll.addWidget(grp_status)

        # ── Export buttons ─────────────────────────────────────────────
        exp_row = QHBoxLayout()
        self.btn_plan_pdf  = QPushButton("Export Floor Plan PDF…")
        self.btn_label_pdf = QPushButton("Export Crate Labels PDF…")
        self.btn_plan_pdf.setToolTip("Preview the floor plan PDF, then save it")
        self.btn_label_pdf.setToolTip("Preview the crate labels, then save them")
        for b in [self.btn_plan_pdf, self.btn_label_pdf]:
            b.setEnabled(False)
            exp_row.addWidget(b)
        self.btn_plan_pdf.clicked.connect(self._export_floor_plan)
        self.btn_label_pdf.clicked.connect(self._export_labels)
        ll.addLayout(exp_row)
        self.chk_pdf_3d = QCheckBox("Include 3D views page in floor plan PDF")
        self.chk_pdf_3d.setToolTip("Adds a landscape page with the load shown from four 3D angles")
        ll.addWidget(self.chk_pdf_3d)

        splitter.addWidget(left)

        # ══ RIGHT PANEL — truck bar + tabbed trailer views ═══════════
        right = QWidget()
        rl = QVBoxLayout(right)
        rl.setContentsMargins(0, 0, 0, 0)
        truck_row = QHBoxLayout()
        truck_row.addWidget(QLabel("<b>Truck:</b>"))
        self.cmb_truck = QComboBox()
        self.cmb_truck.setMinimumWidth(320)
        self.cmb_truck.currentIndexChanged.connect(self._select_truck)
        truck_row.addWidget(self.cmb_truck)
        btn_resplit = QPushButton("Re-split Trucks")
        btn_resplit.setToolTip("Re-plan which stops go on which truck (asks when a stop won't fit)")
        btn_resplit.clicked.connect(lambda: self._calculate(resplit=True))
        btn_move = QPushButton("Move Stop…")
        btn_move.setToolTip("Move all of a stop's freight to another truck, a new truck, "
                            "or ship it separately")
        btn_move.clicked.connect(self._move_stop)
        self.btn_other = QPushButton("Ship Separately (0)")
        self.btn_other.clicked.connect(self._show_ship_other)
        for b in (btn_resplit, btn_move, self.btn_other):
            truck_row.addWidget(b)
        truck_row.addStretch()
        rl.addLayout(truck_row)

        rig_row = QHBoxLayout()
        rig_row.addWidget(QLabel("<b>Axles:</b>"))
        self.spin_fw = QSpinBox()
        self.spin_fw.setSuffix("\"")
        self.spin_fw.setToolTip("Fifth wheel position ahead of the drive tandem centre — "
                                "forward puts more weight on the steer axle")
        self.spin_tandem = QSpinBox()
        self.spin_tandem.setSuffix("\"")
        self.spin_tandem.setToolTip("Trailer tandem centre behind the kingpin — "
                                    "back puts more weight on the drive axles")
        for text, spin in (("5th wheel", self.spin_fw), ("Tandem", self.spin_tandem)):
            lbl = QLabel(text)
            lbl.setToolTip(spin.toolTip())
            rig_row.addWidget(lbl)
            rig_row.addWidget(spin)
        self.spin_fw.valueChanged.connect(self._rig_spin_changed)
        self.spin_tandem.valueChanged.connect(self._rig_spin_changed)
        self.btn_rig_auto = QPushButton("Auto")
        self.btn_rig_auto.setToolTip("Let the planner place the fifth wheel and tandem for this truck")
        self.btn_rig_auto.clicked.connect(lambda: self._set_rig(None))
        rig_row.addWidget(self.btn_rig_auto)
        self.lbl_rig = QLabel()
        self.lbl_rig.setStyleSheet("color:#555;")
        rig_row.addWidget(self.lbl_rig)
        rig_row.addStretch()
        btn_rig = QPushButton("Truck && Trailer…")
        btn_rig.setToolTip("Empty weights, wheelbase, slide ranges and the kingpin-to-rear-axle limit")
        btn_rig.clicked.connect(self._edit_rig)
        rig_row.addWidget(btn_rig)
        rl.addLayout(rig_row)

        tabs = QTabWidget()
        tabs.setDocumentMode(True)

        # Tab 1 — Top view (zoomable)
        self.trailer_view = TrailerView()
        self.trailer_view.plan_changed.connect(self._on_plan_changed_by_drag)
        btn_reset_moves = QPushButton("Reset Moves")
        btn_reset_moves.setToolTip("Put crates you dragged on this truck back where the planner placed them")
        btn_reset_moves.clicked.connect(self._reset_moves)
        tabs.addTab(self._zoom_tab(self.trailer_view, [btn_reset_moves],
                                   "Drag crate: shift  ·  Ctrl+drag: swap  ·  "), "Top View")

        # Tab 2 — Side elevation (zoomable)
        self.side_view = SideElevationView()
        self.side_view.rig_moved.connect(self._set_rig)
        tabs.addTab(self._zoom_tab(self.side_view, hint_prefix="Drag axles: slide  ·  "),
                    "Side View")

        # Tab 3 — Rotatable 3D view
        view3d_tab = QWidget()
        v3l = QVBoxLayout(view3d_tab)
        v3l.setContentsMargins(4, 4, 4, 4)
        preset_row = QHBoxLayout()
        self.view_3d = Trailer3DView()
        for name in Trailer3DView.PRESETS:
            b = QPushButton(name)
            b.clicked.connect(lambda _=False, n=name: self.view_3d.set_preset(n))
            preset_row.addWidget(b)
        btn_reset3d = QPushButton("Reset")
        btn_reset3d.clicked.connect(self.view_3d.reset_view)
        preset_row.addWidget(btn_reset3d)
        preset_row.addStretch()
        v3l.addLayout(preset_row)

        tool_row = QHBoxLayout()
        tool_row.addWidget(QLabel("Securement:"))
        self._tool_group = QButtonGroup(self)
        for tool, text in [("rotate", "Rotate"), ("load_bar", "+ Load Bar"),
                           ("strap", "+ Strap"), ("airbag", "+ Airbag"), ("remove", "Remove")]:
            b = QPushButton(text)
            b.setCheckable(True)
            b.setChecked(tool == "rotate")
            b.clicked.connect(lambda _=False, t=tool: self.view_3d.set_tool(t))
            self._tool_group.addButton(b)
            tool_row.addWidget(b)
        btn_clear_sec = QPushButton("Clear All")
        btn_clear_sec.setToolTip("Remove all load bars, straps and airbags")
        btn_clear_sec.clicked.connect(self._clear_securement)
        tool_row.addWidget(btn_clear_sec)
        self.btn_remove_hidden = QPushButton("Remove Hidden")
        self.btn_remove_hidden.setToolTip("Delete load bars / straps / airbags that no longer fit "
                                          "this truck's layout (they aren't shown)")
        self.btn_remove_hidden.setStyleSheet("color:#c0392b;")
        self.btn_remove_hidden.clicked.connect(lambda: self.view_3d.remove_unfit())
        self.btn_remove_hidden.setVisible(False)
        tool_row.addWidget(self.btn_remove_hidden)
        tool_row.addStretch()
        self.lbl_securement = QLabel()
        self.lbl_securement.setStyleSheet("color:#555;")
        tool_row.addWidget(self.lbl_securement)
        v3l.addLayout(tool_row)

        self.view_3d.set_securement(self._securement)
        self.trailer_view.set_securement(self._securement)
        self.side_view.set_securement(self._securement)
        self.view_3d.securement_changed.connect(self._on_securement_changed)
        v3l.addWidget(self.view_3d)
        tabs.addTab(view3d_tab, "3D View")

        rl.addWidget(tabs)
        splitter.addWidget(right)
        splitter.setSizes([500, 610])

    def _zoom_tab(self, view: QWidget, extra: list = None, hint_prefix: str = "") -> QWidget:
        """Wrap a 2D view in a zoomable scroll area with − / Fit / + controls."""
        tab = QWidget()
        lay = QVBoxLayout(tab)
        lay.setContentsMargins(4, 4, 4, 4)
        area = ZoomScrollArea(view)
        row = QHBoxLayout()
        lbl = QLabel("100%")
        lbl.setMinimumWidth(44)
        for text, tip, fn in [("−", "Zoom out", lambda: area.set_zoom(area.zoom / 1.25)),
                              ("Fit", "Fit to window", lambda: area.set_zoom(1.0)),
                              ("+", "Zoom in", lambda: area.set_zoom(area.zoom * 1.25))]:
            b = QPushButton(text)
            b.setToolTip(tip)
            b.setFixedWidth(44)
            b.clicked.connect(fn)
            row.addWidget(b)
        row.addWidget(lbl)
        for wdg in extra or []:
            row.addWidget(wdg)
        hint = QLabel(hint_prefix + "Ctrl+wheel: zoom  ·  Right-drag: pan")
        hint.setStyleSheet("color:#777;")
        row.addWidget(hint)
        row.addStretch()
        area.zoom_changed.connect(lambda z: lbl.setText(f"{z * 100:.0f}%"))
        lay.addLayout(row)
        lay.addWidget(area)
        return tab

    def _securement_text(self) -> str:
        """This truck's securement that fits the layout, plus how many items are hidden."""
        ids = {sl.piece.piece_id for sl in self._plan.slots if sl.piece} if self._plan else set()
        hidden = set(self.view_3d.unfit_indices()) if self._plan else set()
        shown = [it for i, it in enumerate(self._securement)
                 if it["piece_id"] in ids and i not in hidden]
        text = securement_summary(shown)
        if hidden:
            text += f"{' · ' if text else ''}{len(hidden)} hidden"
        return text

    def _on_securement_changed(self):
        self.lbl_securement.setText(self._securement_text())
        n_hidden = len(self.view_3d.unfit_indices())
        self.btn_remove_hidden.setText(f"Remove Hidden ({n_hidden})")
        self.btn_remove_hidden.setVisible(n_hidden > 0)
        self.trailer_view.update()
        self.side_view.update()
        if self._plan:
            self._update_summary(self._plan)

    def _clear_securement(self):
        if self._securement and QMessageBox.question(
                self, "Clear Securement", "Remove all load bars, straps and airbags?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                ) != QMessageBox.StandardButton.Yes:
            return
        self._securement.clear()
        self.view_3d.update()
        self._on_securement_changed()

    def _on_plan_changed_by_drag(self):
        if self._plan:
            for sl in self._plan.slots:
                if sl.tier == 0 and sl.piece and not (sl.is_center and sl.col == 1):
                    if sl.x_in is not None:
                        self._manual[sl.piece.piece_id] = (sl.x_in, sl.y_in)
                    else:
                        self._manual.pop(sl.piece.piece_id, None)
            self._show_plan_change()
            self._on_securement_changed()   # moved crates can open or close gaps

    def _reset_moves(self):
        if not self._plan:
            return
        ids = {sl.piece.piece_id for sl in self._plan.slots if sl.piece}
        if not any(i in self._manual for i in ids):
            return
        for i in ids:
            self._manual.pop(i, None)
        self._build_plans(self._read_pieces(), silent=True)

    def _apply_manual_moves(self, plan: LoadPlan, truck_no: int):
        """Re-apply Top View drags to a freshly built plan; reset them if they no longer fit."""
        moved = [sl for sl in plan.slots if sl.tier == 0 and sl.piece
                 and not (sl.is_center and sl.col == 1) and sl.piece.piece_id in self._manual]
        if not moved:
            return
        for sl in moved:
            sl.x_in, sl.y_in = self._manual[sl.piece.piece_id]
        if _layout_problems(plan.slots):
            for sl in moved:
                sl.x_in = sl.y_in = None
                self._manual.pop(sl.piece.piece_id, None)
            plan.notices.append(f"Crates you moved on Truck {truck_no} were put back — "
                                "the load changed and they no longer fit where you placed them.")
        _refresh_plan_metrics(plan)

    # ── Fifth wheel / tandem positions ───────────────────────────────────────
    def _sync_rig_controls(self, plan: Optional[LoadPlan]):
        """Show `plan`'s positions in the Axles row without triggering a change."""
        fw_opts, holes = rig_positions(legal_only=False)
        rig = plan.rig if plan and plan.rig else default_rig()
        for spin, opts, val in ((self.spin_fw, fw_opts, rig["fw_in"]),
                                (self.spin_tandem, holes, rig["tandem_in"])):
            spin.blockSignals(True)
            spin.setRange(int(min(opts)), int(max(opts)))
            spin.setSingleStep(int(opts[1] - opts[0]) if len(opts) > 1 else 1)
            spin.setValue(int(val))
            spin.blockSignals(False)
            spin.setEnabled(plan is not None)
        auto = plan is None or plan.rig_auto
        self.btn_rig_auto.setEnabled(not auto)
        rear = rig["tandem_in"] + RIG["tandem_spread_in"] / 2
        bad = not _rear_axle_ok(rig["tandem_in"])
        self.lbl_rig.setText(
            f"<span style='color:{'#c0392b' if bad else '#555'}'>KP→rear axle {int(rear // 12)}'"
            f"{rear % 12:.0f}\"</span>  ·  " + ("auto" if auto else "<b>set by you</b>"))
        self.lbl_rig.setToolTip(rig_text(rig))

    def _rig_spin_changed(self):
        fw_opts, holes = rig_positions(legal_only=False)
        self._set_rig({"fw_in": min(fw_opts, key=lambda f: abs(f - self.spin_fw.value())),
                       "tandem_in": min(holes, key=lambda t: abs(t - self.spin_tandem.value()))})

    def _set_rig(self, rig: Optional[dict]):
        """Fix this truck's fifth wheel / tandem at `rig`, or None to let the planner place them."""
        if not self._plan or not self._trucks:
            return
        truck = self._trucks[self._cur_truck]
        if rig is None:
            truck.pop("rig", None)
        else:
            truck["rig"] = dict(rig)
            self._plan.rig = dict(rig)
        self._plan.rig_auto = rig is None
        _refresh_plan_metrics(self._plan)
        self._show_plan_change()

    def _show_plan_change(self):
        """Redraw after the current plan's metrics changed (axles moved, crates dragged)."""
        self._sync_rig_controls(self._plan)
        self._update_summary(self._plan)
        self.trailer_view.update()
        self.side_view.update()
        self.view_3d.update()

    def _edit_rig(self):
        if RigDialog(self).exec() != QDialog.DialogCode.Accepted:
            return
        for t in self._trucks:
            t.pop("rig", None)    # old positions may not exist on the new slide ranges
        if self.table.rowCount():
            self._calculate(silent=True)
        else:
            self._sync_rig_controls(None)

    def _open_popout(self):
        pieces = self._read_pieces()
        if not pieces:
            return
        if self._popout and self._popout.isVisible():
            self._popout.refresh(pieces, self._excluded_ids)
            self._popout.raise_()
            self._popout.activateWindow()
            return
        self._popout = FreightPopout(pieces, self._excluded_ids,
                                     self._apply_exclusions, parent=self)
        self._popout.show()

    def _apply_exclusions(self, excluded_ids: list):
        self._excluded_ids = excluded_ids
        self._calculate(silent=True)

    # ── Actions ──────────────────────────────────────────────────────────────
    def _get_shipment_info(self) -> ShipmentInfo:
        return ShipmentInfo(
            customer=self.f_customer.text().strip(),
            shipment_name=self.f_shipment_name.text().strip(),
            trailer_num=self.f_trailer_num.text().strip(),
            bol_num=self.f_bol.text().strip(),
            date=self.f_date.text().strip(),
            stop_names=list(self._stop_names),
        )

    def _refresh_stop_label(self):
        used = [f"S{i + 1}{' ✓' if st.get('lat') is not None else ''}"
                + (f" {st['name']}" if st.get("name") else "")
                for i, st in enumerate(self._stops) if st.get("name") or st.get("location")]
        self.lbl_stops.setText("  ·  ".join(used) if used else "Set a Stop # per piece in the table")
        self._stop_names = [st.get("name", "") for st in self._stops]

    def _coords(self) -> dict:
        out = {}
        if self._origin.get("lat") is not None:
            out["origin"] = (self._origin["lat"], self._origin["lon"])
        for i, st in enumerate(self._stops, start=1):
            if st.get("lat") is not None:
                out[i] = (st["lat"], st["lon"])
        return out

    def _route(self, stop_ids) -> list:
        coords = self._coords()
        return route_order(sorted(set(stop_ids)), coords,
                           shipment_distance_fn(coords, self._use_road, self._road))

    def _stop_text(self, sid: int) -> str:
        name = self._stops[sid - 1].get("name") if 0 < sid <= MAX_STOPS else ""
        return f"S{sid}" + (f" — {name}" if name else "")

    def _edit_stops(self):
        dlg = StopsDialog(self._origin, self._stops, self._use_road, self._road, self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        moved = ([(d.get("lat"), d.get("lon")) for d in [dlg.origin] + dlg.stops] !=
                 [(d.get("lat"), d.get("lon")) for d in [self._origin] + self._stops]
                 or dlg.chk_road.isChecked() != self._use_road)
        self._origin, self._stops = dlg.origin, dlg.stops
        self._use_road, self._road = dlg.chk_road.isChecked(), dlg.road
        self._refresh_stop_label()
        if self.table.rowCount():
            # New locations change the route, so trucks are re-split
            self._calculate(resplit=moved)

    def _open_template(self):
        path = open_excel_template(self)
        if path:
            self._last_packing_list = path

    def _import_excel(self):
        dlg = ImportDialog(self, getattr(self, "_last_packing_list", ""))
        accepted = dlg.exec() == QDialog.DialogCode.Accepted
        if dlg.template_path:
            self._last_packing_list = dlg.template_path
        if accepted and dlg.pieces:
            self._load_pieces(dlg.pieces)
            self._excluded_ids = dlg.excluded_ids
            self._securement.clear()     # items were anchored to the previous crates
            self._manual.clear()         # so were hand-placed positions
            self._on_securement_changed()
            if dlg.excluded_ids:
                self.lbl_status.setText(
                    f"Imported {len(dlg.pieces)} crates. "
                    f"Excluded: {', '.join('Crate ' + str(i) for i in dlg.excluded_ids)}")

    def _load_pieces(self, pieces: list):
        self.table.setRowCount(0)
        self._next_id = 1
        half_w       = TRAILER["interior_width_in"] / 2
        alt_w_thresh = TRAILER["interior_width_in"] * 0.70
        for p in pieces:
            if p.width_in > alt_w_thresh:
                placement = "alt_center"
            elif p.width_in > half_w:
                placement = "center"
            else:
                placement = p.placement
            self._add_row(
                crate_id=p.piece_id,
                label=p.crate_label,
                part_num=p.part_number,
                length=p.length_in,
                width=p.width_in,
                height=p.height_in,
                weight=p.weight_lbs,
                pcs=p.pcs,
                placement=placement,
                stackable=p.stackable,
                notes=p.notes,
                stop=p.stop,
            )

    def _add_row(self, crate_id=None, label="", part_num="",
                 length=48.0, width=48.0, height=48.0, weight=1000.0, pcs=0,
                 placement="auto", stackable=True, notes="", stop=1):
        r = self.table.rowCount()
        self.table.insertRow(r)

        if crate_id is None:
            crate_id = self._next_id
        if not label:
            label = f"Crate {crate_id}"

        self._next_id = max(self._next_id, crate_id) + 1

        for col, val in enumerate([crate_id, label, part_num,
                                    length, width, height, weight, pcs]):
            item = QTableWidgetItem(str(val))
            item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            if col == 0:
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.table.setItem(r, col, item)

        # Col 8: Pos. placement combo box — changing it auto-recalculates
        combo = QComboBox()
        combo.addItems(["Auto", "Center", "Alt. Center"])
        if placement == "alt_center":
            combo.setCurrentIndex(2)
        elif placement == "center":
            combo.setCurrentIndex(1)
        else:
            combo.setCurrentIndex(0)
        combo.currentIndexChanged.connect(lambda: self._calculate(silent=True))
        self.table.setCellWidget(r, 8, combo)

        # Col 9: Stack? checkbox
        stk_widget = QWidget()
        stk_layout = QHBoxLayout(stk_widget)
        stk_chk = QCheckBox()
        stk_chk.setChecked(stackable)
        stk_chk.stateChanged.connect(lambda: self._calculate(silent=True))
        stk_layout.addWidget(stk_chk)
        stk_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        stk_layout.setContentsMargins(0, 0, 0, 0)
        self.table.setCellWidget(r, 9, stk_widget)

        # Col 10: Notes (hidden — preserves notes through calculate cycles)
        notes_item = QTableWidgetItem(notes)
        self.table.setItem(r, 10, notes_item)

        # Col 11: delivery stop
        stop_combo = QComboBox()
        stop_combo.addItems([str(i) for i in range(1, MAX_STOPS + 1)])
        stop_combo.setCurrentIndex(max(1, min(MAX_STOPS, int(stop))) - 1)
        stop_combo.currentIndexChanged.connect(lambda: self._calculate(silent=True))
        self.table.setCellWidget(r, 11, stop_combo)

    def _remove_row(self):
        rows = sorted(set(i.row() for i in self.table.selectedItems()), reverse=True)
        for r in rows:
            self.table.removeRow(r)

    def _clear_table(self):
        if QMessageBox.question(self, "Clear All", "Remove all freight rows?",
                                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                                ) == QMessageBox.StandardButton.Yes:
            self.table.setRowCount(0)
            self._next_id = 1

    def _read_pieces(self) -> list:
        pieces = []
        for r in range(self.table.rowCount()):
            try:
                crate_id = int(self.table.item(r, 0).text())
                label    = self.table.item(r, 1).text() if self.table.item(r, 1) else f"Crate {crate_id}"
                part     = self.table.item(r, 2).text() if self.table.item(r, 2) else ""
                l_in     = float(self.table.item(r, 3).text())
                w_in     = float(self.table.item(r, 4).text())
                h_in     = float(self.table.item(r, 5).text())
                weight   = float(self.table.item(r, 6).text())
                pcs_item = self.table.item(r, 7)
                pcs      = int(pcs_item.text()) if pcs_item and pcs_item.text().strip() else 0
                combo    = self.table.cellWidget(r, 8)
                txt = combo.currentText() if combo else "Auto"
                if txt == "Alt. Center":
                    placement = "alt_center"
                elif txt == "Center":
                    placement = "center"
                else:
                    placement = "auto"
                stk_w = self.table.cellWidget(r, 9)
                stk_chk = stk_w.findChild(QCheckBox) if stk_w else None
                stackable = stk_chk.isChecked() if stk_chk else True
                notes_item = self.table.item(r, 10)
                notes = notes_item.text() if notes_item else ""
                stop_combo = self.table.cellWidget(r, 11)
                stop = int(stop_combo.currentText()) if stop_combo else 1
                pieces.append(FreightPiece(crate_id, label, part, l_in, w_in, h_in, weight, pcs,
                                           notes=notes, placement=placement, stackable=stackable,
                                           stop=stop))
            except (ValueError, AttributeError) as e:
                QMessageBox.warning(self, "Input Error", f"Row {r+1}: {e}")
                return []
        return pieces

    def _calculate(self, silent: bool = False, resplit: bool = False):
        pieces = self._read_pieces()
        if not pieces:
            return

        TRAILER["steer_limit"]   = self._limit_fields["steer"].value()
        TRAILER["drive_limit"]   = self._limit_fields["drive"].value()
        TRAILER["trailer_limit"] = self._limit_fields["trailer"].value()
        TRAILER["gross_limit"]   = self._limit_fields["gross"].value()

        active = [pc for pc in pieces if pc.piece_id not in self._excluded_ids]
        key = {pc.piece_id: pc.stop for pc in active}
        stale = resplit or self._split_key != key or not self._trucks
        if stale and not silent:
            # The split may stop to ask about a stop; don't leave the old plan's status up
            self.lbl_status.setText("<span style='color:#555'>Splitting freight into trucks…</span>")
            QApplication.processEvents()
            self._split_trucks(active)
            self._split_key = key
        elif stale:
            self._quick_assign(active)   # no dialogs while editing; Calculate re-splits
        self._build_plans(pieces, silent)

    def _split_trucks(self, active: list):
        """Assign every active piece to a truck (or ship-separately) along the route."""
        route = self._route({pc.stop for pc in active})
        trucks, other, impossible = split_into_trucks(
            [FreightPiece(**vars(pc)) for pc in active], route, self._ask_overflow)
        old_trailers = [t.get("trailer_num", "") for t in self._trucks]
        self._trucks = [{"trailer_num": old_trailers[i] if i < len(old_trailers) else "",
                         "piece_ids": ids} for i, ids in enumerate(trucks)]
        self._ship_reasons = {pc.piece_id: why for pc, why in impossible}
        self._ship_other = [pc.piece_id for pc, _ in impossible] + other
        self._cur_truck = 0
        if impossible:
            QMessageBox.warning(
                self, "Can't Go On Any Truck",
                "These pieces can't go in a 53' van at all and are listed under "
                "Ship Separately:\n\n" +
                "\n".join(f"  •  {pc.crate_label} — {why}" for pc, why in impossible))

    def _quick_assign(self, active: list):
        """Keep existing assignments; put new/changed pieces with their stop's truck."""
        if not self._trucks:
            self._trucks = [{"trailer_num": "", "piece_ids": []}]
        ids = {pc.piece_id for pc in active}
        for t in self._trucks:
            t["piece_ids"] = [i for i in t["piece_ids"] if i in ids]
        self._ship_other = [i for i in self._ship_other if i in ids]
        stop_of = {pc.piece_id: pc.stop for pc in active}
        old = self._split_key or {}
        placed = {i for t in self._trucks for i in t["piece_ids"]} | set(self._ship_other)
        for pc in active:
            if pc.piece_id in placed and old.get(pc.piece_id) == pc.stop:
                continue
            for t in self._trucks:
                if pc.piece_id in t["piece_ids"]:
                    t["piece_ids"].remove(pc.piece_id)
            if pc.piece_id in self._ship_other:
                self._ship_other.remove(pc.piece_id)
            home = next((t for t in self._trucks
                         if any(stop_of.get(i) == pc.stop for i in t["piece_ids"])),
                        self._trucks[min(self._cur_truck, len(self._trucks) - 1)])
            home["piece_ids"].append(pc.piece_id)
        # _split_key is left alone: only a real split counts, so the next
        # Calculate still re-splits and asks about stops that don't fit

    def _ask_overflow(self, info: dict) -> str:
        """User's choice when a stop won't fully fit on the current truck."""
        sid, t, fit = info["stop"], info["truck"], info["fit"]
        rem = info["pieces"]
        rest = len(rem) - fit
        name = self._stop_text(sid)
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Icon.Question)
        box.setWindowTitle(f"{name} doesn't fit on Truck {t}")
        if fit:
            box.setText(f"<b>Truck {t}</b> has room for <b>{fit} of {len(rem)}</b> pieces for "
                        f"<b>{name}</b> ({info['fit_lbs']:,.0f} lb).<br>"
                        f"<b>{rest}</b> piece(s), {info['rest_lbs']:,.0f} lb, won't fit.")
        else:
            box.setText(f"<b>Truck {t}</b> has no room for {name}'s {len(rem)} piece(s) "
                        f"({info['rest_lbs']:,.0f} lb).")
        box.setInformativeText("What should happen with this stop?")
        box.setDetailedText("Won't fit:\n" + "\n".join(
            f"{pc.crate_label}  {pc.length_in:.0f}×{pc.width_in:.0f}×{pc.height_in:.0f}\"  "
            f"{pc.weight_lbs:,.0f} lb" for pc in rem[fit:]))
        choices = {}
        if not info["truck_empty"]:
            choices[box.addButton(f"Move all {len(rem)} pcs to a new truck",
                                  QMessageBox.ButtonRole.AcceptRole)] = "new_truck"
        else:
            choices[box.addButton(f"Load {fit}, put the other {rest} on another truck",
                                  QMessageBox.ButtonRole.AcceptRole)] = "continue"
        label = (f"Load {fit}, ship the other {rest} separately (UPS / hot-shot)" if fit
                 else f"Ship all {len(rem)} separately (UPS / hot-shot)")
        choices[box.addButton(label, QMessageBox.ButtonRole.ActionRole)] = "ship_other"
        box.setDefaultButton(next(iter(choices)))
        box.exec()
        return choices.get(box.clickedButton(), next(iter(choices.values())))

    def _build_plans(self, pieces: list, silent: bool):
        override = self.chk_override.isChecked()
        by_id = {pc.piece_id: pc for pc in pieces if pc.piece_id not in self._excluded_ids}
        self._plans = []
        n = len(self._trucks)
        for i, t in enumerate(self._trucks):
            tp = [by_id[x] for x in t["piece_ids"] if x in by_id]
            plan = plan_load(tp, override=override, stop_order=self._route({pc.stop for pc in tp}),
                             rig=t.get("rig"))
            plan.stop_names = list(self._stop_names)
            plan.truck_label = f"Truck {i + 1} of {n}" if n > 1 else ""
            plan.notices = []
            self._apply_manual_moves(plan, i + 1)
            self._plans.append(plan)
        self._cur_truck = min(self._cur_truck, max(0, n - 1))

        if not silent:
            unfit = [(i + 1, pc, why) for i, pl in enumerate(self._plans)
                     for pc, why in (pl.unplaced or [])]
            if unfit:
                wt = sum(pc.weight_lbs for _, pc, _ in unfit)
                QMessageBox.warning(
                    self, "Freight Doesn't Fit",
                    f"{len(unfit)} piece(s), {wt:,.0f} lb, don't fit on their truck:\n\n" +
                    "\n".join(f"  •  Truck {t}: {pc.crate_label} ({pc.weight_lbs:,.0f} lb) — {why}"
                              for t, pc, why in unfit) +
                    "\n\nUse Re-split Trucks or Move Stop… to rearrange, or ship them separately.")
            weight_v = [f"Truck {i + 1}: {v}" if n > 1 else v
                        for i, pl in enumerate(self._plans) for v in pl.violations
                        if not v.startswith("DOES NOT FIT")]
            if weight_v and not override:
                QMessageBox.warning(self, "Weight Violations",
                                    "Weight violations detected:\n\n" + "\n".join(weight_v) +
                                    "\n\nCheck 'Allow override' to proceed anyway, or adjust the load.")
        self._refresh_truck_bar()
        self._select_truck(self._cur_truck)
        self.btn_plan_pdf.setEnabled(True)
        self.btn_label_pdf.setEnabled(True)

    def _refresh_truck_bar(self):
        self.cmb_truck.blockSignals(True)
        self.cmb_truck.clear()
        for i, pl in enumerate(self._plans):
            pcs = {sl.piece.piece_id: sl.piece for sl in pl.slots if sl.piece}
            stops = ", ".join(f"S{sid}" for sid in plan_stops(pl))
            self.cmb_truck.addItem(f"Truck {i + 1} of {len(self._plans)}  ·  {stops or '—'}  ·  "
                                   f"{len(pcs)} pcs  ·  {sum(x.weight_lbs for x in pcs.values()):,.0f} lb")
        self.cmb_truck.setCurrentIndex(self._cur_truck)
        self.cmb_truck.blockSignals(False)
        self.btn_other.setText(f"Ship Separately ({len(self._ship_other)})")

    def _select_truck(self, idx: int):
        if not self._plans or idx < 0:
            return
        self._cur_truck = idx
        plan = self._plans[idx]
        self._plan = plan
        self.f_trailer_num.setText(self._trucks[idx].get("trailer_num", ""))
        self.trailer_view.set_plan(plan)
        self.side_view.set_plan(plan)
        self.view_3d.set_plan(plan)
        self._sync_rig_controls(plan)
        self._update_summary(plan)
        self._on_securement_changed()

    def _trailer_edited(self, text: str):
        if self._trucks and self._cur_truck < len(self._trucks):
            self._trucks[self._cur_truck]["trailer_num"] = text.strip()

    def _move_stop(self):
        pieces = [pc for pc in self._read_pieces() if pc.piece_id not in self._excluded_ids]
        stops = sorted({pc.stop for pc in pieces})
        if not stops:
            return
        dlg = QDialog(self)
        dlg.setWindowTitle("Move Stop")
        form = QFormLayout(dlg)
        cmb_stop = QComboBox()
        for sid in stops:
            cmb_stop.addItem(f"{self._stop_text(sid)}  ({sum(1 for pc in pieces if pc.stop == sid)} pcs)", sid)
        cmb_dest = QComboBox()
        for i in range(len(self._trucks)):
            cmb_dest.addItem(f"Truck {i + 1}", i)
        cmb_dest.addItem("New truck", "new")
        cmb_dest.addItem("Ship separately (UPS / hot-shot)", "other")
        form.addRow("Stop:", cmb_stop)
        form.addRow("Move all of its pieces to:", cmb_dest)
        btns = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok |
                                QDialogButtonBox.StandardButton.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        form.addRow(btns)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        sid, dest = cmb_stop.currentData(), cmb_dest.currentData()
        ids = [pc.piece_id for pc in pieces if pc.stop == sid]
        self._move_pieces(ids, dest)

    def _move_pieces(self, ids: list, dest):
        """dest: truck index, "new" or "other"."""
        for t in self._trucks:
            t["piece_ids"] = [i for i in t["piece_ids"] if i not in ids]
        self._ship_other = [i for i in self._ship_other if i not in ids]
        if dest == "other":
            self._ship_other += ids
            for i in ids:
                self._ship_reasons.setdefault(i, "moved here by you")
        else:
            if dest == "new" or dest >= len(self._trucks):
                self._trucks.append({"trailer_num": "", "piece_ids": []})
                dest = len(self._trucks) - 1
            self._trucks[dest]["piece_ids"] += ids
            for i in ids:
                self._ship_reasons.pop(i, None)
            self._cur_truck = dest
        self._trucks = [t for t in self._trucks if t["piece_ids"]] or \
            [{"trailer_num": "", "piece_ids": []}]
        self._cur_truck = min(self._cur_truck, len(self._trucks) - 1)
        self._build_plans(self._read_pieces(), silent=False)

    def _show_ship_other(self):
        by_id = {pc.piece_id: pc for pc in self._read_pieces()}
        pcs = [by_id[i] for i in self._ship_other if i in by_id]
        if not pcs:
            QMessageBox.information(self, "Ship Separately", "Every piece is on a truck.")
            return
        dlg = ShipSeparatelyDialog(pcs, self._ship_reasons, len(self._trucks), self._stop_text, self)
        dlg.exec()
        if dlg.returned:
            dest = dlg.returned[0][1]
            self._move_pieces([pid for pid, _ in dlg.returned],
                              "new" if dest >= len(self._trucks) else dest)

    def _shipment_dict(self) -> dict:
        return {"origin": self._origin, "stops": self._stops, "use_road": self._use_road,
                "road": self._road, "trucks": self._trucks, "ship_other": self._ship_other,
                "ship_reasons": {str(k): v for k, v in self._ship_reasons.items()},
                "split_key": {str(k): v for k, v in (self._split_key or {}).items()},
                "manual": {str(k): list(v) for k, v in self._manual.items()},
                "rig_specs": dict(RIG)}

    def _update_summary(self, plan: LoadPlan):
        aw = plan.axle_weights
        rows_html = ""
        if len(self._plans) > 1:
            rows_html += (f"<tr><td colspan='3'><b style='color:#1f4e79'>"
                          f"{getattr(plan, 'truck_label', '')}</b></td></tr>")
        for key, label, lim_key in [
            ("steer",   "Steer",          "steer_limit"),
            ("drive",   "Drive",          "drive_limit"),
            ("trailer", "Trailer Tandem", "trailer_limit"),
        ]:
            val   = aw.get(key, 0)
            limit = TRAILER[lim_key]
            pct   = val / limit * 100 if limit else 0
            color = "#c0392b" if val > limit else "#27ae60"
            rows_html += (f"<tr>"
                          f"<td><b>{label}:</b></td>"
                          f"<td align='right'><span style='color:{color}'>{val:,} lbs</span></td>"
                          f"<td align='right'>&nbsp;{pct:.0f}% of {limit:,}</td>"
                          f"</tr>")
        gross, glim = aw.get("gross", 0), TRAILER["gross_limit"]
        rows_html += (f"<tr><td><b>Gross:</b></td>"
                      f"<td align='right'><span style='color:{'#c0392b' if gross > glim else '#27ae60'}'>"
                      f"{gross:,} lbs</span></td>"
                      f"<td align='right'>&nbsp;{gross / glim * 100 if glim else 0:.0f}% of {glim:,}</td></tr>")
        rows_html += (f"<tr><td><b>Freight:</b></td>"
                      f"<td align='right'>{aw.get('freight_gross', 0):,} lbs</td>"
                      f"<td align='right' style='color:#555'>&nbsp;empty rig {rig_tare():,.0f}</td></tr>")
        if plan.rig and aw:
            rows_html += (f"<tr><td colspan='3' style='color:#555'>"
                          f"{rig_text(plan.rig)}{'' if plan.rig_auto else ' (set by you)'}</td></tr>")

        # Side weights: split the cargo by its lateral centre of gravity
        W_in = TRAILER["interior_width_in"]
        boxes = _crate_boxes(plan)
        total = sum(b[0].weight_lbs for b in boxes) or 1.0
        cg_off = sum(b[0].weight_lbs * (b[2] + b[5] / 2 - W_in / 2) for b in boxes) / total
        left_w = round(total * (0.5 - cg_off / W_in))
        right_w = round(total - left_w)
        diff    = abs(left_w - right_w)
        heavier = "Left" if left_w > right_w else "Right"
        diff_color = "#c0392b" if diff > 2000 else "#e67e22" if diff > 500 else "#27ae60"
        rows_html += (
            f"<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
            f"<tr><td><b>Left (A):</b></td><td align='right'>{left_w:,} lbs</td><td></td></tr>"
            f"<tr><td><b>Right (B):</b></td><td align='right'>{right_w:,} lbs</td><td></td></tr>"
            f"<tr><td><b>Side diff:</b></td>"
            f"<td align='right'><span style='color:{diff_color}'>{diff:,} lbs</span></td>"
            f"<td align='right'><span style='color:{diff_color}'>{heavier} heavier</span></td>"
            f"</tr>"
        )
        if plan.unplaced:
            nl_wt = sum(p.weight_lbs for p, _ in plan.unplaced)
            names = [p.crate_label for p, _ in plan.unplaced]
            shown = ", ".join(names[:8]) + (f" +{len(names) - 8} more" if len(names) > 8 else "")
            rows_html += (
                f"<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
                f"<tr><td><b style='color:#c0392b'>NOT LOADED:</b></td>"
                f"<td align='right' style='color:#c0392b'>{len(plan.unplaced)} pc, {nl_wt:,.0f} lbs</td>"
                f"<td align='right' style='color:#c0392b'>"
                f"{shown}</td></tr>"
            )
        if self._plan:
            n_stacked = sum(1 for s in self._plan.slots if s.tier == 1 and s.piece)
            if n_stacked:
                stacked_labels = ", ".join(
                    s.piece.crate_label for s in self._plan.slots
                    if s.tier == 1 and s.piece)
                rows_html += (
                    f"<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
                    f"<tr><td><b>Stacked:</b></td>"
                    f"<td align='right'>{n_stacked} crate(s)</td>"
                    f"<td align='right' style='color:#8B6914'>{stacked_labels}</td></tr>"
                )
        stops = plan_stops(plan)
        if len(stops) > 1:
            rows_html += "<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
            for st in stops:
                pcs = [sl.piece for sl in plan.slots if sl.piece and sl.piece.stop == st
                       and not (sl.is_center and sl.col == 1)]
                color = STOP_COLORS[(st - 1) % len(STOP_COLORS)]
                rows_html += (f"<tr><td><b style='color:{color}'>{stop_label(plan, st)}"
                              f"{' (doors)' if st == stops[0] else ''}:</b></td>"
                              f"<td align='right'>{len(pcs)} pc</td>"
                              f"<td align='right'>{sum(pc.weight_lbs for pc in pcs):,.0f} lbs</td></tr>")
            for delivered, steer, drive, trailer in plan.stop_axles or []:
                bad = (drive > TRAILER["drive_limit"] or trailer > TRAILER["trailer_limit"]
                       or steer > TRAILER["steer_limit"])
                c = "#c0392b" if bad else "#555"
                rows_html += (f"<tr><td style='color:{c}'>After S{delivered}:</td>"
                              f"<td colspan='2' align='right' style='color:{c}'>"
                              f"Drive {drive:,} · Trailer {trailer:,}</td></tr>")
        sec_text = self._securement_text()
        if sec_text:
            rows_html += (
                f"<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
                f"<tr><td><b>Securement:</b></td>"
                f"<td colspan='2' align='right'>{sec_text}</td></tr>"
            )
        if len(self._plans) > 1 or self._ship_other:
            by_id = {pc.piece_id: pc for pc in self._read_pieces()}
            other_wt = sum(by_id[i].weight_lbs for i in self._ship_other if i in by_id)
            rows_html += (f"<tr><td colspan='3'><hr style='margin:2px'></td></tr>"
                          f"<tr><td><b>Shipment:</b></td><td colspan='2' align='right'>"
                          f"{len(self._plans)} truck(s)"
                          + (f" · <span style='color:#b03a2e'>{len(self._ship_other)} pcs, "
                             f"{other_wt:,.0f} lb ship separately</span>" if self._ship_other else "")
                          + "</td></tr>")
        self.lbl_summary.setText(f"<table>{rows_html}</table>")

        if plan.violations:
            v_html = "<br>".join(f"<span style='color:red'>⚠ {v}</span>" for v in plan.violations)
            if plan.override:
                v_html += "<br><i>Override active.</i>"
        else:
            v_html = "<span style='color:green'>✓ All axle weights within legal limits.</span>"
        moved = sum(1 for sl in plan.slots if sl.tier == 0 and sl.x_in is not None
                    and not (sl.is_center and sl.col == 1))
        if moved:
            v_html += (f"<br><span style='color:#555'>{moved} crate(s) positioned by hand — "
                       "Reset Moves (Top View) puts them back.</span>")
        for note in getattr(plan, "notices", []):
            v_html += f"<br><span style='color:#7a5c12'>ⓘ {note}</span>"
        self.lbl_status.setText(v_html)

    # ── File: New / Open / Save ───────────────────────────────────────────────
    def _new_plan(self):
        if QMessageBox.question(self, "New Plan", "Clear everything and start a new plan?",
                                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                                ) != QMessageBox.StandardButton.Yes:
            return
        self.table.setRowCount(0)
        self._next_id = 1
        self._plan = None
        self._excluded_ids = []
        self._securement.clear()
        self.lbl_securement.setText("")
        self._current_file = None
        for f in [self.f_customer, self.f_shipment_name, self.f_trailer_num, self.f_bol]:
            f.clear()
        self.f_date.setText(datetime.date.today().strftime("%Y-%m-%d"))
        self._origin = {"name": "", "location": "", "lat": None, "lon": None}
        self._stops = [{"name": "", "location": "", "lat": None, "lon": None}
                       for _ in range(MAX_STOPS)]
        self._use_road, self._road = False, None
        self._trucks, self._ship_other, self._ship_reasons = [], [], {}
        self._split_key, self._plans, self._cur_truck = None, [], 0
        self._manual = {}
        self.cmb_truck.clear()
        self.btn_other.setText("Ship Separately (0)")
        self._refresh_stop_label()
        for key, default in [("steer", 12000), ("drive", 34000), ("trailer", 34000),
                             ("gross", 80000)]:
            self._limit_fields[key].setValue(default)
        self.chk_override.setChecked(False)
        self.lbl_summary.setText("—")
        self.lbl_status.setText("No plan calculated.")
        self.trailer_view.set_plan(None)
        self.side_view.set_plan(None)
        self.view_3d.set_plan(None)
        load_rig_settings()          # an opened plan may have brought its own truck specs
        self._sync_rig_controls(None)
        self.btn_plan_pdf.setEnabled(False)
        self.btn_label_pdf.setEnabled(False)
        self._set_title()

    def _open_plan(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Open Load Plan", self._save_dir(), "Load Plan Files (*.json)")
        if not path:
            return
        try:
            pieces, info, limits, excluded_ids, override, securement, shipment = load_plan_json(path)
        except Exception as e:
            QMessageBox.critical(self, "Open Failed", str(e))
            return

        # Populate shipment info
        self.f_customer.setText(info.customer)
        self.f_shipment_name.setText(info.shipment_name)
        self.f_trailer_num.setText(info.trailer_num)
        self.f_bol.setText(info.bol_num)
        self.f_date.setText(info.date)
        blank = {"name": "", "location": "", "lat": None, "lon": None}
        self._origin = dict(blank, **shipment.get("origin", {}))
        saved_stops = shipment.get("stops") or [dict(blank, name=n) for n in info.stop_names]
        self._stops = [dict(blank, **(saved_stops[i] if i < len(saved_stops) else {}))
                       for i in range(MAX_STOPS)]
        self._use_road = bool(shipment.get("use_road", False))
        self._road = shipment.get("road")
        self._trucks = shipment.get("trucks", [])
        self._ship_other = list(shipment.get("ship_other", []))
        self._ship_reasons = {int(k): v for k, v in shipment.get("ship_reasons", {}).items()}
        self._manual = {int(k): tuple(v) for k, v in shipment.get("manual", {}).items()}
        specs = shipment.get("rig_specs")
        if specs:
            RIG.update({k: v for k, v in specs.items() if k in RIG_DEFAULTS})
        else:
            load_rig_settings()
        sk = shipment.get("split_key")
        self._split_key = {int(k): v for k, v in sk.items()} if sk else None
        self._cur_truck = 0
        self._refresh_stop_label()

        # Populate axle limits
        for key in ("steer", "drive", "trailer", "gross"):
            if key in limits:
                self._limit_fields[key].setValue(limits[key])
        self.chk_override.setChecked(override)

        # Populate freight table
        self._load_pieces(pieces)
        self._excluded_ids = excluded_ids
        self._securement[:] = securement
        self._on_securement_changed()

        # Recalculate and display
        self._current_file = path
        self._set_title()
        self._calculate()

    def _save_plan(self):
        if self._current_file:
            self._write_plan(self._current_file)
        else:
            self._save_plan_as()

    def _save_plan_as(self):
        info = self._get_shipment_info()
        default_name = make_base_filename(info) + ".json"
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Load Plan", os.path.join(self._save_dir(), default_name),
            "Load Plan Files (*.json)")
        if path:
            self._write_plan(path)
            self._current_file = path
            self._set_title()

    def _write_plan(self, path: str):
        pieces = self._read_pieces()
        if not pieces and self.table.rowCount() > 0:
            return   # parse error already shown by _read_pieces
        info = self._get_shipment_info()
        limits = {k: self._limit_fields[k].value() for k in ("steer", "drive", "trailer", "gross")}
        try:
            save_plan_json(pieces, info, limits, self._excluded_ids,
                           self.chk_override.isChecked(), path, self._securement,
                           self._shipment_dict())
            self._remember_dir(path)
            self.lbl_status.setText(
                self.lbl_status.text() +
                f"<br><span style='color:#555'>Saved: {os.path.basename(path)}</span>")
        except Exception as e:
            QMessageBox.critical(self, "Save Failed", str(e))

    def _set_title(self):
        name = os.path.basename(self._current_file) if self._current_file else "Unsaved"
        self.setWindowTitle(f"Freight Loader v{APP_VERSION} — {name}")

    # ── PDF exports ───────────────────────────────────────────────────────────
    # ── Save folder (remembered between sessions) ────────────────────────────
    def _save_dir(self) -> str:
        """Folder of the user's last save, or Load Plans if none (or it's gone)."""
        last = QSettings("FreightLoader", "FreightLoader").value("last_save_dir", "")
        return last if last and os.path.isdir(last) else LOAD_PLANS_DIR

    def _remember_dir(self, path: str):
        QSettings("FreightLoader", "FreightLoader").setValue("last_save_dir",
                                                             os.path.dirname(path))

    def _preview_and_save(self, title: str, default_name: str, build):
        """Build the PDF into a temp file, preview it, then save where the user picks."""
        fd, tmp = tempfile.mkstemp(suffix=".pdf", prefix="freight_loader_")
        os.close(fd)
        try:
            build(tmp)
            if HAS_QTPDF:
                dlg = PdfPreviewDialog(tmp, title, self)
                dlg.exec()
                if not dlg.save_requested:
                    return
            else:   # no PDF module: open the system viewer, then offer to save
                if sys.platform == "win32":
                    os.startfile(tmp)
                if QMessageBox.question(self, title, "Save this PDF?") != \
                        QMessageBox.StandardButton.Yes:
                    return
            path, _ = QFileDialog.getSaveFileName(
                self, f"Save {title}", os.path.join(self._save_dir(), default_name),
                "PDF Files (*.pdf)")
            if not path:
                return
            if not path.lower().endswith(".pdf"):
                path += ".pdf"
            shutil.copyfile(tmp, path)
            self._remember_dir(path)
            QMessageBox.information(self, "Saved", f"{title} saved:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "PDF Error", str(e))
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass

    def _export_floor_plan(self):
        if not self._plan:
            return
        info = self._get_shipment_info()
        by_id = {pc.piece_id: pc for pc in self._read_pieces()}
        other = [(by_id[i], self._stop_text(by_id[i].stop),
                  self._ship_reasons.get(i, "didn't fit — ship separately"))
                 for i in self._ship_other if i in by_id]
        entries = self._truck_entries(info)
        self._preview_and_save(
            "Floor Plan PDF", make_base_filename(info) + "_load_plan.pdf",
            lambda path: export_shipment_pdf(entries, path, self._securement,
                                             include_3d=self.chk_pdf_3d.isChecked(),
                                             ship_other=other))

    def _truck_entries(self, info: ShipmentInfo) -> list:
        """[(plan, info)] for every truck, each with its own trailer # and truck label."""
        out = []
        for i, pl in enumerate(self._plans):
            ti = ShipmentInfo(**{**vars(info),
                                 "trailer_num": self._trucks[i].get("trailer_num", ""),
                                 "truck_label": getattr(pl, "truck_label", "")})
            out.append((pl, ti))
        return out or [(self._plan, info)]

    def _export_labels(self):
        if not self._plan:
            return
        info = self._get_shipment_info()
        entries = self._truck_entries(info)
        self._preview_and_save("Crate Labels PDF", make_base_filename(info) + "_labels.pdf",
                               lambda path: export_shipment_labels(entries, path))


# ── Entry point ───────────────────────────────────────────────────────────────
def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    ensure_folders()
    load_rig_settings()
    win = FreightLoaderApp()
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
